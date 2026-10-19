from pathlib import Path

import openpyxl

from foresight_os.analysis.workbook_inputs import (
    MaterialView,
    WorkbookLayoutSpec,
    build_material_view,
    resolve_workbook_layout,
)
from foresight_os.analysis.workbook_engine import DomainAnalysis, WorkbookAnalysisEngine
from foresight_os.domain.analysis_inputs import AnalysisInputContract
from foresight_os.runtime.local_toolbelt import SheetSelection, spreadsheet_read


def _workbook(tmp_path: Path) -> Path:
    source = tmp_path / "material-view.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Entity", "Value", "Scenario"])
    sheet.append(["North", 10, "Base"])
    sheet.append(["South", 20, "Base"])
    workbook.save(source)
    return source


def test_material_view_resolves_layout_and_retains_row_provenance(tmp_path):
    source = _workbook(tmp_path)
    contract = AnalysisInputContract(
        domain="test",
        required_tables=("Data",),
        entity_scope=("North",),
        scenarios=("Base",),
        workbook_selection={"sheet": "Data"},
    )
    workbook = spreadsheet_read(source, SheetSelection(sheet_names=("Data",)))
    layout = resolve_workbook_layout(workbook, WorkbookLayoutSpec.from_contract(contract))
    material = build_material_view(workbook, layout, contract)

    assert isinstance(material, MaterialView)
    assert material.table("Data").rows == (("North", 10, "Base"),)
    assert material.table("Data").row_provenance[0][0].workbook_sha256 == workbook.workbook_sha256
    assert material.scope_fingerprint
    assert material.layout_fingerprint


def test_material_view_changes_scope_without_recapturing_source(tmp_path):
    source = _workbook(tmp_path)
    first_contract = AnalysisInputContract(domain="test", required_tables=("Data",), entity_scope=("North",))
    second_contract = AnalysisInputContract(domain="test", required_tables=("Data",), entity_scope=("South",))
    workbook = spreadsheet_read(source, SheetSelection(sheet_names=("Data",)))
    layout = resolve_workbook_layout(workbook, WorkbookLayoutSpec.from_contract(first_contract))

    first = build_material_view(workbook, layout, first_contract)
    second = build_material_view(workbook, layout, second_contract)

    assert first.source_sha256 == second.source_sha256 == workbook.workbook_sha256
    assert first.table("Data").rows != second.table("Data").rows
    assert first.scope_fingerprint != second.scope_fingerprint


def test_engine_returns_typed_domain_result_without_persistence_side_effects(tmp_path):
    source = _workbook(tmp_path)
    contract = AnalysisInputContract(domain="test", required_tables=("Data",))
    engine = WorkbookAnalysisEngine()

    result = engine.run(
        source,
        contract,
        lambda material, normalized: DomainAnalysis(
            result=material.table("Data").rows[0][1],
            diagnostics=(normalized.domain,),
        ),
    )

    assert result.status == "completed"
    assert result.domain_result is not None
    assert result.domain_result.result == 10
    assert result.diagnostics == ("test",)
    assert result.metrics["capture_count"] == 1


def test_engine_does_not_re_resolve_effective_contract_with_explicit_empty_tables(tmp_path):
    source = _workbook(tmp_path)
    contract = AnalysisInputContract(domain="test", required_tables=())

    result = WorkbookAnalysisEngine().run(
        source,
        contract,
        lambda material, normalized: DomainAnalysis(result=normalized.required_tables),
    )

    assert result.status == "completed"
    assert result.domain_result is not None
    assert result.domain_result.result == ()


def test_engine_maps_missing_selection_to_blocked_input(tmp_path):
    source = _workbook(tmp_path)
    contract = AnalysisInputContract(
        domain="test",
        required_tables=("Missing",),
        workbook_selection={"sheet": "Missing"},
    )

    result = WorkbookAnalysisEngine().run(source, contract, lambda material, normalized: material)

    assert result.status == "blocked_input"
    assert "missing workbook sheet" in (result.failure or "")


def test_explicit_table_selection_fails_closed_when_table_is_missing(tmp_path):
    source = _workbook(tmp_path)
    contract = AnalysisInputContract(
        domain="test",
        required_tables=("Data",),
        workbook_selection={"tables": [["Data", "MissingTable"]]},
    )

    result = WorkbookAnalysisEngine().run(source, contract, lambda material, normalized: material)

    assert result.status == "blocked_input"
    assert "missing workbook table" in (result.failure or "")


def test_scope_requires_applicable_dimensions_and_filters_period_and_units(tmp_path):
    source = tmp_path / "scoped.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Entity", "Period", "Unit", "Value"])
    sheet.append(["North", "2025-02", "EUR", 1])
    sheet.append(["North", "2025-10", "EUR", 2])
    sheet.append(["North", "2025-02", "USD", 3])
    sheet.append([None, "2025-02", "EUR", 4])
    workbook.save(source)
    contract = AnalysisInputContract(
        domain="test",
        required_tables=("Data",),
        entity_scope=("North",),
        periods=("2025-01:2025-09",),
        units="EUR",
        workbook_selection={
            "sheet": "Data",
            "required_scope_dimensions": {"Data": ["entity", "period", "unit"]},
        },
    )

    result = WorkbookAnalysisEngine().run(source, contract, lambda material, normalized: material)

    assert result.status == "completed"
    assert result.material_view is not None
    assert result.material_view.table("Data").rows == (("North", "2025-02", "EUR", 1),)
    assert {reason for _, _, reason, _ in result.material_view.data_quality_rows} == {
        "out_of_scope_period",
        "out_of_scope_unit",
        "required_scope_entity_missing",
    }


def test_period_scope_requires_full_containment_and_rejects_coarse_rows(tmp_path):
    source = tmp_path / "periods.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Period", "Value"])
    sheet.append(["2025", 1])
    sheet.append(["2025-Q1", 2])
    sheet.append(["2025-02", 3])
    workbook.save(source)
    contract = AnalysisInputContract(
        domain="test",
        required_tables=("Data",),
        periods=("2025-01:2025-03",),
        workbook_selection={"sheet": "Data", "required_scope_dimensions": {"Data": ["period"]}},
    )

    result = WorkbookAnalysisEngine().run(source, contract, lambda material, normalized: material)

    assert result.status == "completed"
    assert result.material_view is not None
    assert result.material_view.table("Data").rows == (("2025-Q1", 2), ("2025-02", 3))


def test_declared_table_alias_compiles_physical_selection(tmp_path):
    source = tmp_path / "aliased.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Physical Data"
    sheet.append(["Value"])
    sheet.append([7])
    workbook.save(source)
    contract = AnalysisInputContract(
        domain="test",
        required_tables=("LogicalData",),
        workbook_selection={"table_aliases": {"LogicalData": "Physical Data"}},
    )

    result = WorkbookAnalysisEngine().run(source, contract, lambda material, normalized: material.table("LogicalData").rows[0][0])

    assert result.status == "completed"
    assert result.domain_result is not None
    assert result.domain_result.result == 7
