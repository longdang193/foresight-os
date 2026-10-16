from __future__ import annotations

from pathlib import Path

import openpyxl

from foresight_os.analysis.portfolio_resilience import (
    analyze_portfolio_resilience,
    run_portfolio_resilience_workflow,
)
from foresight_os.domain.challenge import derive_finding_views
from foresight_os.persistence.analysis_store import AnalysisStore


def _write_portfolio_fixture(path: Path) -> None:
    workbook = openpyxl.Workbook()
    brief = workbook.active
    brief.title = "Brief"
    brief.append(["Portfolio Resilience"])
    brief.append([])
    brief.append(["Decision question", "Which segments should be prioritized?"])
    brief.append(["Horizon", "2027–2029"])
    brief.append(["Required deliverables", "Baseline economics; 3 scenarios; sensitivity thresholds; competing options; conditional recommendation; monitoring triggers."])

    sales = workbook.create_sheet("Sales_History")
    sales.append(["Sales History"])
    sales.append([])
    sales.append(["Quarter", "Region", "Segment", "Units", "List_Price", "Discount_pct", "Variable_Cost_per_Unit", "Revenue", "Gross_Profit", "Currency", "Source_ID"])
    sales.append(["2026-Q1", "DACH", "Core Hardware", 100, 100, 0.10, 60, 9000, 3000, "EUR", "SRC-001"])
    sales.append(["2026-Q1", "DACH", "Core Hardware", 100, 100, 0.10, 60, 9000, 3000, "EUR", "SRC-001"])
    sales.append(["2026-Q2", "DACH", "Analytics Software", 100, 100, 0.10, 30, 9000, 6000, "EUR", "SRC-002"])
    sales.append(["2026-Q2", "DACH", "Service Contracts", 100, 80, 0.05, 20, 7600, 5600, "EUR", "SRC-003"])
    sales.append(["2026-Q3", "DACH", "Core Hardware", 100, 100, 0.10, None, 9000, None, "EUR", "SRC-004"])
    sales.append(["2026-Q4", "DACH", "Core Hardware", 100, 100, 15, 60, -9000, -15000, "EUR", "SRC-005"])
    sales.append(["2025-Q4", "DACH", "Analytics Software", 100, 100, 0.10, 30, 9000, 6000, "EUR", "SRC-006"])

    fixed = workbook.create_sheet("Fixed_Costs")
    fixed.append(["Fixed Costs"])
    fixed.append([])
    fixed.append(["Segment", "Annual_Fixed_Cost_EUR"])
    fixed.append(["Core Hardware", 1000])
    fixed.append(["Analytics Software", 1000])
    fixed.append(["Service Contracts", 1000])

    scenarios = workbook.create_sheet("Scenario_Assumptions")
    scenarios.append(["Scenario Assumptions"])
    scenarios.append([])
    scenarios.append(["Scenario", "Demand_Growth", "Cost_Inflation", "Discount_Shift", "Fixed_Cost_Change", "Narrative"])
    scenarios.append(["Downside", -0.08, 0.055, 0.025, 0.02, "Demand contraction"])
    scenarios.append(["Base", 0.04, 0.035, 0, 0, "Moderate growth"])
    scenarios.append(["Upside", 0.11, 0.025, -0.01, -0.01, "Stronger adoption"])

    options = workbook.create_sheet("Strategic_Options")
    options.append(["Strategic Options"])
    options.append([])
    options.append(["Option_ID", "Option", "Investment_EUR", "Time_to_Impact_Months", "Reversibility", "Rationale"])
    options.append(["A", "Defend Core Hardware", 1000, 18, "Low", "Protect installed base"])
    options.append(["B", "Scale Analytics Software", 1200, 10, "Medium", "Higher growth"])
    options.append(["C", "Expand Service Contracts", 800, 8, "High", "Recurring revenue"])

    kpis = workbook.create_sheet("Reported_KPIs")
    kpis.append(["Reported KPIs"])
    kpis.append([])
    kpis.append(["KPI", "Reported_Value", "Unit", "Period", "Source_Note", "Validation_Task"])
    kpis.append(["2026 Revenue", 25000, "EUR", "2026", "Management pack", "Recalculate"])
    kpis.append(["2026 Gross Profit", 15000, "EUR", "2026", "Management pack", "Recalculate"])
    kpis.append(["Analytics Software CAGR", 0.20, "%", "2025–2026", "Strategy slide", "Recalculate"])
    kpis.append(["Service Contract renewal proxy", 0.91, "%", "2026", "Commercial estimate", "Treat as assumption"])

    workbook.save(path)


def test_portfolio_recipe_detects_anomalies_and_builds_required_analysis(tmp_path: Path):
    source = tmp_path / "portfolio.xlsx"
    _write_portfolio_fixture(source)

    analysis = analyze_portfolio_resilience(source)

    assert {item.kind for item in analysis.data_quality_issues} == {
        "DUPLICATE_ROW",
        "MISSING_COST",
        "MALFORMED_DISCOUNT",
        "NEGATIVE_REVENUE",
    }
    assert {item.scenario for item in analysis.scenarios} == {"Downside", "Base", "Upside"}
    assert analysis.sensitivity_thresholds
    assert len(analysis.options) == 3
    assert {item.kpi for item in analysis.kpi_reconciliations} == {
        "2026 Revenue",
        "2026 Gross Profit",
        "Analytics Software CAGR",
        "Service Contract renewal proxy",
    }
    assert next(item for item in analysis.kpi_reconciliations if item.kpi == "Analytics Software CAGR").status == "MISMATCH"
    assert next(item for item in analysis.kpi_reconciliations if item.kpi == "Service Contract renewal proxy").status == "ASSUMPTION_ONLY"
    assert any("do not silently repair" in item.lower() for item in analysis.assumptions)


def test_portfolio_workflow_persists_reviewed_finding_and_cited_report(tmp_path: Path):
    source = tmp_path / "portfolio.xlsx"
    _write_portfolio_fixture(source)
    store = AnalysisStore(tmp_path / "case")

    result = run_portfolio_resilience_workflow(store, source, human_decision="ACCEPT")

    assert result.status == "accepted"
    assert result.report_path is not None and result.report_path.exists()
    persisted = store.load_analysis()
    findings = derive_finding_views(persisted)
    assert findings
    assert any(item.accepted and item.review_status == "CURRENT" for item in findings)
    report = result.report_path.read_text(encoding="utf-8")
    assert "Baseline economics" in report
    assert "Sensitivity thresholds" in report
    assert "Alternative futures" in report
    assert "KPI reconciliation" in report
    assert "Monitoring triggers" in report
    assert "Recommended action" in report
    assert "Sales_History" in report
    assert "Counterevidence" in report


def test_portfolio_reject_policy_blocks_workflow_acceptance(tmp_path: Path):
    source = tmp_path / "portfolio.xlsx"
    _write_portfolio_fixture(source)
    store = AnalysisStore(tmp_path / "case")

    result = run_portfolio_resilience_workflow(
        store,
        source,
        human_decision="ACCEPT",
        input_contract={"domain": "portfolio_resilience", "duplicate_policy": "reject"},
    )

    assert result.status == "failed"
    assert result.failure and "duplicate policy reject" in result.failure

    missing_value_result = run_portfolio_resilience_workflow(
        store,
        source,
        human_decision="ACCEPT",
        input_contract={"domain": "portfolio_resilience", "missing_value_policy": "reject"},
    )

    assert missing_value_result.status == "failed"
    assert missing_value_result.failure and "missing-value policy reject" in missing_value_result.failure


def test_portfolio_report_uses_requested_reporting_currency(tmp_path: Path):
    source = tmp_path / "portfolio.xlsx"
    _write_portfolio_fixture(source)
    store = AnalysisStore(tmp_path / "case")

    result = run_portfolio_resilience_workflow(
        store,
        source,
        human_decision="ACCEPT",
        input_contract={
            "domain": "portfolio_resilience",
            "reporting_currency": "USD",
            "conversion_rates": {"EUR": 2.0, "USD": 1.0},
        },
    )

    assert result.status == "accepted"
    report = result.report_path.read_text(encoding="utf-8") if result.report_path else ""
    assert "Reporting currency: USD" in report
    assert "USD " in report
    assert "€" not in report
