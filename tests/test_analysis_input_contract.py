from pathlib import Path
from dataclasses import replace

import pytest

from foresight_os.analysis.supplier_resilience import analyze_supplier_resilience
from foresight_os.analysis.portfolio_resilience import analyze_portfolio_resilience
from foresight_os.analysis.powerbi_strategic_portfolio import analyze_powerbi_strategic_portfolio
from foresight_os.analysis.technology_radar import analyze_technology_radar
from foresight_os.domain.analysis_inputs import AnalysisInputContract, normalize_work_brief
from foresight_os.domain.models import WorkBrief, WorkStage
from foresight_os.domain.validation import ValidationError, validate_work_brief
from foresight_os.runtime.local_toolbelt import ToolbeltError
from foresight_os.persistence.analysis_store import AnalysisStore


def _brief(**analysis_inputs):
    return WorkBrief(
        "contract-test",
        "Which inputs should be assessed?",
        [],
        ["analysis"],
        [],
        [],
        "assess",
        "execute",
        analysis_inputs=analysis_inputs,
    )


def test_analysis_input_contract_round_trips_and_fingerprints_deterministically():
    contract = AnalysisInputContract(
        domain="supplier_resilience",
        required_tables=("Supplier_Master",),
        periods=("2026",),
        reporting_currency="EUR",
        conversion_rates={"USD": 0.92, "EUR": 1.0},
        assumptions=("legacy compatibility",),
    )
    value = validate_work_brief(_brief(supplier_resilience=contract.to_dict()))
    normalized = normalize_work_brief(value)

    assert normalized.contract_for_domain("supplier_resilience") == contract
    assert normalized.contract_for_domain("supplier_resilience").fingerprint == contract.fingerprint
    assert validate_work_brief(value)["analysis_inputs"][0]["domain"] == "supplier_resilience"


def test_sparse_declaration_survives_validation_persistence_and_workbook_resolution(tmp_path):
    sparse = {"domain": "supplier_resilience", "entity_scope": ["Cetus Electronics"]}
    brief = _brief(supplier_resilience={"entity_scope": ["Cetus Electronics"]})

    validated = validate_work_brief(brief)
    assert validated["analysis_inputs"] == [sparse]

    store = AnalysisStore(tmp_path / "case")
    store.save_brief(brief)
    restored = store.load_brief()
    assert restored.to_dict()["analysis_inputs"] == [sparse]

    normalized = normalize_work_brief(restored)
    declaration = normalized.declared_contract_for_domain("supplier_resilience")
    assert declaration is not None
    assert declaration.to_dict() == sparse

    from foresight_os.analysis.workbook_inputs import resolve_workbook_contract

    effective = resolve_workbook_contract("supplier_resilience", declaration)
    assert effective.entity_scope == ("Cetus Electronics",)
    assert effective.required_tables


def test_invalid_contract_policy_is_rejected_before_analysis():
    with pytest.raises((ValidationError, ValueError), match="duplicate policy"):
        validate_work_brief(_brief(supplier_resilience={"domain": "supplier_resilience", "duplicate_policy": "invented"}))


def test_unknown_analysis_input_contract_fields_are_rejected_before_normalization():
    with pytest.raises(ValidationError, match="unsupported fields: entity_scoep"):
        validate_work_brief(_brief(supplier_resilience={"domain": "supplier_resilience", "entity_scoep": ["Power Electronics"]}))


def test_analysis_input_contract_freezes_mappings_and_validates_periods():
    contract = AnalysisInputContract(
        domain="supplier_resilience",
        periods=("2026-01:2026-09",),
        column_mappings={"Supplier": "Supplier_ID"},
        conversion_rates={"EUR": 1.0},
    )

    with pytest.raises(TypeError):
        contract.column_mappings["Supplier"] = "Other"

    with pytest.raises(ValueError, match="invalid period"):
        AnalysisInputContract(domain="supplier_resilience", periods=("2026-13",))


@pytest.mark.parametrize("period", ("2026-09:2026-03", "2026-03:2025-09"))
def test_analysis_input_contract_rejects_reversed_period_ranges(period):
    with pytest.raises(ValueError, match="period range"):
        AnalysisInputContract(domain="supplier_resilience", periods=(period,))


def test_period_spec_preserves_cross_year_range_meaning():
    contract = AnalysisInputContract(domain="supplier_resilience", periods=("2025-10:2026-03",))

    spec = contract.period_specs[0]

    assert (spec.start_year, spec.start_month) == (2025, 10)
    assert (spec.end_year, spec.end_month) == (2026, 3)
    assert spec.granularity == "range"
    assert spec.is_cross_year


@pytest.mark.parametrize(
    ("period", "start", "end", "granularity"),
    (
        ("2026", (2026, 1), (2026, 12), "year"),
        ("2026-01", (2026, 1), (2026, 1), "month"),
        ("2026-Q1", (2026, 1), (2026, 3), "quarter"),
        ("2025-01:2026-03", (2025, 1), (2026, 3), "range"),
        ("2025-Q4:2026-Q2", (2025, 10), (2026, 6), "range"),
    ),
)
def test_period_spec_normalizes_year_month_quarter_and_range(period, start, end, granularity):
    spec = AnalysisInputContract(domain="supplier_resilience", periods=(period,)).period_specs[0]

    assert spec.start_key == start
    assert spec.end_key == end
    assert spec.granularity == granularity


def test_contract_resolver_fails_closed_for_missing_stage_binding():
    normalized = normalize_work_brief({
        "decision_question": "Question",
        "stages": [{"stage_id": "stage-a", "objective": "Assess A"}],
        "analysis_inputs": [{"domain": "shared", "stage_id": "stage-a"}],
    })

    with pytest.raises(ValueError, match="missing analysis input contract"):
        normalized.resolve_contract_for_domain("other", "stage-a")


def test_contract_resolver_fails_closed_for_ambiguous_manual_intent():
    contract = AnalysisInputContract(domain="shared")
    normalized = normalize_work_brief({
        "decision_question": "Question",
        "stages": [{"stage_id": "stage-a", "objective": "Assess A"}],
        "analysis_inputs": [{"domain": "shared", "stage_id": "stage-a"}],
    })
    object.__setattr__(normalized, "contracts", (contract, contract))

    with pytest.raises(ValueError, match="ambiguous analysis input contract"):
        normalized.resolve_contract_for_domain("shared", "stage-a")


def test_workbook_contract_resolution_preserves_omitted_vs_explicit_empty_tables():
    from foresight_os.analysis.workbook_inputs import resolve_workbook_contract

    omitted = resolve_workbook_contract("supplier_resilience", {"domain": "supplier_resilience"})
    explicit_empty = resolve_workbook_contract("supplier_resilience", {"domain": "supplier_resilience", "required_tables": []})

    assert omitted.required_tables
    assert explicit_empty.required_tables == ()


def test_normalized_intent_freezes_nested_policy_and_stage_values():
    normalized = normalize_work_brief({
        "decision_question": "Question",
        "horizon": "2030",
        "resource_limits": {"max_total_agent_calls": 4},
        "stages": [{"stage_id": "stage-1", "objective": "Assess", "limits": {"max_timeout_seconds": 30}}],
    })

    with pytest.raises(TypeError):
        normalized.resource_limits["max_total_agent_calls"] = 8
    with pytest.raises(TypeError):
        normalized.stages[0]["objective"] = "Changed"

    assert normalized.resource_limits["max_timeout_seconds"] == 30


def test_stage_scoped_contracts_prefer_exact_scope_and_preserve_same_domain_identity():
    normalized = normalize_work_brief({
        "decision_question": "Question",
        "horizon": "2030",
        "stages": [
            {"stage_id": "stage-a", "objective": "Assess A"},
            {"stage_id": "stage-b", "objective": "Assess B"},
        ],
        "analysis_inputs": [
            {"domain": "shared", "required_tables": ["Global"]},
            {"domain": "shared", "stage_id": "stage-a", "required_tables": ["A"]},
            {"domain": "shared", "stage_id": "stage-b", "required_tables": ["B"]},
        ],
    })

    assert normalized.contract_for_domain("shared", "stage-a").required_tables == ("A",)
    assert normalized.contract_for_domain("shared", "stage-b").required_tables == ("B",)
    assert normalized.contract_for_domain("shared", "unscoped").required_tables == ("Global",)


def test_duplicate_or_unknown_stage_scoped_contracts_fail():
    duplicate = {
        "decision_question": "Question",
        "stages": [{"stage_id": "stage-a", "objective": "Assess A"}],
        "analysis_inputs": [
            {"domain": "shared", "stage_id": "stage-a"},
            {"domain": "shared", "stage_id": "stage-a"},
        ],
    }
    unknown = {
        "decision_question": "Question",
        "stages": [{"stage_id": "stage-a", "objective": "Assess A"}],
        "analysis_inputs": [{"domain": "shared", "stage_id": "missing"}],
    }

    with pytest.raises(ValidationError, match="duplicate analysis input contract scope"):
        validate_work_brief(WorkBrief(
            "duplicate",
            "Question",
            [],
            [],
            [],
            [],
            "inspect",
            "execute",
            stages=[WorkStage("stage-a", "Assess A")],
            analysis_inputs=duplicate["analysis_inputs"],
        ))
    with pytest.raises(ValueError, match="unknown analysis input contract stage"):
        normalize_work_brief(unknown)


def test_stage_scoped_contracts_round_trip_through_brief_store(tmp_path):
    brief = WorkBrief(
        "round-trip",
        "Question",
        [],
        [],
        [],
        [],
        "inspect",
        "execute",
            stages=[WorkStage("stage-a", "Assess A"), WorkStage("stage-b", "Assess B")],
        analysis_inputs=[
            {"domain": "shared", "stage_id": "stage-a", "required_tables": ["A"]},
            {"domain": "shared", "stage_id": "stage-b", "required_tables": ["B"]},
        ],
    )
    store = AnalysisStore(tmp_path / "case")
    store.save_brief(brief)
    restored = store.load_brief()
    normalized = normalize_work_brief(restored)

    assert normalized.contract_for_domain("shared", "stage-a").required_tables == ("A",)
    assert normalized.contract_for_domain("shared", "stage-b").required_tables == ("B",)


def test_unsupported_domain_period_fails_before_workbook_capture(monkeypatch):
    import foresight_os.analysis.ev_capacity_investment as ev_module

    monkeypatch.setattr(ev_module.SourceSnapshot, "capture", lambda *_args, **_kwargs: pytest.fail("workbook capture should not run"))

    with pytest.raises(ValueError, match="unsupported period granularity"):
        ev_module.analyze_ev_capacity_investment(
            Path("tests/fixtures/exercise/03_EV_Capacity_Investment.xlsx"),
            AnalysisInputContract(domain="ev_capacity_investment", periods=("2030-01:2030-03",)),
        )


def test_explicit_column_mapping_to_missing_source_fails_before_row_defaults():
    from foresight_os.analysis.workbook_inputs import apply_column_mappings

    with pytest.raises(ToolbeltError, match="MISSING_MAPPED_COLUMN.*Missing"):
        apply_column_mappings([{"Actual": 1}], {"Canonical": "Missing"})


def test_explicit_column_mapping_does_not_fallback_to_existing_canonical_name():
    from foresight_os.analysis.workbook_inputs import apply_column_mappings

    with pytest.raises(ToolbeltError, match="MISSING_MAPPED_COLUMN.*Missing"):
        apply_column_mappings([{"Canonical": 1}], {"Canonical": "Missing"})


def test_explicit_incomplete_supplier_conversion_does_not_use_legacy_fx():
    source = Path("tests/fixtures/exercise/02_Supplier_Resilience.xlsx")
    analysis = analyze_supplier_resilience(
        source,
        AnalysisInputContract(
            domain="supplier_resilience",
            conversion_rates={"EUR": 1.0},
            reporting_currency="EUR",
        ),
    )

    assert any(issue.kind == "UNSUPPORTED_CONVERSION" for issue in analysis.data_quality_issues)


def test_legacy_supplier_contract_retains_documented_fx_assumption():
    analysis = analyze_supplier_resilience(Path("tests/fixtures/exercise/02_Supplier_Resilience.xlsx"))

    assert any("USD=0.92" in assumption for assumption in analysis.assumptions)


def test_explicit_empty_supplier_conversion_map_never_uses_legacy_rates():
    analysis = analyze_supplier_resilience(
        Path("tests/fixtures/exercise/02_Supplier_Resilience.xlsx"),
        AnalysisInputContract(domain="supplier_resilience", reporting_currency="EUR"),
    )

    assert any(issue.kind == "UNSUPPORTED_CONVERSION" for issue in analysis.data_quality_issues)


def test_supplier_entity_scope_and_reject_policy_are_enforced():
    source = Path("tests/fixtures/exercise/02_Supplier_Resilience.xlsx")
    scoped = analyze_supplier_resilience(source, AnalysisInputContract(domain="supplier_resilience", entity_scope=("DOES_NOT_EXIST",)))
    assert "EMPTY_ENTITY_SCOPE" in {issue.kind for issue in scoped.data_quality_issues}
    with pytest.raises(ValueError, match="duplicate policy reject"):
        analyze_supplier_resilience(source, AnalysisInputContract(domain="supplier_resilience", duplicate_policy="reject"))


def test_supplier_scenario_scope_filters_rows_and_reports_unknown_requests():
    source = Path("tests/fixtures/exercise/02_Supplier_Resilience.xlsx")
    contract = AnalysisInputContract(
        domain="supplier_resilience",
        scenarios=("Short outage",),
        reporting_currency="EUR",
        conversion_rates={"EUR": 1.0, "USD": 0.92, "CNY": 0.13},
    )
    scoped = analyze_supplier_resilience(source, contract)
    assert {item.scenario for item in scoped.scenarios} == {"Short outage"}

    unknown = analyze_supplier_resilience(source, replace(contract, scenarios=("NO_SUCH_SCENARIO",)))
    assert unknown.scenarios == ()
    assert any(issue.kind == "UNKNOWN_SCENARIO_SCOPE" for issue in unknown.data_quality_issues)


def test_technology_scenario_scope_filters_generated_scenarios_and_reports_unknown_requests():
    source = Path("tests/fixtures/exercise/04_Technology_Radar.xlsx")
    scoped = analyze_technology_radar(source, AnalysisInputContract(domain="technology_radar", scenarios=("Base",)))
    assert {item.scenario for item in scoped.scenarios} == {"Base"}

    unknown = analyze_technology_radar(source, AnalysisInputContract(domain="technology_radar", scenarios=("NO_SUCH_SCENARIO",)))
    assert unknown.scenarios == ()
    assert any(issue.kind == "UNKNOWN_SCENARIO_SCOPE" for issue in unknown.data_quality_issues)


def test_portfolio_column_mapping_and_reporting_units_reach_calculations(tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "portfolio.xlsx"
    source.write_bytes(Path("tests/fixtures/exercise/01_Portfolio_Resilience.xlsx").read_bytes())
    workbook = load_workbook(source)
    workbook["Sales_History"].cell(3, 1).value = "FiscalQuarter"
    workbook.save(source)

    analysis = analyze_portfolio_resilience(
        source,
        AnalysisInputContract(
            domain="portfolio_resilience",
            column_mappings={"Quarter": "FiscalQuarter"},
            reporting_currency="USD",
            units="USD",
            conversion_rates={"EUR": 2.0, "USD": 1.0},
        ),
    )

    assert sum(item.revenue for item in analysis.segment_baselines) > 0
    assert analysis.state.calculations[0].units == "USD"
    assert analysis.state.calculations[0].key.units == "USD"


@pytest.mark.parametrize("conversion_rates", [{}, {"GBP": 1.3}])
def test_portfolio_missing_conversion_does_not_relabel_eur_as_usd(conversion_rates):
    analysis = analyze_portfolio_resilience(
        Path("tests/fixtures/exercise/01_Portfolio_Resilience.xlsx"),
        AnalysisInputContract(domain="portfolio_resilience", reporting_currency="USD", conversion_rates=conversion_rates),
    )

    assert any(issue.kind == "UNSUPPORTED_CONVERSION" for issue in analysis.data_quality_issues)
    assert sum(item.revenue for item in analysis.segment_baselines) == 0


def test_portfolio_missing_canonical_column_fails_before_numeric_defaults(tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "portfolio.xlsx"
    source.write_bytes(Path("tests/fixtures/exercise/01_Portfolio_Resilience.xlsx").read_bytes())
    workbook = load_workbook(source)
    workbook["Sales_History"].cell(3, 8).value = "Revenue_USD"
    workbook.save(source)

    with pytest.raises(ToolbeltError, match="MISSING_MAPPED_COLUMN.*canonical column"):
        analyze_portfolio_resilience(source)


def test_powerbi_column_mapping_reaches_revenue_calculation(tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "powerbi.xlsx"
    source.write_bytes(Path("tests/fixtures/exercise/05_PowerBI_Strategic_Portfolio.xlsx").read_bytes())
    workbook = load_workbook(source)
    workbook["Fact_Performance"].cell(3, 9).value = "Revenue_USD"
    workbook.save(source)

    analysis = analyze_powerbi_strategic_portfolio(
        source,
        AnalysisInputContract(domain="powerbi_strategic_portfolio", column_mappings={"Revenue_EUR": "Revenue_USD"}),
    )

    assert analysis.revenue_2025_jan_sep > 0
    assert analysis.revenue_2026_jan_sep > 0


def test_powerbi_missing_column_mapping_does_not_fallback_to_canonical_data():
    with pytest.raises(ToolbeltError, match="MISSING_MAPPED_COLUMN.*Missing"):
        analyze_powerbi_strategic_portfolio(
            Path("tests/fixtures/exercise/05_PowerBI_Strategic_Portfolio.xlsx"),
            AnalysisInputContract(domain="powerbi_strategic_portfolio", column_mappings={"Revenue_EUR": "Missing"}),
        )


def test_powerbi_mapping_source_in_another_table_fails_for_consuming_table():
    with pytest.raises(ToolbeltError, match="MISSING_MAPPED_COLUMN.*Market_Type"):
        analyze_powerbi_strategic_portfolio(
            Path("tests/fixtures/exercise/05_PowerBI_Strategic_Portfolio.xlsx"),
            AnalysisInputContract(domain="powerbi_strategic_portfolio", column_mappings={"Revenue_EUR": "Market_Type"}),
        )


def test_powerbi_column_mapping_reaches_optional_refresh_timestamp_consumption(tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "powerbi.xlsx"
    source.write_bytes(Path("tests/fixtures/exercise/05_PowerBI_Strategic_Portfolio.xlsx").read_bytes())
    workbook = load_workbook(source)
    fact_sheet = workbook["Fact_Performance"]
    header_row = next(row for row in fact_sheet.iter_rows() if any(cell.value == "Refresh_Timestamp" for cell in row))
    refresh_column = next(index for index, cell in enumerate(header_row, start=1) if cell.value == "Refresh_Timestamp")
    header_row[refresh_column - 1].value = "UpdatedAt"
    dashboard_sheet = workbook["Dashboard_Snapshot"]
    dashboard_header_row = next(row for row in dashboard_sheet.iter_rows() if any(cell.value == "Refresh_Timestamp" for cell in row))
    dashboard_refresh_column = next(index for index, cell in enumerate(dashboard_header_row, start=1) if cell.value == "Refresh_Timestamp")
    dashboard_header_row[dashboard_refresh_column - 1].value = "UpdatedAt"
    first_data_row = header_row[0].row + 1
    fact_sheet.cell(first_data_row, refresh_column).value = "2026-11-01"
    workbook.save(source)

    analysis = analyze_powerbi_strategic_portfolio(
        source,
        AnalysisInputContract(domain="powerbi_strategic_portfolio", column_mappings={"Refresh_Timestamp": "UpdatedAt"}),
    )

    assert any(issue.kind == "STALE_REFRESH" for issue in analysis.data_quality_issues)


def test_portfolio_selected_scenarios_are_not_invented():
    analysis = analyze_portfolio_resilience(
        Path("tests/fixtures/exercise/01_Portfolio_Resilience.xlsx"),
        AnalysisInputContract(domain="portfolio_resilience", scenarios=("Base",)),
    )

    statement = next(item.statement for item in analysis.state.claims if item.id == "portfolio-claim-analytics-software")
    assert "Base" in statement
    assert "Downside" not in statement
    assert "Upside" not in statement


def test_powerbi_changed_period_narrative_follows_computed_direction():
    analysis = analyze_powerbi_strategic_portfolio(
        Path("tests/fixtures/exercise/05_PowerBI_Strategic_Portfolio.xlsx"),
        AnalysisInputContract(
            domain="powerbi_strategic_portfolio",
            baseline_period="2026",
            comparison_period="2025",
            periods=("2025-01:2025-09", "2026-01:2026-09"),
        ),
    )

    assert analysis.yoy_revenue_growth < 0
    assert "growth is negative" in analysis.state.implications[0].statement


def test_powerbi_same_year_period_specs_remain_distinct():
    analysis = analyze_powerbi_strategic_portfolio(
        Path("tests/fixtures/exercise/05_PowerBI_Strategic_Portfolio.xlsx"),
        AnalysisInputContract(
            domain="powerbi_strategic_portfolio",
            baseline_period="2025-01:2025-03",
            comparison_period="2025-04:2025-06",
            periods=("2025-01:2025-03", "2025-04:2025-06"),
        ),
    )

    assert analysis.baseline_period_label == "2025-01:2025-03"
    assert analysis.comparison_period_label == "2025-04:2025-06"
    assert analysis.revenue_2025_jan_sep != analysis.revenue_2026_jan_sep


def test_powerbi_explicit_periods_override_year_matching_configured_windows():
    analysis = analyze_powerbi_strategic_portfolio(
        Path("tests/fixtures/exercise/05_PowerBI_Strategic_Portfolio.xlsx"),
        AnalysisInputContract(
            domain="powerbi_strategic_portfolio",
            baseline_period="2025-01:2025-03",
            comparison_period="2026-04:2026-06",
            periods=("2025-01:2025-09", "2026-01:2026-09"),
        ),
    )

    assert analysis.baseline_period_label == "2025-01:2025-03"
    assert analysis.comparison_period_label == "2026-04:2026-06"
