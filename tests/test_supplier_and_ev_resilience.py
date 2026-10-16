from __future__ import annotations

from pathlib import Path

import openpyxl

from foresight_os.analysis.ev_capacity_investment import (
    analyze_ev_capacity_investment,
    run_ev_capacity_investment_workflow,
)
from foresight_os.analysis.supplier_resilience import (
    analyze_supplier_resilience,
    run_supplier_resilience_workflow,
)
from foresight_os.domain.challenge import derive_finding_views
from foresight_os.persistence.analysis_store import AnalysisStore


def _write_supplier_fixture(path: Path) -> None:
    workbook = openpyxl.Workbook()
    brief = workbook.active
    brief.title = "Brief"
    brief.append(["Supplier Resilience"])
    brief.append([])
    brief.append(["Decision question", "Does concentration create unacceptable exposure?"])
    brief.append(["Horizon", "12–36 months"])
    brief.append(["Required deliverables", "Concentration metrics; criticality; outage scenarios; sourcing alternatives; staged actions; early-warning indicators."])

    suppliers = workbook.create_sheet("Supplier_Master")
    suppliers.append(["Supplier Master"])
    suppliers.append([])
    suppliers.append(["Line_ID", "Supplier", "Country", "Part_Family", "Annual_Spend", "Currency", "Criticality_1_5", "Lead_Time_Days", "Defect_PPM", "Qualified_Alternatives", "Geo_Risk_1_5", "Unit_Cost"])
    suppliers.append(["SP-001", "Taiwan Power", "Taiwan", "Power Electronics", 1000000, "USD", 5, 70, 320, 0, 5, 20])
    suppliers.append(["SP-002", "German Power", "Germany", "Power Electronics", 400000, "EUR", 5, 42, 120, 2, 2, 85])
    suppliers.append(["SP-003", "Low Cost Risk", "China", "Sensors", 100000, "CNY", 4, 70, 450, 1, 4, 19])
    suppliers.append(["SP-004", "Alternate Sensors", "Germany", "Sensors", 200000, "EUR", 4, 28, 110, None, 2, 30])
    suppliers.append(["SP-002", "German Power", "Germany", "Power Electronics", 400000, "EUR", 5, 42, 120, 2, 2, 85])

    incidents = workbook.create_sheet("Incidents")
    incidents.append(["Incidents"])
    incidents.append([])
    incidents.append(["Incident_ID", "Date", "Supplier", "Part_Family", "Disruption_Days", "Expedite_Premium_pct", "Cause"])
    incidents.append(["INC-001", "2026-01-01", "Taiwan Power", "Power Electronics", 21, 0.15, "Port disruption"])
    incidents.append(["INC-002", "2026-02-01", "Alternate Sensors", "Sensors", 3, 0.05, "Quality hold"])

    scenarios = workbook.create_sheet("Scenario_Assumptions")
    scenarios.append(["Scenario Assumptions"])
    scenarios.append([])
    scenarios.append(["Scenario", "Outage_Days", "Expedite_Premium", "Demand_at_Risk_pct", "Narrative"])
    scenarios.append(["Short outage", 7, 0.08, 0.20, "One-week disruption"])
    scenarios.append(["Extended outage", 28, 0.22, 0.45, "Four-week disruption"])
    scenarios.append(["Correlated region shock", 56, 0.35, 0.75, "Regional shock"])

    options = workbook.create_sheet("Mitigation_Options")
    options.append(["Mitigation Options"])
    options.append([])
    options.append(["Action", "One_Time_Cost_EUR", "Lead_Time_Months", "Reversibility", "Description"])
    options.append(["Dual-source", 450000, 9, "Medium", "Qualify second supplier"])
    options.append(["Safety-stock", 280000, 2, "High", "Raise stock for critical parts"])
    options.append(["Supplier development", 220000, 12, "Medium", "Improve quality"])
    options.append(["Redesign for substitution", 900000, 18, "Low", "Reduce dependence"])
    workbook.save(path)


def _write_ev_fixture(path: Path) -> None:
    workbook = openpyxl.Workbook()
    brief = workbook.active
    brief.title = "Brief"
    brief.append(["EV Capacity Investment"])
    brief.append([])
    brief.append(["Decision question", "Should the company expand, pilot, outsource, or defer EV service capacity?"])
    brief.append(["Horizon", "2026–2030"])
    brief.append(["Required deliverables", "Observed-vs-projected separation; scenarios; capacity gap; option economics; staged commitment; decision triggers."])

    history = workbook.create_sheet("Demand_History")
    history.append(["Demand History"])
    history.append([])
    history.append(["Year", "Region", "Service_Demand_Units", "Basis", "Source", "Source_ID"])
    history.append([2025, "DACH", 100, "Observed", "Internal", "EV-2025-DA"])
    history.append([2025, "Nordics", 60, "Observed", "Internal", "EV-2025-NO"])
    history.append([2025, "Southern Europe", 40, "Observed", "Internal", "EV-2025-SE"])
    history.append([2026, "DACH", 120, "Observed", "Internal", "EV-2026-DA"])
    history.append([2026, "Nordics", 70, "Observed", "Internal", "EV-2026-NO"])
    history.append([2026, "Southern Europe", 50, "Observed", "Internal", "EV-2026-SE"])

    scenarios = workbook.create_sheet("Demand_Scenarios")
    scenarios.append(["Demand Scenarios"])
    scenarios.append([])
    scenarios.append(["Scenario", "Year", "Region", "Demand_Units", "Annual_Growth", "Basis", "Source_Note"])
    for scenario, values in {
        "Low": [120, 70, 50],
        "Base": [150, 90, 70],
        "High": [180, 110, 100],
    }.items():
        for region, demand in zip(["DACH", "Nordics", "Southern Europe"], values, strict=True):
            scenarios.append([scenario, 2030, region, demand, 0.10, "Projected", "Scenario assumption"])

    capacity = workbook.create_sheet("Capacity")
    capacity.append(["Capacity"])
    capacity.append([])
    capacity.append(["Site", "Region", "Year", "Nominal_Capacity", "Practical_Availability", "Status"])
    capacity.append(["Plant North", "DACH", 2030, 100, 0.90, "Existing/committed"])
    capacity.append(["Plant Nordic", "Nordics", 2030, 60, 0.90, "Existing/committed"])
    capacity.append(["Plant South", "Southern Europe", 2030, 40, 0.90, "Existing/committed"])

    infrastructure = workbook.create_sheet("Infrastructure")
    infrastructure.append(["Infrastructure"])
    infrastructure.append([])
    infrastructure.append(["Region", "Charging_Readiness_1_5", "Grid_Readiness_1_5", "Service_Network_Readiness_1_5", "Observation"])
    infrastructure.append(["DACH", 4.0, 3.8, 4.0, "Grid queues"])
    infrastructure.append(["Nordics", 4.5, 4.4, 4.5, "Strong coverage"])
    infrastructure.append(["Southern Europe", 2.9, 3.2, 2.7, "Readiness lags demand"])

    policy = workbook.create_sheet("Policy_Signals")
    policy.append(["Policy Signals"])
    policy.append([])
    policy.append(["Region", "Policy", "Current_End_Year", "Direction", "Demand_Relevance", "Implication"])
    policy.append(["DACH", "Purchase incentives", 2028, "Step-down", "Medium", "Base/high may overstate acceleration"])
    policy.append(["Southern Europe", "Charging subsidy", 2027, "Uncertain renewal", "High", "Demand is policy-sensitive"])

    options = workbook.create_sheet("Strategic_Options")
    options.append(["Strategic Options"])
    options.append([])
    options.append(["Option", "Investment_EUR", "Lead_Time_Months", "Added_Capacity", "Operating_Cost_Premium", "Reversibility", "Notes"])
    options.append(["Full expansion", 42000000, 24, 42, 0.25, "Low", "Irreversible build"])
    options.append(["Modular pilot", 9500000, 10, 12, 0.08, "High", "Expandable"])
    options.append(["Outsource overflow", 3500000, 4, 18, 0.18, "High", "Fast but margin dilution"])
    options.append(["Defer", 0, 0, 0, 0, "High", "No capacity increase"])
    options.append(["South regional expansion", None, 20, 18, 0.22, "Low", "Cost estimate not available"])

    kpis = workbook.create_sheet("Reported_KPIs")
    kpis.append(["Reported KPIs"])
    kpis.append([])
    kpis.append(["KPI", "Reported_Value", "Unit", "Scenario_Year", "Narrative", "Validation_Task"])
    kpis.append(["2030 Base Demand", 300, "units", "Base / 2030", "Expansion required", "Recalculate"])
    kpis.append(["2030 High Demand", 390, "units", "High / 2030", "Full expansion", "Recalculate"])
    kpis.append(["2030 Existing Practical Capacity", 180, "units", "2030", "Existing network", "Recalculate"])
    workbook.save(path)


def test_supplier_recipe_detects_anomalies_and_required_analysis(tmp_path: Path):
    source = tmp_path / "supplier.xlsx"
    _write_supplier_fixture(source)

    analysis = analyze_supplier_resilience(source)

    assert {item.kind for item in analysis.data_quality_issues} == {
        "DUPLICATE_LINE",
        "MISSING_ALTERNATIVES",
        "LOW_COST_HIGH_RISK",
    }
    assert analysis.concentration_metrics
    assert {item.scenario for item in analysis.scenarios} == {"Short outage", "Extended outage", "Correlated region shock"}
    assert len(analysis.options) == 4
    assert any(metric.criticality_weighted >= 4 for metric in analysis.concentration_metrics)


def test_supplier_low_cost_screen_normalizes_mixed_currency(tmp_path: Path):
    source = tmp_path / "supplier.xlsx"
    _write_supplier_fixture(source)
    analysis = analyze_supplier_resilience(source)

    assert any("Low Cost Risk" in item.message for item in analysis.data_quality_issues if item.kind == "LOW_COST_HIGH_RISK")
    assert any("USD=0.92" in item for item in analysis.assumptions)


def test_supplier_workflow_persists_reviewed_finding_and_report(tmp_path: Path):
    source = tmp_path / "supplier.xlsx"
    _write_supplier_fixture(source)
    store = AnalysisStore(tmp_path / "case")

    result = run_supplier_resilience_workflow(store, source, human_decision="ACCEPT")

    assert result.status == "accepted"
    assert result.report_path is not None and result.report_path.exists()
    assert all(item.accepted for item in derive_finding_views(store.load_analysis()))
    report = result.report_path.read_text(encoding="utf-8")
    for section in ["Concentration metrics", "Outage scenarios", "Sourcing alternatives", "Counterarguments", "Early-warning indicators"]:
        assert section in report
    assert "Supplier_Master" in report


def test_supplier_report_uses_requested_reporting_currency(tmp_path: Path):
    source = tmp_path / "supplier.xlsx"
    _write_supplier_fixture(source)
    result = run_supplier_resilience_workflow(
        AnalysisStore(tmp_path / "case"),
        source,
        human_decision="ACCEPT",
        input_contract={
            "domain": "supplier_resilience",
            "reporting_currency": "USD",
            "conversion_rates": {"EUR": 2.0, "USD": 1.0, "CNY": 0.2},
        },
    )

    assert result.status == "accepted"
    assert result.supplier is not None and result.supplier.reporting_currency == "USD"
    assert result.analysis is not None and any(item.units == "USD" for item in result.analysis.calculations)
    report = result.report_path.read_text(encoding="utf-8") if result.report_path else ""
    assert "Reporting currency: USD" in report
    assert "USD " in report
    assert "€" not in report


def test_ev_recipe_separates_observed_forecast_and_flags_conflicts(tmp_path: Path):
    source = tmp_path / "ev.xlsx"
    _write_ev_fixture(source)

    analysis = analyze_ev_capacity_investment(source)

    assert {item.kind for item in analysis.data_quality_issues} == {
        "KPI_MISMATCH",
        "MISSING_EXPANSION_COST",
        "POLICY_CONFLICT",
    }
    assert {item.scenario for item in analysis.scenarios} == {"Low", "Base", "High"}
    assert any(item.capacity_gap > 0 for item in analysis.scenarios)
    assert len(analysis.options) == 5
    assert any(item.basis == "Observed" for item in analysis.demand_history)
    assert any(item.basis == "Projected" for item in analysis.scenarios)


def test_ev_workflow_persists_reviewed_finding_and_report(tmp_path: Path):
    source = tmp_path / "ev.xlsx"
    _write_ev_fixture(source)
    store = AnalysisStore(tmp_path / "case")

    result = run_ev_capacity_investment_workflow(store, source, human_decision="ACCEPT")

    assert result.status == "accepted"
    assert result.report_path is not None and result.report_path.exists()
    assert all(item.accepted for item in derive_finding_views(store.load_analysis()))
    report = result.report_path.read_text(encoding="utf-8")
    for section in ["Observed demand", "Alternative futures", "Capacity gap", "Option economics", "Decision triggers"]:
        assert section in report
    assert "Demand_Scenarios" in report
