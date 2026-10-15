import hashlib
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import get_type_hints

import openpyxl
import pytest

from foresight_os.runtime.local_toolbelt import (
    AnalysisOperation,
    FormulaStatus,
    SheetSelection,
    ToolbeltLimits,
    ToolbeltError,
    TypedMeasurement,
    _implicit_xlsx_bounds,
    calculation,
    data_analysis,
    spreadsheet_read,
    WorkbookReadCache,
    bind_local_toolbelt,
)
from foresight_os.domain.models import Evidence, SourceLocator


def test_provider_tool_paths_are_json_string_boundaries():
    bindings = bind_local_toolbelt()

    assert get_type_hints(bindings["spreadsheet_read"].value)["source"] is str
    assert get_type_hints(bindings["powerbi_project_read"].value)["source"] is str


def test_provider_spreadsheet_arguments_are_json_coerced(workbook_path: Path):
    result = spreadsheet_read(
        str(workbook_path),
        "Demand",
        {"max_rows": 10},
    )

    assert result.sheets[0].name == "Demand"
    assert result.measurements


def test_workbook_read_compacts_implicit_padding_and_provider_serialization(tmp_path: Path):
    source = tmp_path / "padded.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Case"
    sheet.append(["Label", "Value"])
    sheet.append(["North", 10])
    sheet["Z1000"].number_format = "0.00"
    workbook.save(source)

    result = spreadsheet_read(str(source), "Case")
    payload = result.to_dict()
    encoded = json.dumps(payload, sort_keys=True)

    assert result.tables[0].source_range == "A1:B2"
    assert len(result.measurements) == 4
    assert payload["measurement_count"] == 4
    assert payload["omitted_blank_measurement_count"] == 0
    assert result.workbook_sha256 in encoded
    assert "B2" in encoded
    assert len(encoded) < 10_000


def test_implicit_xlsx_bounds_uses_sparse_cells_without_iterating():
    class Cell:
        def __init__(self, row, column, value):
            self.row = row
            self.column = column
            self.value = value

    class Worksheet:
        _cells = {(1, 1): Cell(1, 1, "header"), (4, 3): Cell(4, 3, 10)}

        def iter_rows(self):
            raise AssertionError("implicit bounds must not scan the worksheet grid")

    assert _implicit_xlsx_bounds(Worksheet()) == (1, 1, 3, 4)


def test_selected_sheet_does_not_inherit_unselected_sheet_row_limit(tmp_path: Path):
    workbook = openpyxl.Workbook()
    large = workbook.active
    large.title = "Large"
    for _ in range(20):
        large.append(["value"])
    small = workbook.create_sheet("Small")
    small.append(["value"])
    source = tmp_path / "bounded.xlsx"
    workbook.save(source)

    result = spreadsheet_read(str(source), {"sheet": "Small"}, {"max_rows": 2})

    assert result.tables[0].sheet == "Small"


def test_unselected_sheet_merge_allocation_is_preflighted(tmp_path: Path):
    workbook = openpyxl.Workbook()
    selected = workbook.active
    selected.title = "Selected"
    selected.append(["value"])
    selected.append([1])
    unselected = workbook.create_sheet("Unselected")
    unselected.merge_cells("A1:J1000")
    source = tmp_path / "merged-unselected.xlsx"
    workbook.save(source)

    with pytest.raises(ToolbeltError, match="allocation exceeds cell limit"):
        spreadsheet_read(source, {"sheet": "Selected"}, {"max_cells": 5})


def test_workbook_read_cache_evicts_old_entries_without_unpacking_error(tmp_path: Path):
    cache = WorkbookReadCache(max_entries=1)
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    for source, value in ((first, "first"), (second, "second")):
        workbook = openpyxl.Workbook()
        workbook.active.append([value])
        workbook.save(source)

    spreadsheet_read(str(first), _cache=cache)
    spreadsheet_read(str(second), _cache=cache)

    assert cache.size > 0


def _make_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    demand = workbook.active
    demand.title = "Demand"
    demand.append(["Demand inputs", "", "", "", "", ""])
    demand.append(["Region", "Metric", "Value", "Date", "Formula", "Error"])
    demand.append(["North", "demand", 10, datetime(2026, 1, 1), "=Assumptions!B2*C3", "#N/A"])
    demand.append(["North", "demand", 10, datetime(2026, 1, 1), "=C4*2", None])
    demand.row_dimensions[4].hidden = True
    demand.merge_cells("A1:B1")
    demand.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A2:F4"))
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Secret", 99])
    assumptions = workbook.create_sheet("Assumptions")
    assumptions.append(["Scenario", "Multiplier"])
    assumptions.append(["base", 1.0])
    assumptions.append(["high", 1.2])
    workbook.save(path)


@pytest.fixture
def workbook_path(tmp_path: Path) -> Path:
    path = tmp_path / "analyst_multisheet.xlsx"
    _make_workbook(path)
    return path


def test_xlsx_read_inspects_sheets_hidden_rows_headers_types_formulas_and_provenance(workbook_path):
    before = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    result = spreadsheet_read(workbook_path, SheetSelection(tables=(("Demand", "DemandTable"),)))
    after = hashlib.sha256(workbook_path.read_bytes()).hexdigest()

    assert before == after
    assert {sheet.name for sheet in result.sheets} == {"Demand", "Assumptions"}
    assert result.sheets[0].hidden_rows == (4,)
    table = result.tables[0]
    assert table.name == "DemandTable"
    assert table.headers[0][:2] == ("Region", "Metric")
    assert len(table.rows) == 1
    formula_cells = [item for item in result.measurements if item.formula]
    assert {item.formula_status for item in formula_cells} == {FormulaStatus.FORMULA_WITHOUT_CACHE}
    assert any(item.cross_sheet_references == ("Assumptions",) for item in formula_cells)
    assert any(item.data_type == "date" for item in result.measurements)
    assert any(item.data_type == "error" for item in result.measurements)
    assert all(item.provenance.workbook_sha256 == result.workbook_sha256 for item in result.measurements)


def test_xlsx_selection_can_include_hidden_sheet_and_rows_and_analysis_is_deterministic(workbook_path):
    result = spreadsheet_read(
        workbook_path,
        SheetSelection(sheet_names=("Demand", "Hidden"), include_hidden=True, include_hidden_rows=True),
    )
    demand = next(table for table in result.tables if table.sheet == "Demand")
    assert len(demand.rows) == 2
    assert any(table.sheet == "Hidden" for table in result.tables)

    first = data_analysis(demand, AnalysisOperation("duplicates", key_columns=("Region", "Metric", "Value")))
    second = data_analysis(demand, AnalysisOperation("duplicates", key_columns=("Region", "Metric", "Value")))
    assert first.deterministic_key == second.deterministic_key
    assert first.rows[0]["count"] == 2

    scenarios = data_analysis(demand, AnalysisOperation("scenario", value_column="Value", scenarios=(("base", 1), ("high", 1.2))))
    assert [row["scenario"] for row in scenarios.rows] == ["base", "high"]


def test_formula_without_cache_blocks_decision_critical_calculation(workbook_path):
    result = spreadsheet_read(workbook_path, SheetSelection(tables=(("Demand", "DemandTable"),)))
    formula = next(item for item in result.measurements if item.formula)
    with pytest.raises(ToolbeltError, match="BLOCKED_RECALCULATION"):
        calculation([formula], formula=__import__("foresight_os.runtime.local_toolbelt", fromlist=["FormulaSpec"]).FormulaSpec("formula", "sum"), units=__import__("foresight_os.runtime.local_toolbelt", fromlist=["UnitSpec"]).UnitSpec("units"))


def test_false_evidence_excerpt_cannot_be_recalculated_as_workbook_truth():
    evidence = Evidence("evidence-1", SourceLocator("hash", "Demand", "C3", "units", "none", "source"), "v1", "2026", "=999")

    with pytest.raises(ToolbeltError, match="authoritative workbook measurement required"):
        calculation(
            [evidence],
            formula={"expression": "Demand!C3", "operation": "sum"},
            units={"units": "units"},
        )


def test_calculation_rejects_untrusted_json_shaped_measurements():
    measurement = {
        "id": "measurement-1",
        "value": 12,
        "data_type": "number",
        "formula": None,
        "cached_value": 12,
        "formula_status": "VALUE",
        "provenance": {
            "workbook_sha256": "workbook-hash",
            "sheet": "Demand",
            "coordinate": "C3",
            "range": "C3",
        },
    }

    with pytest.raises(ToolbeltError, match="provenance is untrusted"):
        calculation(
            [measurement],
            formula={"expression": "Demand!C3", "operation": "sum"},
            units={"units": "units"},
        )


def test_spreadsheet_measurement_survives_json_provider_boundary(workbook_path):
    result = spreadsheet_read(workbook_path, SheetSelection(tables=(("Demand", "DemandTable"),)))
    measurement = next(item for item in result.measurements if item.provenance.coordinate == "C3")

    record = calculation(
        [json.loads(json.dumps(measurement.to_dict(), default=str))],
        formula={"expression": "Demand!C3", "operation": "sum"},
        units={"units": "units"},
        approved_workbook_sha256=result.workbook_sha256,
    )

    assert record.verification_status.value == "VERIFIED"


def test_fabricated_calculation_mapping_cannot_be_verified():
    with pytest.raises(ToolbeltError, match="authoritative calculation"):
        calculation(
            [{
                "id": "calculation-fake",
                "key": {"region": "toolbelt", "measure": "fake", "period": "current", "units": "units", "scenario": "default", "version": "v1"},
                "formula": "fake",
                "input_ids": ["measurement-fake"],
                "output": 12,
                "units": "units",
                "source_locators": ["workbook-hash:Demand!C3"],
            }],
            formula={"expression": "fake", "operation": "sum"},
            units={"units": "units"},
            approved_workbook_sha256="workbook-hash",
        )


def test_serialized_formula_measurements_keep_recalculation_guard(tmp_path):
    source = tmp_path / "formula.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Value"])
    workbook.active.append(["=1+1"])
    workbook.save(source)

    result = spreadsheet_read(source)
    payload = result.tables[0].to_dict()

    with pytest.raises(ToolbeltError, match="BLOCKED_RECALCULATION"):
        data_analysis(payload, AnalysisOperation("aggregate", value_column="Value"))


def test_formula_cache_and_manual_recalculation_state_are_truthful(workbook_path, tmp_path):
    cached_path = tmp_path / "cached.xlsx"
    cached_path.write_bytes(workbook_path.read_bytes())
    with zipfile.ZipFile(cached_path, "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    sheet = entries["xl/worksheets/sheet1.xml"].decode("utf-8")
    sheet, replacements = re.subn(
        r'(<c[^>]*\br="E3"[^>]*>.*?<f>Assumptions!B2\*C3</f>.*?<v)(?:\s*/>|>.*?</v>)',
        r'\g<1>>10</v>',
        sheet,
        count=1,
        flags=re.DOTALL,
    )
    assert replacements == 1
    entries["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")
    workbook_xml, replacements = re.subn(r'<calcPr([^>]*)/>', r'<calcPr\1 calcMode="manual"/>', entries["xl/workbook.xml"].decode("utf-8"), count=1)
    assert replacements == 1
    entries["xl/workbook.xml"] = workbook_xml.encode("utf-8")
    with zipfile.ZipFile(cached_path, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    result = spreadsheet_read(cached_path, SheetSelection(tables=(("Demand", "DemandTable"),)))
    formula = next(item for item in result.measurements if item.coordinate == "E3") if False else next(item for item in result.measurements if item.formula and item.provenance.coordinate == "E3")
    assert formula.cached_value == 10
    assert formula.formula_status == FormulaStatus.FORMULA_STALE_OR_UNKNOWN


def test_analysis_supports_filters_and_duplicate_join_key_guard(workbook_path):
    result = spreadsheet_read(workbook_path, SheetSelection(tables=(("Demand", "DemandTable"),), include_hidden_rows=True))
    demand = next(table for table in result.tables if table.sheet == "Demand")
    filtered = data_analysis(demand, AnalysisOperation("missing", filters=(("Region", "eq", "North"),)))
    assert filtered.rows[0]["column"] == "Region"
    with pytest.raises(ToolbeltError, match="duplicate join keys"):
        data_analysis(demand, AnalysisOperation("join", join_table=demand, join_keys=("Region", "Region")))


@pytest.mark.parametrize("kind", ["sum", "scenario", "sensitivity"])
def test_analysis_blocks_numeric_analysis_when_rows_mix_units(tmp_path, kind):
    source = tmp_path / "mixed_units.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Metric", "Unit", "Value"])
    sheet.append(["share", "percent", 25])
    sheet.append(["sales", "Vehicles", 100])
    workbook.save(source)

    table = spreadsheet_read(source).tables[0]
    operation = AnalysisOperation(
        kind,
        value_column="Value",
        scenarios=(("base", 1.0),),
        sensitivity=(("scale", (1.0,)),),
    )
    with pytest.raises(ToolbeltError, match="unit mismatch"):
        data_analysis(table, operation)


def test_analysis_ignores_units_on_rows_without_numeric_values(tmp_path):
    source = tmp_path / "nullable_units.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Metric", "Unit", "Value"])
    sheet.append(["sales", "Vehicles", 100])
    sheet.append(["share", "percent", None])
    workbook.save(source)

    table = spreadsheet_read(source).tables[0]
    result = data_analysis(table, AnalysisOperation("sum", value_column="Value"))
    assert result.rows[0]["sum"] == 100.0


def test_spreadsheet_and_analysis_honor_cancellation_checkpoints(workbook_path):
    cancelled = lambda: True
    with pytest.raises(ToolbeltError, match="cancelled"):
        spreadsheet_read(workbook_path, cancel_check=cancelled)
    result = spreadsheet_read(workbook_path, SheetSelection(tables=(("Demand", "DemandTable"),)))
    with pytest.raises(ToolbeltError, match="cancelled"):
        data_analysis(result.tables[0], AnalysisOperation("duplicates"), cancel_check=cancelled)


def test_csv_read_is_supported_with_typed_values_and_provenance(tmp_path):
    source = tmp_path / "signals.csv"
    source.write_text("Region,Value\nNorth,10\nSouth,20\n", encoding="utf-8")
    result = spreadsheet_read(source)
    table = result.tables[0]
    assert table.columns == ("Region", "Value")
    assert table.rows == (("North", 10), ("South", 20))
    assert result.measurements[1].provenance.sheet == "signals"
    assert data_analysis(table, AnalysisOperation("sum", value_column="Value")).rows[0]["sum"] == 30.0


def test_csv_range_selection_preserves_original_coordinates(tmp_path):
    source = tmp_path / "signals.csv"
    source.write_text("Region,Value,Flag\nNorth,10,Y\nSouth,20,N\n", encoding="utf-8")
    result = spreadsheet_read(source, SheetSelection(ranges=(("signals", "B2:C3"),)))
    coordinates = [item.provenance.coordinate for item in result.measurements]
    assert coordinates == ["B2", "C2", "B3", "C3"]


def test_csv_range_selection_honors_row_limit(tmp_path):
    source = tmp_path / "signals.csv"
    source.write_text("Region,Value\n" + "\n".join(f"R{index},{index}" for index in range(1, 101)) + "\n", encoding="utf-8")
    with pytest.raises(ToolbeltError, match="row limit"):
        spreadsheet_read(source, SheetSelection(ranges=(("signals", "A1:B101"),)), ToolbeltLimits(max_rows=2))


def test_xlsx_actual_cells_are_counted_before_eager_load(workbook_path):
    with pytest.raises(ToolbeltError, match="cell limit"):
        spreadsheet_read(workbook_path, SheetSelection(sheet_names=("Demand",)), ToolbeltLimits(max_cells=1))


def test_xlsx_expanded_archive_limit_is_checked_before_loading(workbook_path):
    with pytest.raises(ToolbeltError, match="expanded content"):
        spreadsheet_read(workbook_path, limits=ToolbeltLimits(max_expanded_bytes=1))


def test_xlsx_allocation_limits_are_cumulative_across_sheets(tmp_path):
    source = tmp_path / "multi-sheet-allocation.xlsx"
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["A", "B"])
    second = workbook.create_sheet("Second")
    second.append(["C", "D"])
    workbook.save(source)

    with pytest.raises(ToolbeltError, match="allocation exceeds cell limit"):
        spreadsheet_read(source, limits=ToolbeltLimits(max_cells=3))
    with pytest.raises(ToolbeltError, match="allocation exceeds row limit"):
        spreadsheet_read(source, limits=ToolbeltLimits(max_rows=1))


def test_xlsx_merged_range_count_is_limited_before_loading(tmp_path):
    source = tmp_path / "merged-range-limit.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.merge_cells("A1:B1")
    sheet.merge_cells("A2:B2")
    workbook.save(source)

    with pytest.raises(ToolbeltError, match="merged range limit"):
        spreadsheet_read(source, limits=ToolbeltLimits(max_merged_ranges=1))


def test_xlsx_merged_area_is_limited_before_loading(tmp_path):
    source = tmp_path / "merged-area-limit.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.merge_cells("A1:C4")
    workbook.save(source)

    with pytest.raises(ToolbeltError, match="merged area exceeds per-range"):
        spreadsheet_read(source, limits=ToolbeltLimits(max_merged_cells_per_range=3))


def test_compare_preserves_missing_values_and_reports_incomplete_status(tmp_path):
    source = tmp_path / "compare.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Left", "Right"])
    sheet.append([10, None])
    workbook.save(source)

    result = data_analysis(spreadsheet_read(source).tables[0], AnalysisOperation("compare", columns=("Left", "Right")))
    assert result.rows[0]["delta"] is None
    assert result.rows[0]["status"] == "INCOMPLETE_INPUT"


@pytest.mark.parametrize("kind", ["aggregate", "sum", "mean"])
def test_aggregate_aliases_return_deterministic_numeric_results(tmp_path, kind):
    source = tmp_path / f"{kind}.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Value"])
    sheet.append([2])
    sheet.append([4])
    workbook.save(source)

    result = data_analysis(spreadsheet_read(source).tables[0], AnalysisOperation(kind, value_column="Value"))
    assert result.rows[0]["sum"] == 6.0
    assert result.rows[0]["mean"] == 3.0


def test_run_local_workbook_cache_keeps_path_provenance(tmp_path, workbook_path):
    second = tmp_path / "same-content.xlsx"
    second.write_bytes(workbook_path.read_bytes())
    cache = WorkbookReadCache()
    first_result = spreadsheet_read(workbook_path, _cache=cache)
    second_result = spreadsheet_read(second, _cache=cache)
    assert cache.size == 2
    assert first_result.source == str(workbook_path)
    assert second_result.source == str(second)
    assert first_result.workbook_sha256 == second_result.workbook_sha256


def test_overlapping_workbook_reads_preserve_measurement_authority(workbook_path):
    narrow = spreadsheet_read(workbook_path, SheetSelection(ranges=(("Demand", "A2:C3"),)))
    measurement = next(item for item in narrow.measurements if item.provenance.coordinate == "C3")
    spreadsheet_read(workbook_path, SheetSelection(ranges=(("Demand", "A2:F4"),)))

    record = calculation(
        [json.loads(json.dumps(measurement.to_dict(), default=str))],
        formula={"expression": "Demand!C3", "operation": "sum"},
        units={"units": "units"},
        approved_workbook_sha256=narrow.workbook_sha256,
    )

    assert record.verification_status.value == "VERIFIED"


def test_run_local_workbook_cache_reuses_deterministic_analysis(workbook_path):
    cache = WorkbookReadCache()
    table = spreadsheet_read(workbook_path, _cache=cache).tables[0]
    operation = AnalysisOperation("duplicates", key_columns=("Region", "Metric", "Value"))

    first = data_analysis(table, operation, _cache=cache)
    second = data_analysis(table, operation, _cache=cache)

    assert first is second
    assert cache.size == 2


def test_calculation_provenance_qualifies_workbook_sheet_and_coordinate(workbook_path):
    result = spreadsheet_read(workbook_path, SheetSelection(sheet_names=("Demand", "Assumptions"), include_hidden=True, include_hidden_rows=True))
    demand_value = next(item for item in result.measurements if item.provenance.sheet == "Demand" and item.provenance.coordinate == "C3")
    assumption_value = next(item for item in result.measurements if item.provenance.sheet == "Assumptions" and item.provenance.coordinate == "B2")
    record = calculation(
        [demand_value, assumption_value],
        __import__("foresight_os.runtime.local_toolbelt", fromlist=["FormulaSpec"]).FormulaSpec("Demand!C3 + Assumptions!B2"),
        __import__("foresight_os.runtime.local_toolbelt", fromlist=["UnitSpec"]).UnitSpec("units"),
        approved_workbook_sha256=result.workbook_sha256,
    )
    assert any(":Demand!C3" in locator for locator in record.source_locators)
    assert any(":Assumptions!B2" in locator for locator in record.source_locators)
    assert record.verification_status.value == "VERIFIED"

def test_xlsx_reader_skips_blank_and_title_rows_before_data_header(tmp_path):
    source = tmp_path / "titled-range.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append([None, None, None])
    sheet.append(["Portfolio Data", None, None])
    sheet.append(["Region", "Value", "Unit"])
    sheet.append(["North", 10, "EUR"])
    sheet.append(["South", 20, "EUR"])
    workbook.save(source)

    table = spreadsheet_read(source).tables[0]

    assert table.columns == ("Region", "Value", "Unit")
    assert table.rows == (("North", 10, "EUR"), ("South", 20, "EUR"))
