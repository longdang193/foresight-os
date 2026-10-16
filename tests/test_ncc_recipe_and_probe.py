import hashlib
import json
from dataclasses import replace
from pathlib import Path
import shutil
import subprocess
import sys

import openpyxl
import pytest

from foresight_os.domain.models import (
    AnalysisState,
    AnalyticalRecord,
    Assessment,
    Basis,
    CalculationRecord,
    Evidence,
    Invalidation,
    MeasurementKey,
    Revision,
    SourceLocator,
)
from foresight_os.persistence.analysis_store import AnalysisStore


ROOT = Path(__file__).parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "ncc-independent-validation.xlsx"
SPEC = ROOT / "tests" / "fixtures" / "ncc-independent-validation-review-spec.json"
ACTUAL_SPEC = ROOT / "evaluation" / "cases" / "ncc-2022-cosluxe-vietnam" / "review-spec.json"
PROBE = ROOT / "evaluation" / "probes" / "run_ncc_independent_validation.py"


def make_analysis(path: Path, *, growth: float = 0.1, threshold: float = 116.64, category: str = "Core") -> Path:
    workbook_sha = hashlib.sha256(FIXTURE.read_bytes()).hexdigest()
    evidence = Evidence("ncc-source", SourceLocator(workbook_sha, "Offline", "A1:C3", "value", "raw", "source"), workbook_sha, "2024", "NCC fixture")
    claim = AnalyticalRecord("ncc-claim", "Core category growth is observed in the fixture.", Basis.OBSERVED, Assessment.SUPPORTED, evidence_ids=[evidence.id])
    calculations = [
        CalculationRecord("ncc-offline-growth-core", MeasurementKey(category, "growth_rate", "2024", "rate"), "independent fixture calculation", [evidence.id], growth, "rate", ["Offline!C2", "Offline!C3"]),
        CalculationRecord("ncc-offline-planning-threshold", MeasurementKey(category, "planning_threshold", "2024", "value"), "independent fixture planning threshold", [evidence.id], threshold, "value", ["Offline!C2"]),
    ]
    state = AnalysisState(
        "ncc-fixture",
        1,
        evidence=[evidence],
        claims=[claim],
        calculations=calculations,
        revisions=[Revision("revision-1", 1, "fixture analysis", [evidence.id, claim.id, *[item.id for item in calculations]])],
        invalidation=Invalidation(),
    )
    store = AnalysisStore(path)
    store.save_analysis(state)
    return path


def run_probe(tmp_path: Path, *, analysis_dir: Path | None = None, workbook: Path = FIXTURE):
    analysis_dir = analysis_dir or make_analysis(tmp_path / "analysis")
    output = tmp_path / "output"
    completed = subprocess.run(
        [sys.executable, str(PROBE), "--workbook", str(workbook), "--analysis-dir", str(analysis_dir), "--review-spec", str(SPEC), "--output-root", str(output)],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )
    assert completed.returncode == 0, completed.stderr
    return json.loads((output / "receipt.json").read_text(encoding="utf-8"))


def test_actual_review_spec_matches_external_ncc_workbook_scope():
    workbook_path = ROOT / "data" / "exercise1-nielsenIQ" / "NCC 2022 - Semi Final Case - Data set.xlsx"
    if not workbook_path.exists():
        pytest.skip("NCC workbook is an approved external fixture and is not tracked in the repository")
    spec = json.loads(ACTUAL_SPEC.read_text(encoding="utf-8"))
    workbook = openpyxl.load_workbook(workbook_path, read_only=True)

    assert set(spec["source"]["sheets"]) == {"Macro & FMCG", "Personal Care", "Modern Trade", "Ecommerce"}
    assert {tuple(item) for item in spec["source"]["ranges"]} == {
        ("Macro & FMCG", "A1:G24"),
        ("Personal Care", "A1:U18"),
        ("Modern Trade", "A1:O58"),
        ("Ecommerce", "A1:Q65"),
    }
    assert set(spec["source"]["sheets"]).issubset(set(workbook.sheetnames))
    assert all(item["operation"] in {"growth_rate", "ratio", "share"} for item in spec["calculations"])


def test_probe_loads_persisted_analysis_and_keeps_decision_challenge_separate(tmp_path: Path):
    receipt = run_probe(tmp_path)

    assert receipt["status"] == "NOT_VERIFIED"
    assert receipt["verification"] == "NOT_VERIFIED"
    assert receipt["automated_review_outcome"] == "PASS"
    assert receipt["human_acceptance"] == "NOT_RUN"
    assert receipt["calculation_ids"] == ["ncc-offline-growth-core", "ncc-offline-planning-threshold"]
    assert receipt["checks"][-1]["check_id"] == "decision_challenge"
    assert receipt["checks"][-1]["status"] == "NOT_RUN"


def test_probe_detects_wrong_denominator_and_value_volume_mix(tmp_path: Path):
    analysis_dir = make_analysis(tmp_path / "analysis", growth=0.21, threshold=121.0)
    receipt = run_probe(tmp_path, analysis_dir=analysis_dir)

    assert receipt["status"] == "FAILED"
    calculation_checks = [item for item in receipt["checks"] if item["check_id"] == "calculation_reconciliation"]
    assert all(item["status"] == "FAIL" for item in calculation_checks)
    assert all(item["expected"] != item["actual"] for item in calculation_checks)


def test_probe_detects_invalid_category_mapping(tmp_path: Path):
    receipt = run_probe(tmp_path, analysis_dir=make_analysis(tmp_path / "analysis", category="Invalid"))

    assert receipt["status"] == "FAILED"
    assert any(item["status"] == "FAIL" and item["check_id"] == "calculation_reconciliation" for item in receipt["checks"])


def test_probe_detects_changed_workbook_hash_without_copying_source_into_receipt(tmp_path: Path):
    changed = tmp_path / "changed.xlsx"
    shutil.copy2(FIXTURE, changed)
    workbook = openpyxl.load_workbook(changed)
    workbook["Offline"]["C3"] = 120
    workbook.save(changed)

    receipt = run_probe(tmp_path, workbook=changed)

    assert receipt["status"] == "FAILED"
    source_check = next(item for item in receipt["checks"] if item["check_id"] == "source_fidelity")
    assert source_check["status"] == "FAIL"
    assert "bytes" not in json.dumps(receipt).lower()


def test_probe_missing_persisted_analysis_is_blocked_not_passed(tmp_path: Path):
    receipt = run_probe(tmp_path, analysis_dir=tmp_path / "missing")

    assert receipt["status"] == "BLOCKED_INPUT"
    assert receipt["verification"] == "NOT_VERIFIED"


def test_raw_cell_check_supports_explicit_channel_layout_operations():
    from types import SimpleNamespace

    from evaluation.probes.run_ncc_independent_validation import _raw_cell_check

    result = SimpleNamespace(
        measurements=tuple(
            SimpleNamespace(provenance=SimpleNamespace(sheet="Channel Dynamic", coordinate=coordinate), value=value)
            for coordinate, value in {
                "B4": 1_523_412,
                "D4": 1_489_956,
                "B5": 301,
                "D5": 334,
                "B6": 583,
                "D6": 744,
                "B7": 1_513,
                "D7": 2_512,
            }.items()
        )
    )

    values = {(item.provenance.sheet, item.provenance.coordinate): item.value for item in result.measurements}
    percent_change = _raw_cell_check(
        result,
        SimpleNamespace(id="growth", output=-2.2),
        {
            "operation": "relative_change",
            "scale": 100,
            "source_cells": ["Channel Dynamic!B4", "Channel Dynamic!D4"],
            "tolerance": 0.05,
        },
        values,
    )
    store_difference = _raw_cell_check(
        result,
        SimpleNamespace(id="difference", output=161),
        {
            "operation": "difference",
            "source_cells": ["Channel Dynamic!B6", "Channel Dynamic!D6"],
            "tolerance": 0.1,
        },
        values,
    )
    combined_change = _raw_cell_check(
        result,
        SimpleNamespace(id="combined", output=49.8),
        {
            "operation": "sum_relative_change",
            "scale": 100,
            "split_at": 3,
            "source_cells": [
                "Channel Dynamic!B5",
                "Channel Dynamic!B6",
                "Channel Dynamic!B7",
                "Channel Dynamic!D5",
                "Channel Dynamic!D6",
                "Channel Dynamic!D7",
            ],
            "tolerance": 0.05,
        },
        values,
    )

    assert percent_change.status == "PASS"
    assert store_difference.status == "PASS"
    assert combined_change.status == "PASS"
