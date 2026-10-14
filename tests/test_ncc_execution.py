from pathlib import Path

from openpyxl import load_workbook
import pytest

from foresight_os.domain.models import AnalysisState, Artifact, Lifecycle, Revision, RevisionStatus
from foresight_os.persistence.analysis_store import AnalysisStore
from foresight_os.wayfinding import create_work_brief


WORKBOOK = Path(__file__).resolve().parents[1] / "data" / "exercise1-nielsenIQ" / "NCC 2022 - Semi Final Case - Data set.xlsx"


@pytest.mark.skipif(
    not WORKBOOK.exists(),
    reason="NCC workbook is an approved external fixture and is not tracked in the repository",
)
def test_ncc_workbook_boundary_is_available_without_claiming_analytical_completion():
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)

    assert {"Case", "Glossary", "Personal Care", "Consumer", "Channel Dynamic", "Trad.Trade", "Modern Trade", "Ecommerce", "Others"}.issubset(workbook.sheetnames)
    assert workbook["Others"].sheet_state == "hidden"


def test_ncc_deliverables_require_durable_artifacts_and_review(tmp_path: Path):
    store = AnalysisStore(tmp_path / "ncc")
    brief = create_work_brief(
        "Which category and channels should Cosluxe Nature prioritize?",
        investigation_id="ncc",
        desired_deliverables=[
            "category screen",
            "channel recommendation",
            "conditional growth bridge",
        ],
    )
    store.save_brief(brief)
    state = AnalysisState(
        "ncc",
        1,
        artifacts=[Artifact("artifact-category", "category screen", "artifacts/category-screen.md", Lifecycle.CURRENT)],
        revisions=[Revision("revision-1", 1, "controlled plumbing fixture", ["artifact-category"])],
        revision_status=RevisionStatus.PENDING_REVIEW,
    )
    store.save_pending(state)

    reloaded = store.load_analysis()
    statuses = {
        requested: (
            "PERSISTED"
            if any(
                requested.lower() in f"{artifact.kind} {artifact.path}".lower()
                for artifact in reloaded.artifacts
            )
            else "NOT_STARTED"
        )
        for requested in brief.desired_deliverables
    }
    runtime_summary = {"outputs": 6, "settled": True}

    assert reloaded.revision_status == RevisionStatus.PENDING_REVIEW
    assert statuses == {
        "category screen": "PERSISTED",
        "channel recommendation": "NOT_STARTED",
        "conditional growth bridge": "NOT_STARTED",
    }
    assert "status" not in runtime_summary
