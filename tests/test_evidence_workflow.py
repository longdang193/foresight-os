import asyncio
import hashlib
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from threading import Thread

import pytest

from foresight_os.analysis.evidence_workflow import accept_pending_evidence, run_evidence_workflow
from foresight_os.domain.models import AnalysisState, Revision
from foresight_os.persistence.analysis_store import AnalysisStore
from foresight_os.runtime.page_retrieval import PageRetrievalCapability
from foresight_os.runtime.web_search import WebSearchHit
from foresight_os.runtime.local_toolbelt import AnalysisOperation, SheetSelection, SourceSnapshot, ToolbeltLimits, WorkbookContext
import openpyxl


class _Provider:
    name = "fixture-search"
    fingerprint = "fixture-search-v1"

    def __init__(self, url):
        self.url = url
        self.calls = 0

    async def search(self, query, *, max_results=5, freshness=None):
        self.calls += 1
        return [WebSearchHit("Capacity forecast", self.url, "A source excerpt", source=self.name)]


class _Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        body = b"<html><body>The source reports a forecast fact.</body></html>"
        self.send_response(200)
        self.send_header("Content-Type", "text/html")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *_):
        return


def test_evidence_workflow_requires_explicit_acceptance_and_supports_restart(tmp_path: Path):
    server = ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        url = f"http://127.0.0.1:{server.server_port}/source"
        provider = _Provider(url)
        store = AnalysisStore(tmp_path / "case")
        rejected = asyncio.run(run_evidence_workflow(store, "capacity expansion cost", provider, retrieval=PageRetrievalCapability(fixture_hosts={"127.0.0.1"})))
        assert rejected.status == "rejected"
        assert provider.calls == 1
        assert rejected.analysis.revision_status.value == "REJECTED"

        provider = _Provider(url)
        store = AnalysisStore(tmp_path / "restart")
        accepted = asyncio.run(run_evidence_workflow(store, "capacity expansion cost", provider, retrieval=PageRetrievalCapability(fixture_hosts={"127.0.0.1"}), human_decision="ACCEPT"))
        assert accepted.status == "accepted"
        assert accepted.analysis.claims[0].evidence_ids == [accepted.evidence.id]
        assert provider.calls == 1

        pending_store = AnalysisStore(tmp_path / "pending")
        provider = _Provider(url)
        first = asyncio.run(run_evidence_workflow(pending_store, "capacity expansion cost", provider, retrieval=PageRetrievalCapability(fixture_hosts={"127.0.0.1"}), human_decision="PENDING"))
        assert first.status == "pending"
        assert first.analysis.revision_status.value == "PENDING_REVIEW"
        resumed = accept_pending_evidence(pending_store, human_decision="ACCEPT")
        assert resumed.status == "accepted"
        assert provider.calls == 1

        preserved_store = AnalysisStore(tmp_path / "preserved")
        accepted = asyncio.run(run_evidence_workflow(preserved_store, "capacity expansion cost", _Provider(url), retrieval=PageRetrievalCapability(fixture_hosts={"127.0.0.1"}), human_decision="ACCEPT"))
        rejected = asyncio.run(run_evidence_workflow(preserved_store, "capacity expansion cost", _Provider(url), retrieval=PageRetrievalCapability(fixture_hosts={"127.0.0.1"})))
        assert accepted.status == "accepted"
        assert rejected.status == "rejected"
        assert preserved_store.analysis_path.exists()
        assert preserved_store.load_analysis().revision_status.value == "ACCEPTED"
        assert preserved_store.load_analysis().evidence[0].id == accepted.analysis.evidence[0].id

        pending_twice = AnalysisStore(tmp_path / "pending-twice")
        first_pending = asyncio.run(run_evidence_workflow(pending_twice, "capacity expansion cost", _Provider(url), retrieval=PageRetrievalCapability(fixture_hosts={"127.0.0.1"}), human_decision="ACCEPT"))
        second_pending = asyncio.run(run_evidence_workflow(pending_twice, "capacity expansion cost", _Provider(url), retrieval=PageRetrievalCapability(fixture_hosts={"127.0.0.1"}), human_decision="PENDING"))
        third_pending = asyncio.run(run_evidence_workflow(pending_twice, "capacity expansion cost", _Provider(url), retrieval=PageRetrievalCapability(fixture_hosts={"127.0.0.1"}), human_decision="PENDING"))
        assert first_pending.status == "accepted"
        assert second_pending.status == third_pending.status == "pending"
        with pytest.raises(ValueError, match="stale pending"):
            pending_twice.commit_revision(second_pending.analysis, expected_revision=second_pending.analysis.revision, accepted=True)
        rejected = accept_pending_evidence(pending_twice, human_decision="REJECT")
        assert rejected.status == "rejected"
        assert pending_twice.load_analysis().revision_status.value == "ACCEPTED"
        assert pending_twice.load_analysis().evidence[0].id == first_pending.analysis.evidence[0].id
        with pytest.raises(ValueError, match="stale pending identity"):
            pending_twice.commit_revision(third_pending.analysis, expected_revision=third_pending.analysis.revision, accepted=True)
    finally:
        server.shutdown()
        thread.join(timeout=2)


def test_superseded_pending_token_cannot_be_replayed(tmp_path: Path):
    store = AnalysisStore(tmp_path / "superseded")
    state = AnalysisState("case", 1, revisions=[Revision("revision-1", 1, "initial", [])])
    first = store.save_pending(state)
    second = store.save_pending(state)

    store.commit_revision(second, expected_revision=1, accepted=True)

    with pytest.raises(ValueError, match="stale pending identity"):
        store.commit_revision(first, expected_revision=1, accepted=True)


def test_xlsx_finding_workflow_persists_provenance_calculation_and_restart(tmp_path: Path):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "finding.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Region", "Value"])
    sheet.append(["North", 10])
    sheet.append(["South", 20])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:B3"))
    workbook.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    store = AnalysisStore(tmp_path / "xlsx-case")
    first = run_xlsx_finding_workflow(
        store,
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        human_decision="PENDING",
    )
    assert first.status == "pending"
    assert first.analysis.calculations[0].input_ids
    resumed = accept_pending_evidence(store, human_decision="ACCEPT")
    assert resumed.status == "accepted"
    assert resumed.analysis.calculations[0].input_ids
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    from foresight_os.domain.challenge import derive_finding_views

    views = derive_finding_views(resumed.analysis)
    assert len(views) == 1
    assert views[0].accepted is True
    assert views[0].review_status == "CURRENT"


def test_xlsx_finding_workflow_rolls_back_analysis_cache_when_persistence_rejects(tmp_path: Path):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "rejected-publication.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Region", "Value"])
    sheet.append(["North", 10])
    sheet.append(["South", 20])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:B3"))
    workbook.save(source)

    class RejectingStore:
        def save_pending(self, _state):
            raise RuntimeError("persistence rejected")

    context = WorkbookContext()
    result = run_xlsx_finding_workflow(
        RejectingStore(),
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        workbook_context=context,
        human_decision="ACCEPT",
    )

    assert result.status == "failed"
    assert context.cache.analysis_bytes == 0
    assert context.stats.retained_analysis_result_bytes == 0


def test_xlsx_finding_workflow_rejects_durable_publication_before_budget_rollback(tmp_path: Path):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "budget-rejected-publication.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Region", "Value"])
    sheet.append(["North", 10])
    sheet.append(["South", 20])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:B3"))
    workbook.save(source)
    store = AnalysisStore(tmp_path / "budget-case")
    context = WorkbookContext(max_total_bytes=1)

    result = run_xlsx_finding_workflow(
        store,
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        workbook_context=context,
        human_decision="ACCEPT",
    )

    assert result.status == "failed"
    assert not store.analysis_path.exists()
    assert not store.pending_analysis_path.exists()


def test_xlsx_finding_workflow_uses_selected_table_when_workbook_has_multiple_tables(tmp_path: Path):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "multi-table.xlsx"
    workbook = openpyxl.Workbook()
    brief = workbook.active
    brief.title = "Brief"
    brief.append(["Question"])
    brief.append(["Which table should be analyzed?"])
    brief.tables.add(openpyxl.worksheet.table.Table(displayName="BriefTable", ref="A1:A2"))
    data = workbook.create_sheet("Data")
    data.append(["Region", "Value"])
    data.append(["North", 10])
    data.append(["South", 20])
    data.tables.add(openpyxl.worksheet.table.Table(displayName="DataTable", ref="A1:B3"))
    workbook.save(source)

    result = run_xlsx_finding_workflow(
        AnalysisStore(tmp_path / "multi-table-case"),
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Data", "DataTable"),)),
        human_decision="PENDING",
    )

    assert result.status == "pending"
    assert result.selected_source == "Data"
    assert result.analysis is not None
    assert result.analysis.calculations[0].output == 30.0


def test_xlsx_finding_workflow_reuses_run_workbook_context(tmp_path: Path):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "warm-finding.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Region", "Value"])
    sheet.append(["North", 10])
    sheet.append(["South", 20])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:B3"))
    workbook.save(source)
    selection = SheetSelection(tables=(("Demand", "DemandTable"),))
    context = WorkbookContext()

    first = run_xlsx_finding_workflow(
        AnalysisStore(tmp_path / "warm-one"),
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=selection,
        workbook_context=context,
        human_decision="PENDING",
    )
    second = run_xlsx_finding_workflow(
        AnalysisStore(tmp_path / "warm-two"),
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=selection,
        workbook_context=context,
        human_decision="PENDING",
    )

    assert first.status == second.status == "pending"
    assert context.stats.capture_count == 1
    assert context.stats.capture_hash_count == 1
    assert context.stats.xlsx_parse_count == 1
    assert context.stats.cache_hits == 1
    assert context.stats.analysis_cache_hits == 1


def test_xlsx_finding_workflow_rejects_snapshot_and_context_together(tmp_path: Path):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "ambiguous-finding.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Value"])
    sheet.append([10])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:A2"))
    workbook.save(source)

    result = run_xlsx_finding_workflow(
        AnalysisStore(tmp_path / "ambiguous-case"),
        source,
        AnalysisOperation("sum", value_column="Value"),
        snapshot=SourceSnapshot.capture(source),
        workbook_context=WorkbookContext(),
    )

    assert result.status == "failed"
    assert result.failure == "ValueError"


def test_xlsx_finding_workflow_rejects_oversized_standalone_input_before_parse(tmp_path: Path, monkeypatch):
    from foresight_os.analysis import evidence_workflow
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "oversized.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Value"])
    sheet.append([10])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:A2"))
    workbook.save(source)

    monkeypatch.setattr(evidence_workflow, "spreadsheet_read", lambda *_args, **_kwargs: pytest.fail("parse should not run"))
    result = run_xlsx_finding_workflow(
        AnalysisStore(tmp_path / "oversized-case"),
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        limits=ToolbeltLimits(max_bytes=1),
        human_decision="PENDING",
    )

    assert result.status == "failed"
    assert result.failure == "ToolbeltError"


def test_xlsx_finding_workflow_closes_only_workflow_owned_context(tmp_path: Path, monkeypatch):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "owned-context.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Value"])
    sheet.append([10])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:A2"))
    workbook.save(source)
    cleared = []
    original_clear = WorkbookContext.clear

    def record_clear(context):
        cleared.append(context)
        original_clear(context)

    monkeypatch.setattr(WorkbookContext, "clear", record_clear)
    result = run_xlsx_finding_workflow(
        AnalysisStore(tmp_path / "owned-context-case"),
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        human_decision="PENDING",
    )

    assert result.status == "pending"
    assert len(cleared) == 1
    assert cleared[0]._closed is True


def test_xlsx_finding_workflow_keeps_caller_context_open_and_reuses_snapshot(tmp_path: Path, monkeypatch):
    from foresight_os.analysis import evidence_workflow
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "caller-context.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Value"])
    sheet.append([10])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:A2"))
    workbook.save(source)
    context = WorkbookContext()
    snapshot = SourceSnapshot.capture(source)

    def fail_capture(*_args, **_kwargs):
        raise AssertionError("caller snapshot must not be recaptured")

    monkeypatch.setattr(evidence_workflow.SourceSnapshot, "capture", fail_capture)
    first = run_xlsx_finding_workflow(
        AnalysisStore(tmp_path / "caller-context-case"),
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        snapshot=snapshot,
        human_decision="PENDING",
    )
    monkeypatch.undo()
    second = run_xlsx_finding_workflow(
        AnalysisStore(tmp_path / "caller-context-case-two"),
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        workbook_context=context,
        human_decision="PENDING",
    )

    assert first.status == second.status == "pending"
    assert context._closed is False


def test_xlsx_finding_workflow_cancellation_rolls_back_before_publication(tmp_path: Path, monkeypatch):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "cancelled.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Value"])
    sheet.append([10])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:A2"))
    workbook.save(source)
    store = AnalysisStore(tmp_path / "cancelled-case")
    context = WorkbookContext()
    prior = run_xlsx_finding_workflow(
        store,
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        workbook_context=context,
        human_decision="ACCEPT",
    )
    before_retained = (context.cache.size, context.cache.analysis_bytes, context.authority_index.estimated_bytes, len(context.snapshots))
    cancelled = False

    def cancel_after_analysis():
        nonlocal cancelled
        cancelled = True

    original_enforce_budget = context.enforce_budget

    def enforce_then_cancel():
        original_enforce_budget()
        cancel_after_analysis()

    monkeypatch.setattr(context, "enforce_budget", enforce_then_cancel)
    result = run_xlsx_finding_workflow(
        store,
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        workbook_context=context,
        cancel_check=lambda: cancelled,
        human_decision="ACCEPT",
    )

    assert prior.status == "accepted"
    assert result.status == "failed"
    assert result.failure == "ToolbeltError"
    assert store.load_analysis().revision_status.value == "ACCEPTED"
    assert (context.cache.size, context.cache.analysis_bytes, context.authority_index.estimated_bytes, len(context.snapshots)) == before_retained
    assert context._closed is False


def test_xlsx_finding_workflow_rolls_back_store_when_cancelled_after_pending_save(tmp_path: Path, monkeypatch):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "cancel-after-pending.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Value"])
    sheet.append([10])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:A2"))
    workbook.save(source)
    store = AnalysisStore(tmp_path / "cancel-after-pending-case")
    prior = run_xlsx_finding_workflow(
        store,
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        human_decision="ACCEPT",
    )
    accepted_before = store.analysis_path.read_bytes()
    cancelled = False
    original_save_pending = store.save_pending

    def save_then_cancel(state):
        nonlocal cancelled
        result = original_save_pending(state)
        cancelled = True
        return result

    monkeypatch.setattr(store, "save_pending", save_then_cancel)
    result = run_xlsx_finding_workflow(
        store,
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        cancel_check=lambda: cancelled,
        human_decision="ACCEPT",
    )

    assert prior.status == "accepted"
    assert result.status == "failed"
    assert result.failure == "ToolbeltError"
    assert store.analysis_path.read_bytes() == accepted_before
    assert not store.pending_analysis_path.exists()


def test_xlsx_finding_workflow_settles_accepted_when_cancelled_after_commit(tmp_path: Path, monkeypatch):
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    source = tmp_path / "cancel-after-commit.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Demand"
    sheet.append(["Value"])
    sheet.append([10])
    sheet.tables.add(openpyxl.worksheet.table.Table(displayName="DemandTable", ref="A1:A2"))
    workbook.save(source)
    store = AnalysisStore(tmp_path / "cancel-after-commit-case")
    cancelled = False
    original_commit_revision = store.commit_revision

    def commit_then_cancel(state, expected_revision, accepted):
        nonlocal cancelled
        result = original_commit_revision(state, expected_revision, accepted)
        cancelled = True
        return result

    monkeypatch.setattr(store, "commit_revision", commit_then_cancel)
    result = run_xlsx_finding_workflow(
        store,
        source,
        AnalysisOperation("sum", value_column="Value"),
        selection=SheetSelection(tables=(("Demand", "DemandTable"),)),
        cancel_check=lambda: cancelled,
        human_decision="ACCEPT",
    )

    assert result.status == "accepted"
    assert store.load_analysis().revision_status.value == "ACCEPTED"
    assert not store.pending_analysis_path.exists()
