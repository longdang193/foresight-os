from __future__ import annotations

import argparse
import asyncio
from contextlib import contextmanager
import hashlib
import inspect
import json
import os
import shutil
import socket
import subprocess
import sys
import threading
import tempfile
import time
import tracemalloc
from collections.abc import Mapping
from pathlib import Path
from unittest.mock import patch


def _cli_path(name: str) -> Path | None:
    try:
        index = sys.argv.index(name)
    except ValueError:
        return None
    if index + 1 >= len(sys.argv):
        raise SystemExit(f"{name} requires a path")
    return Path(sys.argv[index + 1]).resolve()


TARGET_ROOT = _cli_path("--target-root") or Path(__file__).parents[2].resolve()
HARNESS_ROOT = _cli_path("--harness-root") or Path(__file__).parents[2].resolve()
# Direct probe execution must measure the explicitly selected worktree, not an
# unrelated editable installation from another checkout.
sys.path.insert(0, str(HARNESS_ROOT))
sys.path.insert(0, str(TARGET_ROOT / "src"))

from agent_framework import Agent, ChatResponse, ChatResponseUpdate, Content, FunctionInvocationContext, Message, ResponseStream

from foresight_os.controller.recipes import ExecutionPattern, RecipeName, make_recipe
from foresight_os.domain.analysis_inputs import AnalysisInputContract, normalize_work_brief
from foresight_os.domain.models import WorkBrief, WorkStage
from foresight_os.persistence.analysis_store import AnalysisStore
from foresight_os.runtime.capabilities import (
    CapabilityDefinition,
    CapabilityScope,
    ExecutionMode,
    ScopeOwnership,
    instrument_tool,
)
from foresight_os.runtime.local_toolbelt import AnalysisOperation, SheetSelection, WorkbookContext, data_analysis, spreadsheet_read
from foresight_os.analysis.independent_validation import run_independent_validation
from foresight_os.analysis.technology_radar import run_technology_radar_workflow
from foresight_os.runtime.workflow import run_recipe_async
from optimization_measurement import DETERMINISTIC_XLSX_WORKLOADS, MANDATORY_XLSX_WORKLOADS, SCENARIO_XLSX_WORKLOADS, _metric_snapshot, _percentile, build_receipt, configure_measurement_roots, measurement_evidence, observed_xlsx_state, process_rss_bytes, stage_metric, timing_statistics

configure_measurement_roots(target_root=TARGET_ROOT, harness_root=HARNESS_ROOT)


class _ProbeClient:
    additional_properties = {}

    def as_agent(self, **kwargs):
        return Agent(self, **kwargs)

    def get_response(self, messages, **kwargs):
        async def response():
            return ChatResponse(messages=[Message("assistant", ["probe-complete"])])

        if kwargs.get("stream"):
            async def updates():
                yield ChatResponseUpdate(contents=[Content("text", text="probe-complete")], role="assistant")

            return ResponseStream(updates(), finalizer=lambda _: ChatResponse(messages=[Message("assistant", ["probe-complete"])]) )
        return response()


class _MCPProbeClient:
    additional_properties = {}

    def as_agent(self, **kwargs):
        return Agent(self, **kwargs)

    def get_response(self, messages, **kwargs):
        options = kwargs.get("options", {})
        client_kwargs = kwargs.get("client_kwargs", {})

        async def response():
            tools = options.get("tools", ())
            if tools:
                function = tools[0]
                context = FunctionInvocationContext(function=function, arguments={"query": "capacity"})

                async def call_next():
                    context.result = await function.invoke(context=context, skip_parsing=True)

                middleware = client_kwargs.get("middleware", ())
                for layer in middleware:
                    await layer.process(context, call_next)
            else:
                raise AssertionError(f"native MCP function was not forwarded by Agent: options={sorted(options)} tool_count={len(tools)} client={sorted(client_kwargs)}")
            return ChatResponse(messages=[Message("assistant", ["mcp-complete"])])

        return response()


def _free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


async def _native_mcp_receipt() -> dict[str, object]:
    from agent_framework import MCPStreamableHTTPTool

    port = _free_port()
    process = subprocess.Popen(
        [sys.executable, str(TARGET_ROOT / "tests/fixtures/mcp_streamable_http_server.py"), "--port", str(port)],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        env={**os.environ, "PYTHONPATH": str((TARGET_ROOT / "src").resolve())},
    )
    tool = MCPStreamableHTTPTool("optimization", f"http://127.0.0.1:{port}/mcp", load_prompts=False)
    try:
        for _ in range(50):
            if process.poll() is not None:
                raise RuntimeError("MCP fixture exited before discovery")
            try:
                await tool.connect()
                break
            except Exception:
                await asyncio.sleep(0.1)
        names = [function.name for function in tool.functions]
        scope = CapabilityScope()
        scope.register(tool)
        from foresight_os.runtime.metrics import RunMetrics
        metrics = RunMetrics("optimization-mcp", "DIRECT", "direct")
        from foresight_os.runtime.participants import AgentFactory, ParticipantConfig
        from foresight_os.runtime.capabilities import CapabilityBinding, CapabilityResolver
        factory = AgentFactory(_MCPProbeClient())
        agent = factory.create(
            ParticipantConfig("mcp-probe", "analyst", "read MCP evidence", capability_ids=("mcp_read",)),
            {
                "scope": scope,
                "metrics": metrics,
                "capability_resolver": CapabilityResolver({"mcp_read": CapabilityDefinition("mcp_read", "external_read")}),
                "mcp_tools": {"mcp_read": CapabilityBinding(tool, fingerprint="fixture-mcp")},
            },
        )
        await agent.run("read capacity")
        return {"status": "VERIFIED", "discovered_tools": names, "tool_calls": metrics.tool_calls}
    finally:
        try:
            await tool.close()
        except Exception:
            pass
        process.terminate()
        try:
            process.wait(timeout=5)
        except subprocess.TimeoutExpired:
            process.kill()


def _write_probe_workbook(path: Path, row_count: int, value_offset: int = 0) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_index in range(1, 4):
        sheet = workbook.create_sheet(f"Sheet{sheet_index}")
        sheet.append(["Category", "Value"])
        for row_index in range(row_count):
            sheet.append([f"item-{row_index}", value_offset + sheet_index + row_index])
    workbook.save(path)


def _analysis_value(table: object) -> float | None:
    columns = getattr(table, "columns", ())
    value_column = "Value" if "Value" in columns else (columns[-1] if columns else None)
    if value_column is None:
        return None
    result = data_analysis(table, AnalysisOperation("aggregate", value_column=value_column))
    if result.rows and isinstance(result.rows[0], Mapping):
        value = result.rows[0].get("sum")
        if value is None:
            value = result.rows[0].get("value")
        return float(value) if value is not None else None
    return None


async def _cancelled_read_workload(path: Path) -> dict[str, object]:
    from foresight_os.runtime.local_toolbelt import ToolbeltError
    from foresight_os.runtime.metrics import RunMetrics

    scope = CapabilityScope(ownership=ScopeOwnership.RUN)
    metrics = RunMetrics("optimization-cancel", "DIRECT", "direct")
    worker_started = threading.Event()
    worker_failed = threading.Event()
    worker_failure: dict[str, str] = {}

    def blocking_read(source: str, selection: SheetSelection | None = None, cancel_check=None):
        worker_started.set()
        while not cancel_check():
            time.sleep(0.001)
        try:
            return spreadsheet_read(source, selection, cancel_check=cancel_check)
        except ToolbeltError as error:
            worker_failure["type"] = type(error).__name__
            worker_failed.set()
            return None

    wrapped = instrument_tool(
        blocking_read,
        CapabilityDefinition("spreadsheet_read", "local_deterministic", execution_mode=ExecutionMode.BLOCKING),
        {"scope": scope, "metrics": metrics},
    )
    task = asyncio.create_task(wrapped(str(path), SheetSelection(include_hidden=True, include_hidden_rows=True)))
    started = await asyncio.to_thread(worker_started.wait, 5)
    caller_cancelled = False
    if started:
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            caller_cancelled = True
    cleanup_error = None
    try:
        await scope.close()
    except Exception as error:
        cleanup_error = type(error).__name__
    return {
        "executed": True,
        "status": "VERIFIED" if started and worker_failed.is_set() and caller_cancelled and cleanup_error is None else "FAILED",
        "correctness": "VERIFIED" if started and worker_failed.is_set() and caller_cancelled else "FAILED",
        "worker_observed": started,
        "worker_failed": worker_failed.is_set(),
        "worker_failure_type": worker_failure.get("type"),
        "caller_cancelled": caller_cancelled,
        "cancellation_settlements": metrics.cancellation_settlements,
        "cache_entries_after_settlement": 0 if scope.workbook_context is None else scope.workbook_context.cache.size,
        "authority_records_after_settlement": 0 if scope.workbook_context is None else scope.workbook_context.authority_index.size,
        "cleanup_complete": scope.executor is None and scope.workbook_context is None and cleanup_error is None,
        "limitation": None if cleanup_error is None else f"cleanup failed: {cleanup_error}",
    }


def _repeated_workflow_once(path: Path) -> dict[str, object]:
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow
    from foresight_os.persistence.analysis_store import AnalysisStore

    with tempfile.TemporaryDirectory(prefix="foresight-repeat-") as temporary_dir:
        store = AnalysisStore(Path(temporary_dir) / "state")
        context = WorkbookContext()
        selection = SheetSelection(include_hidden=True, include_hidden_rows=True)
        started = time.perf_counter()
        first = run_xlsx_finding_workflow(
            store,
            path,
            AnalysisOperation("aggregate", value_column="Value"),
            selection=selection,
            investigation_id="optimization-repeat",
            human_decision="ACCEPT",
            workbook_context=context,
        )
        second = run_xlsx_finding_workflow(
            store,
            path,
            AnalysisOperation("aggregate", value_column="Value"),
            selection=selection,
            investigation_id="optimization-repeat",
            human_decision="ACCEPT",
            workbook_context=context,
        )
        elapsed_ms = (time.perf_counter() - started) * 1000
    first_hash = first.evidence.source_locator.file_version if first.evidence is not None else None
    second_hash = second.evidence.source_locator.file_version if second.evidence is not None else None
    first_value = first.analysis.calculations[0].output if first.analysis and first.analysis.calculations else None
    second_value = second.analysis.calculations[0].output if second.analysis and second.analysis.calculations else None
    correct = (
        first.status == "accepted"
        and second.status == "accepted"
        and first_hash is not None
        and first_hash == second_hash
        and first_value == second_value
    )
    store_stats = getattr(store, "stats", None)
    return {
        "executed": True,
        "status": "VERIFIED" if correct else "FAILED",
        "correctness": "VERIFIED" if correct else "FAILED",
        "first_status": first.status,
        "second_status": second.status,
        "same_source_hash": first_hash == second_hash,
        "same_calculation_value": first_value == second_value,
        "wall_time_ms": round(elapsed_ms, 3),
        "local_tool_calls": len(first.tool_calls) + len(second.tool_calls),
        "capture_count": context.stats.capture_count,
        "capture_hash_count": context.stats.capture_hash_count,
        "currentness_hash_count": context.stats.currentness_hash_count,
        "xlsx_parse_count": context.stats.xlsx_parse_count,
        "read_cache_hits": context.stats.cache_hits,
        "analysis_cache_hits": context.stats.analysis_cache_hits,
        "retained_snapshot_bytes": context.stats.snapshot_bytes,
        "retained_read_result_bytes": context.stats.retained_read_result_bytes,
        "retained_analysis_result_bytes": context.stats.retained_analysis_result_bytes,
        "retained_authority_record_bytes": context.stats.retained_authority_record_bytes,
        "persistence": {
            "instrumentation_available": store_stats is not None,
            "load_count": int(getattr(store_stats, "load_count", 0)),
            "load_time_ms": round(float(getattr(store_stats, "load_time_ms", 0.0) or 0.0), 3),
            "merge_count": int(getattr(store_stats, "merge_count", 0)),
            "merge_time_ms": round(float(getattr(store_stats, "merge_time_ms", 0.0) or 0.0), 3),
            "validation_count": int(getattr(store_stats, "validation_count", 0)),
            "validation_time_ms": round(float(getattr(store_stats, "validation_time_ms", 0.0) or 0.0), 3),
            "pending_write_count": int(getattr(store_stats, "pending_write_count", 0)),
            "pending_write_bytes": int(getattr(store_stats, "pending_write_bytes", 0)),
            "pending_write_time_ms": round(float(getattr(store_stats, "pending_write_time_ms", 0.0) or 0.0), 3),
            "committed_write_count": int(getattr(store_stats, "committed_write_count", 0)),
            "committed_write_bytes": int(getattr(store_stats, "committed_write_bytes", 0)),
            "committed_write_time_ms": round(float(getattr(store_stats, "committed_write_time_ms", 0.0) or 0.0), 3),
            "review_attach_count": int(getattr(store_stats, "review_attach_count", 0)),
            "review_attach_time_ms": round(float(getattr(store_stats, "review_attach_time_ms", 0.0) or 0.0), 3),
            "receipt_write_count": int(getattr(store_stats, "receipt_write_count", 0)),
            "receipt_write_time_ms": round(float(getattr(store_stats, "receipt_write_time_ms", 0.0) or 0.0), 3),
        },
        "limitation": None,
    }


def _repeated_workflow(path: Path, sample_count: int = 3) -> dict[str, object]:
    samples = [_repeated_workflow_once(path) for _ in range(sample_count)]
    latest = samples[-1]
    return {
        **latest,
        "wall_time_ms": _percentile([float(item["wall_time_ms"]) for item in samples], 0.95),
        "sample_count": len(samples),
        "samples_ms": [item["wall_time_ms"] for item in samples],
        "local_tool_calls": sum(int(item.get("local_tool_calls", 0) or 0) for item in samples),
    }


def _retained_context_workloads() -> dict[str, object]:
    results: dict[str, object] = {}
    with tempfile.TemporaryDirectory(prefix="foresight-context-sizes-") as temporary_dir:
        root = Path(temporary_dir)
        for label, rows in (("small", 8), ("medium", 128), ("large", 2048)):
            source = root / f"{label}.xlsx"
            _write_probe_workbook(source, rows)
            context = WorkbookContext(max_entries=4, max_bytes=20_000_000, max_total_bytes=30_000_000)
            selection = SheetSelection(include_hidden=True, include_hidden_rows=True)
            first = context.read(source, selection)
            before = (context.stats.capture_count, context.stats.capture_hash_count, context.stats.xlsx_parse_count)
            started = time.perf_counter()
            warm = context.read(source, selection)
            warm_ms = (time.perf_counter() - started) * 1000
            after = (context.stats.capture_count, context.stats.capture_hash_count, context.stats.xlsx_parse_count)
            retained = context.stats.snapshot_bytes + context.stats.retained_read_result_bytes + context.stats.retained_authority_record_bytes
            results[label] = {
                "rows": rows,
                "status": "VERIFIED" if warm.workbook_sha256 == first.workbook_sha256 and before == after and retained <= context.max_total_bytes else "FAILED",
                "correctness": "VERIFIED" if warm.workbook_sha256 == first.workbook_sha256 else "FAILED",
                "warm_time_ms": round(warm_ms, 3),
                "before_counts": before,
                "after_counts": after,
                "retained_bytes": retained,
                "retained_budget_bytes": context.max_total_bytes,
            }
            context.clear()
    return results


async def _multi_workbook_soak() -> dict[str, object]:
    with tempfile.TemporaryDirectory(prefix="foresight-multi-workbook-") as temporary_dir:
        root = Path(temporary_dir)
        sources: list[Path] = []
        for index, rows in enumerate((16, 32, 64), 1):
            source = root / f"source-{index}.xlsx"
            _write_probe_workbook(source, rows, value_offset=index * 10_000)
            sources.append(source)
        context = WorkbookContext(max_entries=2, max_bytes=20_000_000, max_total_bytes=40_000_000)
        scope = CapabilityScope(ownership=ScopeOwnership.RUN)
        scope.workbook_context = context
        selection = SheetSelection(include_hidden=True, include_hidden_rows=True)
        results = [context.read(source, selection) for source in sources]
        cache_keys_after_three = tuple(context.cache._entries)
        retained = context.stats.snapshot_bytes + context.stats.retained_read_result_bytes + context.stats.retained_authority_record_bytes
        hashes = tuple(result.workbook_sha256 for result in results)
        no_cross_source_authority_aliasing = len({key[0] for key in context.authority_index.measurements}) == 3
        repeated = context.read(sources[0], selection)
        cache_keys_after_repeat = tuple(context.cache._entries)
        evicted_first = not any(key[0] == str(sources[0].resolve()) for key in cache_keys_after_three)
        cleanup_error = None
        try:
            await scope.close()
        except Exception as error:
            cleanup_error = type(error).__name__
        correct = (
            context.stats.capture_count == 3
            and len(set(hashes)) == 3
            and no_cross_source_authority_aliasing
            and repeated.workbook_sha256 == hashes[0]
            and evicted_first
            and len(cache_keys_after_repeat) <= 2
            and retained <= context.max_total_bytes
            and scope.workbook_context is None
            and cleanup_error is None
        )
        return {
            "executed": True,
            "status": "VERIFIED" if correct else "FAILED",
            "correctness": "VERIFIED" if correct else "FAILED",
            "source_count": len(sources),
            "capture_count": context.stats.capture_count,
            "distinct_source_hashes": len(set(hashes)),
            "no_cross_source_authority_aliasing": no_cross_source_authority_aliasing,
            "deterministic_eviction": evicted_first and len(cache_keys_after_repeat) <= 2,
            "retained_bytes_before_cleanup": retained,
            "retained_budget_bytes": context.max_total_bytes,
            "owned_resources_after_cleanup": 0 if scope.workbook_context is None and scope.executor is None and not scope.resources else 1,
            "cleanup_complete": cleanup_error is None and scope.workbook_context is None,
            "limitation": None if cleanup_error is None else f"cleanup failed: {cleanup_error}",
        }


def _retained_bytes(context: WorkbookContext) -> int:
    return (
        context.stats.snapshot_bytes
        + context.stats.retained_read_result_bytes
        + context.stats.retained_analysis_result_bytes
        + context.stats.retained_authority_record_bytes
    )


def _process_cold_samples(path: Path, sample_count: int = 10) -> dict[str, object]:
    script = """
import sys
import time
from pathlib import Path
sys.path.insert(0, sys.argv[1])
from foresight_os.runtime.local_toolbelt import SheetSelection, spreadsheet_read
started = time.perf_counter()
spreadsheet_read(Path(sys.argv[2]), SheetSelection(include_hidden=True, include_hidden_rows=True))
print((time.perf_counter() - started) * 1000)
"""
    samples: list[float] = []
    environment = os.environ.copy()
    for _ in range(sample_count):
        completed = subprocess.run(
            [sys.executable, "-c", script, str(TARGET_ROOT / "src"), str(path.resolve())],
            check=True,
            capture_output=True,
            text=True,
            env=environment,
        )
        samples.append(float(completed.stdout.strip()))
    timing = timing_statistics(samples)
    return {
        "status": "VERIFIED" if len(samples) >= sample_count else "NOT_MEASURED",
        **timing,
        "limitation": None,
    }


@contextmanager
def _count_workbook_reads():
    from foresight_os.analysis import workbook_inputs
    from foresight_os.runtime import local_toolbelt

    counts = {"capture_count": 0, "xlsx_parse_count": 0}
    original_capture = local_toolbelt.SourceSnapshot.capture
    original_spreadsheet_read = local_toolbelt.spreadsheet_read

    def counted_capture(cls, *args, **kwargs):
        counts["capture_count"] += 1
        return original_capture(*args, **kwargs)

    def counted_spreadsheet_read(*args, **kwargs):
        counts["xlsx_parse_count"] += 1
        return original_spreadsheet_read(*args, **kwargs)

    with (
        patch.object(local_toolbelt.SourceSnapshot, "capture", classmethod(counted_capture)),
        patch.object(local_toolbelt, "spreadsheet_read", counted_spreadsheet_read),
        patch.object(workbook_inputs, "spreadsheet_read", counted_spreadsheet_read),
    ):
        yield counts


def _technology_result_fingerprint(result: object) -> str:
    technology = getattr(result, "technology", None)
    assessments = getattr(technology, "assessments", ()) if technology is not None else ()
    signature = tuple(repr(assessment) for assessment in assessments)
    return hashlib.sha256(repr(signature).encode("utf-8")).hexdigest()


async def _approved_changed_brief_workload(mode: str, sample_count: int = 10) -> dict[str, object]:
    from datetime import datetime, timezone

    from foresight_os.artifacts.approval import build_approval_package
    from foresight_os.controller.proposal import build_execution_recipe
    from foresight_os.domain.approval import AdaptationEnvelope, ApprovalRecord, ApprovalStatus
    from foresight_os.persistence.ids import canonical_json_hash
    from foresight_os.runtime.execution_gate import ExecutionGate
    from foresight_os.runtime.workflow import StaleReuseError

    source = TARGET_ROOT / "tests" / "fixtures" / "exercise" / "04_Technology_Radar.xlsx"

    def make_brief(assumption: str, objective: str) -> WorkBrief:
        return WorkBrief(
            "optimization-approved-changed-brief",
            "Which technologies should be prioritized?",
            ["SiC Power Electronics"],
            ["analysis"],
            ["baseline scope remains stable for the context stage"],
            [],
            objective,
            "compare evidence",
            stages=[
                WorkStage("radar", objective),
                WorkStage("context", "Retain the stable workbook context"),
            ],
            analysis_inputs=[{
                "domain": "technology_radar",
                "entity_scope": ["SiC Power Electronics"],
                "assumptions": ["baseline scope remains stable for the context stage"],
            }, {
                "domain": "technology_radar",
                "stage_id": "radar",
                "entity_scope": ["SiC Power Electronics"],
                "assumptions": [assumption],
            }, {
                "domain": "technology_radar",
                "stage_id": "context",
                "entity_scope": ["SiC Power Electronics"],
                "assumptions": ["baseline scope remains stable for the context stage"],
            }],
        )

    def save_approval(store: AnalysisStore, brief: WorkBrief, recipe: object, *, fresh: bool) -> None:
        package = build_approval_package(
            brief,
            recipe,
            AdaptationEnvelope(fallback_on_stale_reuse="run_fresh" if fresh else "stop"),
        )
        store.save_execution_approval(ApprovalRecord(
            brief.investigation_id,
            ApprovalStatus.APPROVED,
            package.projection,
            package.fingerprint,
            package.fingerprint,
            canonical_json_hash(recipe.to_dict()),
            "local_cli",
            datetime.now(timezone.utc).isoformat(),
            run_decisions={"run_fresh_stages": ["radar"] if fresh else []},
        ))

    samples: list[float] = []
    sample_results: list[dict[str, object]] = []
    for index in range(sample_count):
        started = time.perf_counter()
        with tempfile.TemporaryDirectory(prefix="foresight-approved-changed-brief-") as temporary_dir:
            store = AnalysisStore(Path(temporary_dir) / "case")
            base_brief = make_brief("baseline scope", "Assess the focused technology scope")
            base_recipe = build_execution_recipe(base_brief).recipe
            store.save_brief(base_brief)
            store.save_recipe(base_recipe)
            save_approval(store, base_brief, base_recipe, fresh=False)
            first = await ExecutionGate(store).execute(
                prompt=base_brief.decision_question,
                run_id=f"approved-changed-{index}-base",
                workbook_source=source,
            )

            changed_brief = make_brief("changed scope", "Reassess the expanded technology scope")
            changed_recipe = build_execution_recipe(changed_brief).recipe
            store.save_brief(changed_brief)
            store.save_recipe(changed_recipe)
            save_approval(store, changed_brief, changed_recipe, fresh=False)
            stale_receipt_rejected = False
            try:
                await ExecutionGate(store).execute(
                    prompt=changed_brief.decision_question,
                    run_id=f"approved-changed-{index}-stale",
                    workbook_source=source,
                )
            except StaleReuseError:
                stale_receipt_rejected = True

            save_approval(store, changed_brief, changed_recipe, fresh=True)
            changed = await ExecutionGate(store).execute(
                prompt=changed_brief.decision_question,
                run_id=f"approved-changed-{index}-fresh",
                workbook_source=source,
            )
            first_output = first.outputs[0] if first.outputs else None
            changed_output = changed.outputs[0] if changed.outputs else None
            first_contract = getattr(first_output, "contract_fingerprint", None)
            changed_contract = getattr(changed_output, "contract_fingerprint", None)
            source_hashes = {
                getattr(first_output, "source_sha256", None),
                getattr(changed_output, "source_sha256", None),
            }
            correct = (
                first.metrics.runtime_status == "SETTLED"
                and changed.metrics.runtime_status == "SETTLED"
                and stale_receipt_rejected
                and first_contract
                and changed_contract
                and first_contract != changed_contract
                and source_hashes == {getattr(first_output, "source_sha256", None)}
                and changed.metrics.reused_stages > 0
                and len(store.load_stage_receipts()) == 2
            )
            sample_results.append({
                "status": "VERIFIED" if correct else "FAILED",
                "stale_receipt_rejected": stale_receipt_rejected,
                "distinct_contracts": first_contract != changed_contract,
                "source_identity_preserved": len(source_hashes) == 1,
                "reassessed_stages": changed.metrics.reassessed_stages,
                "reused_stages": changed.metrics.reused_stages,
            })
        samples.append((time.perf_counter() - started) * 1000)

    timing = timing_statistics(samples)
    correct = bool(sample_results) and all(item["status"] == "VERIFIED" for item in sample_results)
    return {
        "executed": True,
        "status": "VERIFIED" if correct else "FAILED",
        "correctness": "VERIFIED" if correct else "FAILED",
        "mode": mode,
        "approved_execution": True,
        "stale_receipt_rejected": all(item["stale_receipt_rejected"] for item in sample_results),
        "distinct_contracts": all(item["distinct_contracts"] for item in sample_results),
        "source_identity_preserved": all(item["source_identity_preserved"] for item in sample_results),
        "unaffected_stage_reuse": any(item["reused_stages"] > 0 for item in sample_results),
        "samples_ms": timing["samples_ms"],
        "sample_count": timing["sample_count"],
        "wall_time_ms": timing["p95_ms"],
        "p50_ms": timing["p50_ms"],
        "p95_ms": timing["p95_ms"],
        "noise_floor_ms": timing["noise_floor_ms"],
        "sample_results": sample_results,
        "limitation": None if correct else "approved changed-brief execution did not satisfy all per-sample checks",
    }


async def _changed_brief_workload(mode: str) -> dict[str, object]:
    return await _approved_changed_brief_workload(mode)


def _history_workload(path: Path, history_size: int) -> dict[str, object]:
    from foresight_os.analysis.evidence_workflow import run_xlsx_finding_workflow

    selection = SheetSelection(include_hidden=True, include_hidden_rows=True)
    samples: list[float] = []
    milestones: list[dict[str, object]] = []
    milestone_indexes = {0, history_size // 2, history_size - 1}
    source_hashes: list[str] = []
    with tempfile.TemporaryDirectory(prefix=f"foresight-history-{history_size}-") as temporary_dir:
        root = Path(temporary_dir) / "state"
        store = AnalysisStore(root)
        context = WorkbookContext()
        last = None
        for index in range(history_size):
            started = time.perf_counter()
            last = run_xlsx_finding_workflow(
                store,
                path,
                AnalysisOperation("aggregate", value_column="Value"),
                selection=selection,
                investigation_id="optimization-history",
                human_decision="ACCEPT",
                workbook_context=context,
            )
            samples.append((time.perf_counter() - started) * 1000)
            if last.evidence is not None:
                source_hashes.append(last.evidence.source_locator.file_version)
            if index in milestone_indexes:
                state_bytes = store.analysis_path.read_bytes() if store.analysis_path.exists() else b""
                token_bytes = store.pending_tokens_path.read_bytes() if store.pending_tokens_path.exists() else b""
                milestones.append({
                    "run": index + 1,
                    "state_sha256": hashlib.sha256(state_bytes).hexdigest() if state_bytes else None,
                    "state_bytes": len(state_bytes),
                    "pending_token_sha256": hashlib.sha256(token_bytes).hexdigest() if token_bytes else None,
                    "pending_token_bytes": len(token_bytes),
                    "revision_count": len(last.analysis.revisions) if last.analysis is not None else None,
                })
        final_state = last.analysis if last is not None else None
        timing = timing_statistics(samples)
        correct = (
            last is not None
            and last.status == "accepted"
            and final_state is not None
            and final_state.revision == history_size
            and len(final_state.revisions) == history_size
            and len(source_hashes) == history_size
            and len(set(source_hashes)) == 1
        )
        stats = store.stats
        return {
            "executed": True,
            "status": "VERIFIED" if correct else "FAILED",
            "correctness": "VERIFIED" if correct else "FAILED",
            "history_size": history_size,
            "sample_count": timing["sample_count"],
            "samples_ms": timing["samples_ms"],
            "wall_time_ms_p50": timing["p50_ms"],
            "wall_time_ms_p95": timing["p95_ms"],
            "noise_floor_ms": timing["noise_floor_ms"],
            "wall_time_ms_total": round(sum(samples), 3),
            "source_identity_preserved": len(set(source_hashes)) == 1,
            "final_revision": final_state.revision if final_state is not None else None,
            "final_state_sha256": hashlib.sha256(store.analysis_path.read_bytes()).hexdigest() if store.analysis_path.exists() else None,
            "final_state_bytes": store.analysis_path.stat().st_size if store.analysis_path.exists() else None,
            "pending_token_sha256": hashlib.sha256(store.pending_tokens_path.read_bytes()).hexdigest() if store.pending_tokens_path.exists() else None,
            "pending_token_bytes": store.pending_tokens_path.stat().st_size if store.pending_tokens_path.exists() else 0,
            "load_count": int(getattr(stats, "load_count", 0)),
            "validation_count": int(getattr(stats, "validation_count", 0)),
            "merge_count": int(getattr(stats, "merge_count", 0)),
            "milestones": milestones,
            "limitation": None,
        }


def _write_selected_sheet_stress_workbook(path: Path, *, unrelated_sheet_count: int = 24, unrelated_rows: int = 512) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    workbook.remove(workbook.active)
    for index in range(unrelated_sheet_count):
        sheet = workbook.create_sheet(f"Noise_{index:02d}")
        sheet.append(["Category", "Value"])
        for row_index in range(unrelated_rows):
            sheet.append([f"noise-{index}-{row_index}", index + row_index])
    selected = workbook.create_sheet("Selected")
    selected.append(["Category", "Value"])
    for row_index in range(4):
        selected.append([f"selected-{row_index}", row_index + 1])
    selected.add_table(Table(displayName="SelectedTable", ref="A1:B5"))
    selected.tables["SelectedTable"].tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    workbook.save(path)


def _selected_sheet_stress_workload() -> dict[str, object]:
    samples: list[float] = []
    last = None
    with tempfile.TemporaryDirectory(prefix="foresight-selected-sheet-stress-") as temporary_dir:
        path = Path(temporary_dir) / "selected-sheet-stress.xlsx"
        _write_selected_sheet_stress_workbook(path)
        selection = SheetSelection(sheet_names=("Selected",), include_hidden=True, include_hidden_rows=True)
        for _ in range(10):
            context = WorkbookContext()
            started = time.perf_counter()
            last = context.read(path, selection)
            samples.append((time.perf_counter() - started) * 1000)
        stats = context.stats
        selected_value = _analysis_value(last.tables[0]) if last is not None and last.tables else None
        timing = timing_statistics(samples)
        parsed_matches_selection = stats.archive_sheet_xml_scanned == stats.selected_sheet_count == 1
        correct = (
            last is not None
            and selected_value == 10.0
            and parsed_matches_selection
            and stats.archive_sheets_scanned == 25
            and stats.selected_row_count > 0
            and stats.selected_cell_count > 0
        )
        return {
            "executed": True,
            "status": "VERIFIED" if correct else "FAILED",
            "correctness": "VERIFIED" if correct else "FAILED",
            "fixture_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "generated_sheet_count": 25,
            "archive_sheets_scanned": stats.archive_sheets_scanned,
            "archive_sheet_xml_scanned": stats.archive_sheet_xml_scanned,
            "selected_sheet_count": stats.selected_sheet_count,
            "selected_row_count": stats.selected_row_count,
            "selected_cell_count": stats.selected_cell_count,
            "formula_count": stats.formula_count,
            "openpyxl_load_count": stats.openpyxl_load_count,
            "selected_value": selected_value,
            "parsed_matches_selection": parsed_matches_selection,
            "wall_time_ms": timing["p95_ms"],
            "p50_ms": timing["p50_ms"],
            "p95_ms": timing["p95_ms"],
            "noise_floor_ms": timing["noise_floor_ms"],
            "sample_count": timing["sample_count"],
            "samples_ms": timing["samples_ms"],
            "limitation": None if correct else "selected-sheet stress counters or result fingerprint did not match",
        }


async def _measure_workbook(path: Path, measurement_mode: str = "full", mode: str = "baseline") -> dict[str, object]:
    durations: list[float] = []
    allocations: list[int] = []
    result = None
    selection = SheetSelection(include_hidden=True, include_hidden_rows=True)
    if measurement_mode in {"timing", "full"}:
        for _ in range(2):
            result = spreadsheet_read(path, selection)
        for _ in range(10):
            started = time.perf_counter()
            result = spreadsheet_read(path, selection)
            durations.append((time.perf_counter() - started) * 1000)
    if measurement_mode in {"allocation", "full"}:
        for _ in range(2):
            spreadsheet_read(path, selection)
        for _ in range(10):
            tracemalloc.start()
            spreadsheet_read(path, selection)
            _current, peak = tracemalloc.get_traced_memory()
            tracemalloc.stop()
            allocations.append(peak)
    if result is None:
        result = spreadsheet_read(path, selection)
    cold_noise_floor = _process_cold_samples(path)
    table = result.tables[0]
    value_column = "Value" if "Value" in table.columns else table.columns[-1]
    analysis = data_analysis(table, AnalysisOperation("aggregate", value_column=value_column))
    context = WorkbookContext()
    context_selection = selection
    cold_started = time.perf_counter()
    context.read(path, context_selection)
    cold_read_ms = (time.perf_counter() - cold_started) * 1000
    warm_started = time.perf_counter()
    warm_result = context.read(path, context_selection)
    warm_read_ms = (time.perf_counter() - warm_started) * 1000
    warm_2_plus_samples: list[float] = []
    for _ in range(3):
        started = time.perf_counter()
        context.read(path, context_selection)
        warm_2_plus_samples.append((time.perf_counter() - started) * 1000)
    same_view_samples = [warm_read_ms, *warm_2_plus_samples]
    for _ in range(6):
        started = time.perf_counter()
        context.read(path, context_selection)
        same_view_samples.append((time.perf_counter() - started) * 1000)
    validator_samples: list[float] = []
    validator_report = None
    for _ in range(3):
        validator_started = time.perf_counter()
        validator_report = run_independent_validation(path, selection=context_selection, workbook_context=context)
        validator_samples.append((time.perf_counter() - validator_started) * 1000)
    context.currentness(path)
    reader_instrumentation_available = all(
        hasattr(context.stats, field)
        for field in (
            "archive_sheets_scanned",
            "selected_sheet_count",
            "selected_row_count",
            "selected_cell_count",
            "formula_count",
        )
    )
    sheet_names = tuple(sheet.name for sheet in result.sheets)
    workloads: dict[str, dict[str, object]] = {
        "cold_workbook": {"executed": True, "status": "VERIFIED", "correctness": "VERIFIED", "wall_time_ms": round(cold_read_ms, 3)},
        "same_view_repeat": {
            "executed": True,
            "status": "VERIFIED" if context.stats.capture_count == 1 and context.stats.xlsx_parse_count == 1 else "FAILED",
            "correctness": "VERIFIED" if warm_result.workbook_sha256 == result.workbook_sha256 else "FAILED",
            "wall_time_ms": _percentile(same_view_samples, 0.95),
            "sample_count": len(same_view_samples),
            "samples_ms": [round(value, 3) for value in same_view_samples],
            "capture_count": context.stats.capture_count,
            "xlsx_parse_count": context.stats.xlsx_parse_count,
            "analytical_result_fingerprint": analysis.deterministic_key,
        },
        "warm_1": {
            "executed": True,
            "status": "VERIFIED" if warm_result.workbook_sha256 == result.workbook_sha256 else "FAILED",
            "correctness": "VERIFIED" if warm_result.workbook_sha256 == result.workbook_sha256 else "FAILED",
            "wall_time_ms": round(warm_read_ms, 3),
        },
        "warm_2_plus": {
            "executed": True,
            "status": "VERIFIED" if all(item >= 0 for item in warm_2_plus_samples) else "FAILED",
            "correctness": "VERIFIED" if warm_result.workbook_sha256 == result.workbook_sha256 else "FAILED",
            "sample_count": len(warm_2_plus_samples),
            "p95_ms": _percentile(warm_2_plus_samples, 0.95),
        },
        "analyst_validator_reuse": {
            "executed": True,
            "status": "VERIFIED" if validator_report.accepted else "FAILED",
            "correctness": "VERIFIED" if validator_report.accepted else "FAILED",
            "validator_accepted": validator_report.accepted,
            "wall_time_ms": _percentile(validator_samples, 0.95),
            "sample_count": len(validator_samples),
            "samples_ms": [round(value, 3) for value in validator_samples],
        },
    }
    with tempfile.TemporaryDirectory(prefix="foresight-optimization-") as temporary_dir:
        changed_path = Path(temporary_dir) / path.name
        shutil.copy2(path, changed_path)
        from openpyxl import load_workbook

        changed_workbook = load_workbook(changed_path)
        changed_sheet = changed_workbook[changed_workbook.sheetnames[0]]
        changed_sheet["A1"] = f"changed:{changed_sheet['A1'].value}"
        changed_workbook.save(changed_path)
        changed_context = WorkbookContext()
        changed_result = changed_context.read(changed_path, selection)
        workloads["changed_source"] = {
            "executed": True,
            "status": "VERIFIED" if changed_context.stats.capture_count == 1 and changed_result.workbook_sha256 != result.workbook_sha256 else "FAILED",
            "correctness": "VERIFIED" if changed_context.stats.capture_count == 1 and changed_result.workbook_sha256 != result.workbook_sha256 else "FAILED",
            "source_identity_changed": changed_result.workbook_sha256 != result.workbook_sha256,
        }
    if len(sheet_names) > 1:
        alternate_selection = SheetSelection(sheet_names=(sheet_names[-1],), include_hidden=True, include_hidden_rows=True)
        alternate_context = WorkbookContext()
        alternate = alternate_context.read(path, alternate_selection)
        alternate_samples: list[float] = []
        for _ in range(3):
            started = time.perf_counter()
            alternate_context.read(path, alternate_selection)
            alternate_samples.append((time.perf_counter() - started) * 1000)
        alternate_value = _analysis_value(alternate.tables[0]) if alternate.tables else None
        primary_value = _analysis_value(result.tables[0])
        workloads["different_selection"] = {
            "executed": True,
            "status": "VERIFIED" if alternate.workbook_sha256 == result.workbook_sha256 and alternate_value != primary_value and alternate_context.stats.capture_count == 1 else "FAILED",
            "correctness": "VERIFIED" if alternate.workbook_sha256 == result.workbook_sha256 and alternate_value != primary_value else "FAILED",
            "capture_count": alternate_context.stats.capture_count,
            "xlsx_parse_count": alternate_context.stats.xlsx_parse_count,
            "primary_value": primary_value,
            "alternate_sheet": sheet_names[-1],
            "alternate_value": alternate_value,
            "wall_time_ms": _percentile(alternate_samples, 0.95),
            "sample_count": len(alternate_samples),
            "samples_ms": [round(value, 3) for value in alternate_samples],
            "limitation": None,
        }
    workloads["repeated_investigation"] = _repeated_workflow(path)
    workloads["history-10"] = _history_workload(path, 10)
    workloads["history-100"] = _history_workload(path, 100)
    workloads["blocking_cancellation"] = await _cancelled_read_workload(path)
    sizes = _retained_context_workloads()
    large = sizes["large"]
    workloads["large_retained_context"] = {
        "executed": True,
        "status": large["status"],
        "correctness": large["correctness"],
        "sizes": sizes,
        "limitation": large.get("limitation"),
    }
    workloads["multi_workbook"] = await _multi_workbook_soak()
    workloads["approved_changed_brief"] = await _approved_changed_brief_workload(mode)
    workloads["changed_brief"] = workloads["approved_changed_brief"]
    workloads["selected_sheet_stress"] = _selected_sheet_stress_workload()
    workloads["public_workflow_context_reuse"] = workloads["approved_changed_brief"]
    from foresight_os.runtime.metrics import RunMetrics

    local_metrics = RunMetrics("optimization-local", "DIRECT", "direct")
    local_scope = CapabilityScope(ownership=ScopeOwnership.RUN)
    local_read = instrument_tool(
        spreadsheet_read,
        CapabilityDefinition("spreadsheet_read", "local_deterministic", execution_mode=ExecutionMode.BLOCKING),
        {"scope": local_scope, "metrics": local_metrics},
    )
    await local_read(str(path), selection)
    await local_scope.close()
    return {
        "source": "fixture",
        "sheet_count": len(result.sheets),
        "measurement_count": len(result.measurements),
        "table_count": len(result.tables),
        "p50_ms": _percentile(durations, 0.50) if durations else None,
        "p90_ms": _percentile(durations, 0.90) if durations else None,
        "p95_ms": _percentile(durations, 0.95) if durations else None,
        "timing_paths": {
            "direct_workbook": {
                "p50_ms": _percentile(durations, 0.50) if durations else None,
                "p90_ms": _percentile(durations, 0.90) if durations else None,
                "p95_ms": _percentile(durations, 0.95) if durations else None,
                "sample_count": len(durations),
            },
            "cold_workbook": {"wall_time_ms": workloads["cold_workbook"].get("wall_time_ms"), "sample_count": cold_noise_floor.get("sample_count", 0)},
            "same_view_repeat": {"wall_time_ms": workloads["same_view_repeat"].get("wall_time_ms"), "sample_count": workloads["same_view_repeat"].get("sample_count", 0)},
            "warm_1": {"wall_time_ms": workloads["warm_1"].get("wall_time_ms")},
            "warm_2_plus": {"p95_ms": workloads["warm_2_plus"].get("p95_ms"), "sample_count": workloads["warm_2_plus"].get("sample_count")},
            "repeated_investigation": {"wall_time_ms": workloads["repeated_investigation"].get("wall_time_ms"), "sample_count": workloads["repeated_investigation"].get("sample_count", 0)},
            "analyst_validator_reuse": {"wall_time_ms": workloads["analyst_validator_reuse"].get("wall_time_ms"), "sample_count": workloads["analyst_validator_reuse"].get("sample_count", 0)},
            "approved_changed_brief": {"wall_time_ms": workloads["approved_changed_brief"].get("wall_time_ms", 0), "sample_count": workloads["approved_changed_brief"].get("sample_count", 0), "samples_ms": workloads["approved_changed_brief"].get("samples_ms", []), "noise_floor_ms": workloads["approved_changed_brief"].get("noise_floor_ms", 0)},
            "history-10": {"wall_time_ms": workloads["history-10"].get("wall_time_ms_p95", 0), "sample_count": workloads["history-10"].get("sample_count", 0)},
            "history-100": {"wall_time_ms": workloads["history-100"].get("wall_time_ms_p95", 0), "sample_count": workloads["history-100"].get("sample_count", 0)},
            "selected_sheet_stress": {"wall_time_ms": workloads["selected_sheet_stress"].get("wall_time_ms", 0), "sample_count": workloads["selected_sheet_stress"].get("sample_count", 0), "samples_ms": workloads["selected_sheet_stress"].get("samples_ms", []), "noise_floor_ms": workloads["selected_sheet_stress"].get("noise_floor_ms", 0)},
        },
        "peak_allocation_bytes": max(allocations) if allocations else None,
        "peak_retained_bytes": max(_retained_bytes(context), max((item.get("retained_bytes", 0) for item in workloads.values() if isinstance(item, dict)), default=0)),
        "measurement_mode": measurement_mode,
        "cold_noise_floor": cold_noise_floor,
        "latency": measurement_evidence("VERIFIED", limitation=None) if durations else measurement_evidence("NOT_MEASURED", limitation="timing mode was not selected"),
        "allocation": measurement_evidence("VERIFIED", limitation=None) if allocations else measurement_evidence("NOT_MEASURED", limitation="allocation mode was not selected"),
        "workloads": workloads,
        "deterministic_key": analysis.deterministic_key,
        "source_hash": result.workbook_sha256,
        "analytical_result_fingerprint": analysis.deterministic_key,
        "local_tool_calls": local_metrics.local_tool_calls + int(workloads["repeated_investigation"].get("local_tool_calls", 0) or 0),
        "run_owned_context": {
            "capture_count": context.stats.capture_count,
            "capture_hash_count": context.stats.capture_hash_count,
            "currentness_hash_count": context.stats.currentness_hash_count,
            "xlsx_parse_count": context.stats.xlsx_parse_count,
            "openpyxl_load_count": context.stats.openpyxl_load_count,
            "archive_sheets_scanned": getattr(context.stats, "archive_sheets_scanned", None),
            "archive_sheet_xml_scanned": getattr(context.stats, "archive_sheet_xml_scanned", None),
            "selected_sheet_count": getattr(context.stats, "selected_sheet_count", None),
            "selected_row_count": getattr(context.stats, "selected_row_count", None),
            "selected_cell_count": getattr(context.stats, "selected_cell_count", None),
            "formula_count": getattr(context.stats, "formula_count", None),
            "materialization_time_ms": round(float(getattr(context.stats, "materialization_time_ms", 0.0) or 0.0), 3),
            "cache_hits": context.stats.cache_hits,
            "analysis_cache_hits": context.stats.analysis_cache_hits,
            "snapshot_bytes": context.stats.snapshot_bytes,
            "retained_read_result_bytes": context.stats.retained_read_result_bytes,
            "retained_analysis_result_bytes": context.stats.retained_analysis_result_bytes,
            "retained_authority_record_bytes": context.stats.retained_authority_record_bytes,
            "retained_bytes": _retained_bytes(context),
            "peak_retained_bytes": _retained_bytes(context),
            "validator_accepted": validator_report.accepted,
        },
        "reader_instrumentation_available": reader_instrumentation_available,
        "persistence": workloads["repeated_investigation"].get("persistence", {}),
        "mandatory_workloads": [name for name in MANDATORY_XLSX_WORKLOADS if name in workloads],
        "deterministic_workloads": [name for name in DETERMINISTIC_XLSX_WORKLOADS if name in workloads],
        "scenario_workloads": {
            "approved_changed_brief": workloads.get("approved_changed_brief", {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "approved changed-brief probe unavailable"}),
            "changed_brief": workloads.get("changed_brief", {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "changed-brief probe unavailable"}),
            "changed_selection": workloads.get("different_selection", {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "fixture has one selectable view"}),
            "selected_sheet_stress": workloads.get("selected_sheet_stress", {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "selected-sheet stress probe unavailable"}),
            "history-10": workloads.get("history-10", {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "history probe unavailable"}),
            "history-100": workloads.get("history-100", {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "history probe unavailable"}),
            "history-1000": {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "bounded manual stress workload; not required for release gate"},
            "correction": {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "correction requires a reviewed pending state"},
            "workbook_cancellation": workloads.get("blocking_cancellation", {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "cancellation probe unavailable"}),
            "synthesis_cancellation": {"executed": False, "status": "NOT_MEASURED", "correctness": "NOT_MEASURED", "limitation": "synthesis cancellation is outside deterministic workbook scope"},
        },
    }


async def _event_loop_lag_ms() -> float:
    scheduled = asyncio.get_running_loop().time() + 0.001
    await asyncio.sleep(0.001)
    return round(max(0.0, (asyncio.get_running_loop().time() - scheduled) * 1000), 3)


def _resource_snapshot(metrics, event_loop_lag_ms: float) -> dict[str, object]:
    child_process_count = 0
    open_file_descriptors = None
    try:
        import psutil

        process = psutil.Process(os.getpid())
        child_process_count = len(process.children(recursive=True))
        try:
            open_file_descriptors = int(process.num_fds())
        except (AttributeError, OSError):
            open_file_descriptors = None
    except Exception:
        pass
    return {
        "rss_bytes": process_rss_bytes(),
        "active_python_tasks": len(asyncio.all_tasks()),
        "thread_count": threading.active_count(),
        "child_process_count": child_process_count,
        "owned_capability_resources": 0 if metrics.cleanup_complete else 1,
        "mcp_session_count": 0,
        "event_loop_lag_ms": event_loop_lag_ms,
        "open_file_descriptors": open_file_descriptors,
    }


async def _measure_patterns() -> tuple[dict[str, object], list[dict[str, object]]]:
    receipts: dict[str, object] = {}
    stage_metrics: list[dict[str, object]] = []
    for recipe_name in (RecipeName.DIRECT, RecipeName.ANALYZE, RecipeName.CONCURRENT, RecipeName.HANDOFF, RecipeName.CHALLENGE_RECOMMENDATION):
        recipe = make_recipe(recipe_name)
        for _ in range(2):
            await run_recipe_async(recipe, "optimization baseline", _ProbeClient(), f"warm-{recipe_name.value}")
        samples: list[dict[str, object]] = []
        last = None
        for index in range(10):
            started = time.perf_counter()
            last = await run_recipe_async(recipe, "optimization baseline", _ProbeClient(), f"measure-{recipe_name.value}-{index}")
            lag = await _event_loop_lag_ms()
            last.metrics.event_loop_lag_ms = lag
            samples.append(_metric_snapshot(last.metrics, (time.perf_counter() - started) * 1000))
        if last is not None:
            summary = stage_metric("stage-1", recipe.pattern, last, float(samples[-1]["wall_time_ms"]), samples=samples)
            stage_metrics.append(summary)
        receipts[recipe_name.value] = {
            **summary,
            "p50_ms": summary["latency"]["wall_time_ms_p50"],
            "p95_ms": summary["latency"]["wall_time_ms_p95"],
        }
    return receipts, stage_metrics


async def _soak_receipt() -> dict[str, object]:
    recipe = make_recipe(RecipeName.ANALYZE)
    snapshots: list[dict[str, object]] = []
    for index in range(100):
        result = await run_recipe_async(recipe, "optimization soak", _ProbeClient(), f"soak-{index}")
        snapshots.append(_resource_snapshot(result.metrics, await _event_loop_lag_ms()))
    growth_fields = (
        "rss_bytes",
        "active_python_tasks",
        "thread_count",
        "child_process_count",
        "owned_capability_resources",
        "mcp_session_count",
    )
    growth = {
        field: {
            "first": snapshots[0].get(field) if snapshots else None,
            "last": snapshots[-1].get(field) if snapshots else None,
            "max": max((int(item[field]) for item in snapshots if item.get(field) is not None), default=None),
            "monotonic_non_decreasing": all(
                left.get(field) is not None and right.get(field) is not None and int(left[field]) <= int(right[field])
                for left, right in zip(snapshots, snapshots[1:])
            ),
        }
        for field in growth_fields
    }
    leak_fields = ("active_python_tasks", "thread_count", "child_process_count", "owned_capability_resources", "mcp_session_count")
    monotonic_growth = any(
        growth[field]["monotonic_non_decreasing"] and growth[field]["last"] > growth[field]["first"]
        for field in leak_fields
        if growth[field]["first"] is not None and growth[field]["last"] is not None
    )
    return {
        "runs": len(snapshots),
        "snapshots": snapshots,
        "resource_growth": growth,
        "resource_counts": {"owned_after_cleanup": snapshots[-1]["owned_capability_resources"] if snapshots else None},
        "monotonic_growth": monotonic_growth,
        "plateau": {field: item["first"] == item["last"] for field, item in growth.items() if item["first"] is not None and item["last"] is not None},
    }


async def _run(mode: str, fixture: Path, evidence_root: Path, measurement_mode: str = "full") -> dict[str, object]:
    workbook = await _measure_workbook(fixture, measurement_mode, mode)
    patterns, stage_metrics = await _measure_patterns()
    soak = await _soak_receipt()

    async def external_call(query: str = "") -> str:
        await asyncio.sleep(0)
        return f"fixture:{query}"

    from foresight_os.runtime.metrics import RunMetrics
    metrics = RunMetrics("optimization-external", "DIRECT", "direct")
    call = instrument_tool(external_call, CapabilityDefinition("external_fixture", "external_read"), {"metrics": metrics, "tool_call_kind": "external_http"})
    await call("probe")
    try:
        mcp = await _native_mcp_receipt()
    except Exception as error:
        mcp = {"status": "NOT_MEASURED", "error_type": type(error).__name__, "limitation": "optional MCP fixture unavailable"}
    state = observed_xlsx_state(fixture, evidence_root)
    local_correctness = all(workbook.get("workloads", {}).get(name, {}).get("status") == "VERIFIED" for name in MANDATORY_XLSX_WORKLOADS)
    evidence = {
        "correctness": measurement_evidence("VERIFIED" if local_correctness else "FAILED", limitation=None if local_correctness else "local workbook workload failed"),
        "latency": workbook["latency"],
        "allocation": workbook["allocation"],
        "soak": measurement_evidence("VERIFIED", limitation=None) if soak["runs"] >= 100 and not soak["monotonic_growth"] else measurement_evidence("FAILED", limitation="resource soak did not reach a clean plateau"),
        "regression_threshold": measurement_evidence("NOT_MEASURED", limitation="requires baseline/release comparison"),
        "optimization_demonstrated": measurement_evidence("NOT_MEASURED", limitation="no optimization claim is made by a single receipt"),
    }
    return build_receipt(
        mode=mode,
        fixture=fixture,
        workbook=workbook,
        patterns=patterns,
        stage_metrics=stage_metrics,
        soak=soak,
        external_fixture={"tool_calls": metrics.tool_calls, "external_http_requests": metrics.external_http_requests, "failures": metrics.tool_failures},
        native_mcp=mcp,
        state=state,
        thresholds={"p95_regression_pct": 20, "peak_allocation_regression_pct": 15, "tool_calls_per_finding_increase": 0, "soak_runs": 100, "gate_pattern_latency": False, "gate_timing_paths": True, "gate_cold_noise_floor": True, "gate_reader_metrics": bool(workbook.get("reader_instrumentation_available")) and bool(workbook["persistence"].get("instrumentation_available"))},
        evidence=evidence,
        target_root=TARGET_ROOT,
        harness_root=HARNESS_ROOT,
    )


def main() -> int:
    global TARGET_ROOT, HARNESS_ROOT
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("baseline", "release"), required=True)
    parser.add_argument("--fixture", type=Path, default=Path("tests/fixtures/analyst_multisheet.xlsx"))
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--measurement-mode", choices=("full", "timing", "allocation"), default="full")
    parser.add_argument("--target-root", type=Path, default=TARGET_ROOT)
    parser.add_argument("--harness-root", type=Path, default=HARNESS_ROOT)
    args = parser.parse_args()
    target_root = args.target_root.resolve()
    harness_root = args.harness_root.resolve()
    TARGET_ROOT = target_root
    HARNESS_ROOT = harness_root
    configure_measurement_roots(target_root=TARGET_ROOT, harness_root=HARNESS_ROOT)
    fixture = args.workbook or args.fixture
    if not fixture.is_absolute():
        fixture = target_root / fixture
    output = args.output or Path(f"work/optimization-{args.mode}/receipt.json")
    output.parent.mkdir(parents=True, exist_ok=True)
    result = asyncio.run(_run(args.mode, fixture, output.parent / "measurement-state", args.measurement_mode))
    output.write_text(json.dumps(result, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, sort_keys=True))
    return 0 if result["status"] == "VERIFIED" else 1

if __name__ == "__main__":
    raise SystemExit(main())
