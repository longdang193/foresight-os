from __future__ import annotations

import ast
import csv
import hashlib
import io
import json
from contextlib import contextmanager
from itertools import zip_longest
import operator
import posixpath
import re
from threading import RLock, get_ident
from types import MappingProxyType
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field, fields, replace
from datetime import date, datetime
from enum import Enum, StrEnum
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Sequence

import openpyxl

from foresight_os.domain.models import CalculationRecord, CalculationVerificationStatus, Evidence, EvidenceCandidate, MeasurementKey, SourceLocator
from foresight_os.domain.validation import validate_analysis, validate_evidence_candidate


class ToolbeltError(ValueError):
    pass


_MISSING = object()


@dataclass(frozen=True)
class SourceSnapshot:
    source: str
    sha256: str
    byte_count: int
    content: bytes = field(repr=False, compare=False)

    @classmethod
    def capture(
        cls,
        source: str | Path,
        *,
        cancel_check: Callable[[], bool] | None = None,
        chunk_size: int = 1024 * 1024,
        max_bytes: int | None = None,
    ) -> "SourceSnapshot":
        path = Path(source)
        if max_bytes is not None and max_bytes < 1:
            raise ToolbeltError("source byte limit must be positive")
        digest = hashlib.sha256()
        chunks: list[bytes] = []
        byte_count = 0
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(chunk_size), b""):
                _check_cancel(cancel_check)
                if max_bytes is not None and byte_count + len(chunk) > max_bytes:
                    raise ToolbeltError("workbook exceeds byte limit")
                digest.update(chunk)
                chunks.append(chunk)
                byte_count += len(chunk)
        content = b"".join(chunks)
        return cls(str(path), digest.hexdigest(), byte_count, content)

    def current_sha256(self, *, cancel_check: Callable[[], bool] | None = None) -> str:
        return _sha256_file(Path(self.source), cancel_check=cancel_check)

    def is_current(self, *, cancel_check: Callable[[], bool] | None = None) -> bool:
        return self.current_sha256(cancel_check=cancel_check) == self.sha256


class FormulaStatus(StrEnum):
    VALUE = "VALUE"
    FORMULA_WITH_CACHE = "FORMULA_WITH_CACHE"
    FORMULA_WITHOUT_CACHE = "FORMULA_WITHOUT_CACHE"
    FORMULA_RECALCULATED = "FORMULA_RECALCULATED"
    FORMULA_STALE_OR_UNKNOWN = "FORMULA_STALE_OR_UNKNOWN"
    ERROR = "ERROR"


@dataclass(frozen=True, slots=True)
class ExcelError:
    value: str

    def __repr__(self) -> str:
        return f"ExcelError({self.value!r})"

@dataclass(frozen=True)
class ToolbeltLimits:
    max_bytes: int = 25_000_000
    max_expanded_bytes: int = 100_000_000
    max_rows: int = 50_000
    max_cells: int = 250_000
    max_merged_ranges: int = 10_000
    max_merged_cells_per_range: int = 100_000
    max_total_merged_cells: int = 500_000
    max_result_rows: int = 10_000


def _coerce_toolbelt_limits(value: ToolbeltLimits | Mapping[str, Any] | None) -> ToolbeltLimits:
    if value is None or isinstance(value, ToolbeltLimits):
        return value or ToolbeltLimits()
    if not isinstance(value, Mapping):
        raise ToolbeltError("toolbelt limits must be structured")
    fields = (
        "max_bytes",
        "max_expanded_bytes",
        "max_rows",
        "max_cells",
        "max_merged_ranges",
        "max_merged_cells_per_range",
        "max_total_merged_cells",
        "max_result_rows",
    )
    return ToolbeltLimits(**{name: int(value[name]) for name in fields if name in value})


class WorkbookReadCache:
    def __init__(self, max_entries: int = 8, max_bytes: int = 50_000_000) -> None:
        self.max_entries = max(1, max_entries)
        self.max_bytes = max(1, max_bytes)
        self._entries: dict[tuple[Any, ...], tuple[WorkbookReadResult, int]] = {}
        self._analysis_entries: dict[tuple[Any, ...], tuple[AnalysisResult, int]] = {}
        self._bytes = 0
        self._read_bytes = 0
        self._analysis_bytes = 0
        self._transaction: dict[str, Any] | None = None

    def begin_transaction(self) -> dict[str, Any]:
        if self._transaction is not None:
            raise RuntimeError("workbook cache transaction already active")
        self._transaction = {
            "entries": {},
            "analysis_entries": {},
            "bytes": self._bytes,
            "read_bytes": self._read_bytes,
            "analysis_bytes": self._analysis_bytes,
        }
        return self._transaction

    def commit_transaction(self, transaction: dict[str, Any]) -> None:
        if self._transaction is not transaction:
            raise RuntimeError("workbook cache transaction is not active")
        self._transaction = None

    def rollback_transaction(self, transaction: dict[str, Any]) -> None:
        if self._transaction is not transaction:
            raise RuntimeError("workbook cache transaction is not active")
        for name, entries in (("entries", self._entries), ("analysis_entries", self._analysis_entries)):
            for key, previous in transaction[name].items():
                if previous is _MISSING:
                    entries.pop(key, None)
                else:
                    entries[key] = previous
        self._bytes = transaction["bytes"]
        self._read_bytes = transaction["read_bytes"]
        self._analysis_bytes = transaction["analysis_bytes"]
        self._transaction = None

    def _record_transaction(self, name: str, entries: dict[tuple[Any, ...], tuple[Any, int]], key: tuple[Any, ...]) -> None:
        if self._transaction is not None and key not in self._transaction[name]:
            self._transaction[name][key] = entries.get(key, _MISSING)

    def source_hash(self, source: Path) -> str:
        return _sha256_file(source)

    def key(
        self,
        source: Path,
        selection: SheetSelection | None,
        limits: ToolbeltLimits,
        *,
        snapshot: SourceSnapshot | None = None,
        source_hash: str | None = None,
    ) -> tuple[Any, ...]:
        canonical = str(source.resolve())
        limit_key = repr(limits)
        digest = snapshot.sha256 if snapshot is not None else source_hash or self.source_hash(source)
        view_identity = _workbook_view_identity(digest, selection or SheetSelection())
        return canonical, view_identity.material_fingerprint, limit_key

    def get(self, key: tuple[Any, ...]) -> WorkbookReadResult | None:
        entry = self._entries.get(key)
        return entry[0] if entry is not None else None

    def put(self, key: tuple[Any, ...], result: WorkbookReadResult) -> None:
        self._put(self._entries, key, result)

    def analysis_key(self, table: TypedTable, operation: AnalysisOperation, limits: ToolbeltLimits) -> tuple[Any, ...]:
        return ("analysis", table.material_fingerprint, _analysis_operation_identity(operation), repr(limits))

    def get_analysis(self, key: tuple[Any, ...]) -> AnalysisResult | None:
        entry = self._analysis_entries.get(key)
        return entry[0] if entry is not None else None

    def put_analysis(self, key: tuple[Any, ...], result: AnalysisResult) -> None:
        self._put(self._analysis_entries, key, result)

    def _put(self, entries: dict[tuple[Any, ...], tuple[Any, int]], key: tuple[Any, ...], result: Any) -> None:
        estimated = max(1, int(getattr(result, "retained_bytes", 0) or _retained_size(result)))
        if estimated > self.max_bytes:
            return
        name = "entries" if entries is self._entries else "analysis_entries"
        self._record_transaction(name, entries, key)
        previous = entries.pop(key, None)
        if previous is not None:
            self._bytes -= previous[1]
            if name == "entries":
                self._read_bytes -= previous[1]
            else:
                self._analysis_bytes -= previous[1]
        while self._entry_count >= self.max_entries or self._bytes + estimated > self.max_bytes:
            oldest = self._entries if self._entries else self._analysis_entries
            oldest_key = next(iter(oldest))
            oldest_name = "entries" if oldest is self._entries else "analysis_entries"
            self._record_transaction(oldest_name, oldest, oldest_key)
            _, size = oldest.pop(oldest_key)
            self._bytes -= size
            if oldest_name == "entries":
                self._read_bytes -= size
            else:
                self._analysis_bytes -= size
        entries[key] = (result, estimated)
        self._bytes += estimated
        if name == "entries":
            self._read_bytes += estimated
        else:
            self._analysis_bytes += estimated

    def clear(self) -> None:
        self._entries.clear()
        self._analysis_entries.clear()
        self._bytes = 0
        self._read_bytes = 0
        self._analysis_bytes = 0

    @property
    def size(self) -> int:
        return self._entry_count

    @property
    def bytes_used(self) -> int:
        return self._bytes

    @property
    def _entry_count(self) -> int:
        return len(self._entries) + len(self._analysis_entries)

    @property
    def read_bytes(self) -> int:
        return self._read_bytes

    @property
    def analysis_bytes(self) -> int:
        return self._analysis_bytes


@dataclass
class WorkbookStats:
    capture_count: int = 0
    capture_hash_count: int = 0
    currentness_hash_count: int = 0
    xlsx_parse_count: int = 0
    openpyxl_load_count: int = 0
    cache_hits: int = 0
    analysis_cache_hits: int = 0
    snapshot_bytes: int = 0
    retained_read_result_bytes: int = 0
    retained_analysis_result_bytes: int = 0
    retained_authority_record_bytes: int = 0


class AuthorityIndex:
    def __init__(self, max_records: int = ToolbeltLimits().max_cells) -> None:
        self.max_records = max(1, max_records)
        self.measurements: dict[tuple[str, str, str], TypedMeasurement] = {}
        self.calculations: dict[tuple[str, str], CalculationRecord] = {}
        self._measurement_bytes: dict[tuple[str, str, str], int] = {}
        self._calculation_bytes: dict[tuple[str, str], int] = {}
        self._bytes = 0
        self._generation = 0
        self._transaction: dict[str, Any] | None = None

    def begin_transaction(self) -> dict[str, Any]:
        if self._transaction is not None:
            raise RuntimeError("authority transaction already active")
        self._transaction = {
            "measurements": {},
            "calculations": {},
            "bytes": self._bytes,
            "generation": self._generation,
        }
        return self._transaction

    def commit_transaction(self, transaction: dict[str, Any]) -> None:
        if self._transaction is not transaction:
            raise RuntimeError("authority transaction is not active")
        self._transaction = None

    def rollback_transaction(self, transaction: dict[str, Any]) -> None:
        if self._transaction is not transaction:
            raise RuntimeError("authority transaction is not active")
        for name, entries in (("measurements", self.measurements), ("calculations", self.calculations)):
            sizes = self._measurement_bytes if name == "measurements" else self._calculation_bytes
            for key, previous in transaction[name].items():
                if previous is _MISSING:
                    entries.pop(key, None)
                    sizes.pop(key, None)
                else:
                    entries[key] = previous[0]
                    sizes[key] = previous[1]
        self._bytes = transaction["bytes"]
        self._generation = transaction["generation"]
        self._transaction = None

    def _record_transaction(self, name: str, key: tuple[Any, ...], value: Any, size: int) -> None:
        if self._transaction is not None and key not in self._transaction[name]:
            self._transaction[name][key] = _MISSING if value is None else (value, size)

    def remember_measurements(self, measurements: Iterable[TypedMeasurement]) -> None:
        changed = False
        for measurement in measurements:
            key = (measurement.provenance.workbook_sha256, measurement.id, measurement.provenance.range)
            self._record_transaction("measurements", key, self.measurements.get(key), self._measurement_bytes.get(key, 0))
            if self.measurements.get(key) != measurement:
                changed = True
            self.measurements[key] = measurement
            size = _retained_size(measurement)
            self._bytes += size - self._measurement_bytes.get(key, 0)
            self._measurement_bytes[key] = size
        while len(self.measurements) > self.max_records:
            changed = True
            key = next(iter(self.measurements))
            self._record_transaction("measurements", key, self.measurements.get(key), self._measurement_bytes.get(key, 0))
            self.measurements.pop(key)
            self._bytes -= self._measurement_bytes.pop(key, 0)
        if changed:
            self._generation += 1

    def remember_calculation(self, record: CalculationRecord, workbook_hash: str) -> None:
        key = (workbook_hash, record.id)
        self._record_transaction("calculations", key, self.calculations.get(key), self._calculation_bytes.get(key, 0))
        changed = self.calculations.get(key) != record
        self.calculations[key] = record
        size = _retained_size(record)
        self._bytes += size - self._calculation_bytes.get(key, 0)
        self._calculation_bytes[key] = size
        while len(self.calculations) > self.max_records:
            changed = True
            key = next(iter(self.calculations))
            self._record_transaction("calculations", key, self.calculations.get(key), self._calculation_bytes.get(key, 0))
            self.calculations.pop(key)
            self._bytes -= self._calculation_bytes.pop(key, 0)
        if changed:
            self._generation += 1

    def clear(self) -> None:
        changed = bool(self.measurements or self.calculations)
        self.measurements.clear()
        self.calculations.clear()
        self._measurement_bytes.clear()
        self._calculation_bytes.clear()
        self._bytes = 0
        if changed:
            self._generation += 1

    @property
    def size(self) -> int:
        return len(self.measurements) + len(self.calculations)

    @property
    def estimated_bytes(self) -> int:
        return self._bytes

    @property
    def generation(self) -> int:
        return self._generation


def _sha256_file(
    source: Path,
    chunk_size: int = 1024 * 1024,
    *,
    cancel_check: Callable[[], bool] | None = None,
) -> str:
    digest = hashlib.sha256()
    with source.open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_size), b""):
            _check_cancel(cancel_check)
            digest.update(chunk)
    return digest.hexdigest()


@dataclass(frozen=True)
class SheetSelection:
    sheet_names: tuple[str, ...] = ()
    ranges: tuple[tuple[str, str], ...] = ()
    tables: tuple[tuple[str, str], ...] = ()
    include_hidden: bool = False
    include_hidden_rows: bool = False


@dataclass(frozen=True)
class WorkbookViewIdentity:
    source_sha256: str
    canonical_selection: tuple[Any, ...]
    hidden_sheet_policy: str
    hidden_row_policy: str
    selected_ranges: tuple[tuple[str, str], ...]
    selected_tables: tuple[tuple[str, str], ...]
    parser_schema: str = "xlsx-openpyxl-formula-cache-v1"

    @property
    def material_fingerprint(self) -> str:
        payload = {
            "source_sha256": self.source_sha256,
            "canonical_selection": self.canonical_selection,
            "hidden_sheet_policy": self.hidden_sheet_policy,
            "hidden_row_policy": self.hidden_row_policy,
            "selected_ranges": self.selected_ranges,
            "selected_tables": self.selected_tables,
            "parser_schema": self.parser_schema,
        }
        return hashlib.sha256(json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")).hexdigest()


def _workbook_view_identity(source_sha256: str, selection: SheetSelection) -> WorkbookViewIdentity:
    return WorkbookViewIdentity(
        source_sha256=source_sha256,
        canonical_selection=(selection.sheet_names, selection.ranges, selection.tables),
        hidden_sheet_policy="include" if selection.include_hidden else "visible-only",
        hidden_row_policy="include" if selection.include_hidden_rows else "visible-only",
        selected_ranges=selection.ranges,
        selected_tables=selection.tables,
    )


def _coerce_sheet_selection(value: SheetSelection | Mapping[str, Any] | None) -> SheetSelection | None:
    if value is None or isinstance(value, SheetSelection):
        return value
    if isinstance(value, str):
        return SheetSelection(sheet_names=(value,))
    if not isinstance(value, Mapping):
        raise ToolbeltError("spreadsheet selection must be structured")

    def pairs(name: str, singular_name: str) -> tuple[tuple[str, str], ...]:
        result = []
        entries = value.get(name, ()) or ()
        if not entries and value.get(singular_name) is not None:
            entries = (value[singular_name],)
        if isinstance(entries, (str, bytes)) or not isinstance(entries, Sequence):
            raise ToolbeltError(f"spreadsheet selection {name} must be a list of pairs")
        for item in entries:
            if isinstance(item, (str, bytes)) or not isinstance(item, Sequence) or len(item) != 2:
                raise ToolbeltError(f"spreadsheet selection {name} entries must contain two values")
            result.append((str(item[0]), str(item[1])))
        return tuple(result)

    sheet_names = value.get("sheet_names", ()) or ()
    if not sheet_names and value.get("sheet") is not None:
        sheet_names = (value["sheet"],)
    if isinstance(sheet_names, str):
        sheet_names = (sheet_names,)
    elif isinstance(sheet_names, (bytes, bool)) or not isinstance(sheet_names, Sequence):
        raise ToolbeltError("spreadsheet selection sheet_names must be a list of sheet names")
    return SheetSelection(
        sheet_names=tuple(str(item) for item in sheet_names),
        ranges=pairs("ranges", "range"),
        tables=pairs("tables", "table"),
        include_hidden=bool(value.get("include_hidden", False)),
        include_hidden_rows=bool(value.get("include_hidden_rows", False)),
    )


@dataclass(frozen=True)
class CellProvenance:
    workbook_sha256: str
    sheet: str
    coordinate: str
    range: str


@dataclass(frozen=True)
class TypedMeasurement:
    id: str
    value: Any
    data_type: str
    formula: str | None
    cached_value: Any
    formula_status: FormulaStatus
    provenance: CellProvenance
    cross_sheet_references: tuple[str, ...] = ()
    cross_sheet_cells: tuple[tuple[str, str], ...] = ()
    trusted: bool = False

    def to_dict(self) -> dict[str, Any]:
        return {
            "id": self.id,
            "value": self.value.value if isinstance(self.value, ExcelError) else self.value,
            "data_type": self.data_type,
            "formula": self.formula,
            "cached_value": self.cached_value.value if isinstance(self.cached_value, ExcelError) else self.cached_value,
            "formula_status": self.formula_status.value,
            "provenance": {
                "workbook_sha256": self.provenance.workbook_sha256,
                "sheet": self.provenance.sheet,
                "coordinate": self.provenance.coordinate,
                "range": self.provenance.range,
            },
            "cross_sheet_references": list(self.cross_sheet_references),
            "cross_sheet_cells": [list(item) for item in self.cross_sheet_cells],
            "trusted": self.trusted,
        }


def _stable_material_value(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(key): _stable_material_value(item) for key, item in sorted(value.items(), key=lambda item: str(item[0]))}
    if isinstance(value, (list, tuple)):
        return [_stable_material_value(item) for item in value]
    if isinstance(value, (set, frozenset)):
        items = [_stable_material_value(item) for item in value]
        return sorted(items, key=lambda item: json.dumps(item, sort_keys=True, separators=(",", ":"), ensure_ascii=True))
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, bytes):
        return value.hex()
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Enum):
        return _stable_material_value(value.value)
    if hasattr(value, "to_dict"):
        return _stable_material_value(value.to_dict())
    if hasattr(value, "__dataclass_fields__"):
        return {
            field_info.name: _stable_material_value(getattr(value, field_info.name))
            for field_info in fields(value)
            if field_info.name != "retained_bytes"
        }
    return repr(value)


def _retained_size(value: Any) -> int:
    encoded = json.dumps(
        _stable_material_value(value),
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=True,
    ).encode("utf-8")
    return max(1, len(encoded))


TYPED_TABLE_IDENTITY_VERSION = "typed-table-material-v2"


@dataclass(frozen=True)
class TypedTable:
    id: str
    name: str
    sheet: str
    headers: tuple[tuple[str, ...], ...]
    rows: tuple[tuple[Any, ...], ...]
    row_provenance: tuple[tuple[CellProvenance, ...], ...]
    hidden: bool = False
    source_range: str = ""
    related_provenance: tuple[CellProvenance, ...] = ()
    row_measurements: tuple[tuple[TypedMeasurement, ...], ...] = ()
    material_fingerprint: str = field(default="", init=False)

    def __post_init__(self) -> None:
        payload = {
            "identity_version": TYPED_TABLE_IDENTITY_VERSION,
            "name": self.name,
            "sheet": self.sheet,
            "headers": self.headers,
            "rows": self.rows,
            "row_provenance": self.row_provenance,
            "hidden": self.hidden,
            "source_range": self.source_range,
            "related_provenance": self.related_provenance,
            "row_measurements": self.row_measurements,
        }
        encoded = json.dumps(_stable_material_value(payload), sort_keys=True, separators=(",", ":"), ensure_ascii=True)
        object.__setattr__(self, "material_fingerprint", hashlib.sha256(encoded.encode("utf-8")).hexdigest())

    @property
    def columns(self) -> tuple[str, ...]:
        if not self.headers:
            return ()
        if len(self.headers) == 1:
            return tuple(str(item).strip() for item in self.headers[0])
        return tuple(" ".join(str(part) for part in parts if str(part).strip()).strip() for parts in zip(*self.headers, strict=False))

    def to_dict(self) -> dict[str, Any]:
        compact_rows = []
        compact_row_provenance = []
        compact_row_measurements = []
        for index, row in enumerate(self.rows):
            if not any(value not in (None, "") for value in row):
                continue
            provenance = self.row_provenance[index] if index < len(self.row_provenance) else ()
            measurements = self.row_measurements[index] if index < len(self.row_measurements) else ()
            compact_rows.append(list(row))
            compact_row_provenance.append(
                [
                    {
                        "workbook_sha256": item.workbook_sha256,
                        "sheet": item.sheet,
                        "coordinate": item.coordinate,
                        "range": item.range,
                    }
                    for value, item in zip(row, provenance, strict=False)
                    if value not in (None, "")
                ]
            )
            compact_row_measurements.append([item.to_dict() for item in measurements])
        return {
            "id": self.id,
            "name": self.name,
            "sheet": self.sheet,
            "headers": [list(row) for row in self.headers],
            "rows": compact_rows,
            "row_provenance": compact_row_provenance,
            "row_measurements": compact_row_measurements,
            "row_count": len(self.rows),
            "omitted_blank_row_count": len(self.rows) - len(compact_rows),
            "hidden": self.hidden,
            "source_range": self.source_range,
            "related_provenance": [
                {
                    "workbook_sha256": item.workbook_sha256,
                    "sheet": item.sheet,
                    "coordinate": item.coordinate,
                    "range": item.range,
                }
                for item in self.related_provenance
            ],
            "material_fingerprint": self.material_fingerprint,
        }


_AUTHORITATIVE_MEASUREMENTS: dict[tuple[str, str, str], TypedMeasurement] = {}
_AUTHORITATIVE_CALCULATIONS: dict[tuple[str, str], CalculationRecord] = {}
_AUTHORITATIVE_RECORD_LIMIT = ToolbeltLimits().max_cells


def _remember_authoritative_measurements(
    measurements: Iterable[TypedMeasurement],
    authority_index: AuthorityIndex | None = None,
) -> None:
    if authority_index is not None:
        authority_index.remember_measurements(measurements)
        return
    for measurement in measurements:
        key = (measurement.provenance.workbook_sha256, measurement.id, measurement.provenance.range)
        _AUTHORITATIVE_MEASUREMENTS[key] = measurement
    while len(_AUTHORITATIVE_MEASUREMENTS) > _AUTHORITATIVE_RECORD_LIMIT:
        _AUTHORITATIVE_MEASUREMENTS.pop(next(iter(_AUTHORITATIVE_MEASUREMENTS)))


def _remember_authoritative_calculation(
    record: CalculationRecord,
    workbook_hash: str,
    authority_index: AuthorityIndex | None = None,
) -> None:
    if authority_index is not None:
        authority_index.remember_calculation(record, workbook_hash)
        return
    _AUTHORITATIVE_CALCULATIONS[(workbook_hash, record.id)] = record
    while len(_AUTHORITATIVE_CALCULATIONS) > _AUTHORITATIVE_RECORD_LIMIT:
        _AUTHORITATIVE_CALCULATIONS.pop(next(iter(_AUTHORITATIVE_CALCULATIONS)))


def _publish_read_result(
    result: WorkbookReadResult,
    *,
    cache: Any = None,
    cache_key: tuple[Any, ...] | None = None,
    authority_index: AuthorityIndex | None = None,
) -> WorkbookReadResult:
    cache_transaction = None
    authority_transaction = None
    if cache is not None and hasattr(cache, "begin_transaction") and getattr(cache, "_transaction", None) is None:
        cache_transaction = cache.begin_transaction()
    if authority_index is not None and getattr(authority_index, "_transaction", None) is None:
        authority_transaction = authority_index.begin_transaction()
    try:
        _remember_authoritative_measurements(result.measurements, authority_index)
        if cache is not None and cache_key is not None:
            cache.put(cache_key, result)
        if cache_transaction is not None:
            cache.commit_transaction(cache_transaction)
        if authority_transaction is not None:
            authority_index.commit_transaction(authority_transaction)
        return result
    except BaseException:
        if cache_transaction is not None:
            cache.rollback_transaction(cache_transaction)
        if authority_transaction is not None:
            authority_index.rollback_transaction(authority_transaction)
        raise


def _with_read_source(result: WorkbookReadResult, source: str | Path) -> WorkbookReadResult:
    source_text = str(source)
    return result if result.source == source_text else replace(result, source=source_text)


@dataclass(frozen=True)
class SheetSummary:
    name: str
    state: str
    max_row: int
    max_column: int
    hidden_rows: tuple[int, ...]
    merged_ranges: tuple[str, ...]
    tables: tuple[str, ...]

    def to_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "state": self.state,
            "max_row": self.max_row,
            "max_column": self.max_column,
            "hidden_rows": list(self.hidden_rows),
            "merged_ranges": list(self.merged_ranges),
            "tables": list(self.tables),
        }


@dataclass(frozen=True)
class WorkbookReadResult:
    source: str
    workbook_sha256: str
    sheets: tuple[SheetSummary, ...]
    tables: tuple[TypedTable, ...]
    measurements: tuple[TypedMeasurement, ...]
    warnings: tuple[str, ...] = ()
    retained_bytes: int = field(default=0, init=False, repr=False, compare=False)

    def __post_init__(self) -> None:
        object.__setattr__(self, "warnings", tuple(self.warnings))
        object.__setattr__(self, "retained_bytes", _retained_size(self))

    def to_dict(self) -> dict[str, Any]:
        serialized_measurements = [
            item.to_dict()
            for item in self.measurements
            if item.value not in (None, "") or item.formula is not None
        ]
        return {
            "source": self.source,
            "workbook_sha256": self.workbook_sha256,
            "sheets": [item.to_dict() for item in self.sheets],
            "tables": [item.to_dict() for item in self.tables],
            "measurements": serialized_measurements,
            "measurement_count": len(self.measurements),
            "omitted_blank_measurement_count": len(self.measurements) - len(serialized_measurements),
            "warnings": list(self.warnings),
        }


class WorkbookContext:
    """Run-owned source snapshots, derived results, authority, and statistics."""

    def __init__(self, *, max_entries: int = 8, max_bytes: int = 50_000_000, max_total_bytes: int = 100_000_000) -> None:
        self.snapshots: dict[str, SourceSnapshot] = {}
        self.cache = WorkbookReadCache(max_entries=max_entries, max_bytes=max_bytes)
        self.authority_index = AuthorityIndex(max_records=max(1, max_entries * ToolbeltLimits().max_cells))
        self.stats = WorkbookStats()
        self.max_total_bytes = max(1, max_total_bytes)
        self._closed = False
        self._lock = RLock()
        self._transaction: dict[str, Any] | None = None
        self._transaction_owner: int | None = None
        self._authority_validated: dict[tuple[Any, ...], int] = {}

    def ensure_open(self) -> None:
        if self._closed:
            raise ToolbeltError("workbook context is closed")

    @contextmanager
    def transaction(self):
        self.ensure_open()
        owner = get_ident()
        if self._transaction is not None and self._transaction_owner == owner:
            yield self
            return
        with self._lock:
            transaction = {
                "cache": self.cache.begin_transaction(),
                "authority": self.authority_index.begin_transaction(),
                "snapshots": dict(self.snapshots),
                "authority_validated": dict(self._authority_validated),
                "stats": {
                    name: getattr(self.stats, name)
                    for name in (
                        "snapshot_bytes",
                        "retained_read_result_bytes",
                        "retained_analysis_result_bytes",
                        "retained_authority_record_bytes",
                    )
                },
            }
            self._transaction = transaction
            self._transaction_owner = owner
            try:
                yield self
                self._enforce_budget()
            except BaseException:
                self.cache.rollback_transaction(transaction["cache"])
                self.authority_index.rollback_transaction(transaction["authority"])
                self.snapshots.clear()
                self.snapshots.update(transaction["snapshots"])
                self._authority_validated.clear()
                self._authority_validated.update(transaction["authority_validated"])
                for name, value in transaction["stats"].items():
                    setattr(self.stats, name, value)
                raise
            else:
                self.cache.commit_transaction(transaction["cache"])
                self.authority_index.commit_transaction(transaction["authority"])
            finally:
                self._transaction = None
                self._transaction_owner = None

    def _run_transaction(self, operation: Callable[[], Any]) -> Any:
        if self._transaction is not None and self._transaction_owner == get_ident():
            return operation()
        with self.transaction():
            return operation()

    def capture(
        self,
        source: str | Path,
        *,
        cancel_check: Callable[[], bool] | None = None,
        max_bytes: int | None = None,
    ) -> SourceSnapshot:
        self.ensure_open()
        _check_cancel(cancel_check)
        return self._run_transaction(lambda: self._capture_in_transaction(source, cancel_check=cancel_check, max_bytes=max_bytes))

    def _capture_in_transaction(
        self,
        source: str | Path,
        *,
        cancel_check: Callable[[], bool] | None = None,
        max_bytes: int | None = None,
    ) -> SourceSnapshot:
        key = str(Path(source).resolve())
        snapshot = self.snapshots.get(key)
        if snapshot is not None:
            return snapshot
        snapshot = SourceSnapshot.capture(source, cancel_check=cancel_check, max_bytes=max_bytes)
        self.snapshots[key] = snapshot
        self.stats.capture_count += 1
        self.stats.capture_hash_count += 1
        self.stats.snapshot_bytes += snapshot.byte_count
        return snapshot

    def currentness(self, source: str | Path, *, cancel_check: Callable[[], bool] | None = None) -> bool:
        self.ensure_open()
        _check_cancel(cancel_check)
        return self._run_transaction(lambda: self._currentness_in_transaction(source, cancel_check=cancel_check))

    def _currentness_in_transaction(self, source: str | Path, *, cancel_check: Callable[[], bool] | None = None) -> bool:
        snapshot = self._capture_in_transaction(source, cancel_check=cancel_check)
        self.stats.currentness_hash_count += 1
        return snapshot.is_current(cancel_check=cancel_check)

    def read(
        self,
        source: str | Path,
        selection: Mapping[str, Any] | None = None,
        limits: Mapping[str, Any] | None = None,
        *,
        cancel_check: Callable[[], bool] | None = None,
    ) -> WorkbookReadResult:
        self.ensure_open()
        _check_cancel(cancel_check)
        selection = _coerce_sheet_selection(selection)
        read_limits = _coerce_toolbelt_limits(limits)
        source = Path(source)
        with self._lock:
            cached_result = self._try_read_cache_hit(source, selection, read_limits, cancel_check=cancel_check)
            if cached_result is not None:
                return cached_result
            return self._run_transaction(
                lambda: self._read_in_transaction(source, selection, read_limits, cancel_check=cancel_check)
            )

    def _try_read_cache_hit(
        self,
        source: Path,
        selection: SheetSelection,
        read_limits: ToolbeltLimits,
        *,
        cancel_check: Callable[[], bool] | None = None,
    ) -> WorkbookReadResult | None:
        source_key = str(source.resolve())
        cached_snapshot = self.snapshots.get(source_key)
        if cached_snapshot is None:
            return None
        cached_key = self.cache.key(source, selection, read_limits, snapshot=cached_snapshot)
        cached_result = self.cache.get(cached_key)
        if cached_result is None:
            return None
        _check_cancel(cancel_check)
        if (
            self._authority_validated.get(cached_key) == self.authority_index.generation
            and len(self.authority_index.measurements) >= len(cached_result.measurements)
        ):
            self.stats.cache_hits += 1
            return _with_read_source(cached_result, source)
        authority_keys = {
            (item.provenance.workbook_sha256, item.id, item.provenance.range)
            for item in cached_result.measurements
        }
        if all(key in self.authority_index.measurements for key in authority_keys):
            self._authority_validated[cached_key] = self.authority_index.generation
            self.stats.cache_hits += 1
            return _with_read_source(cached_result, source)
        return None

    def _read_in_transaction(
        self,
        source: Path,
        selection: SheetSelection,
        read_limits: ToolbeltLimits,
        *,
        cancel_check: Callable[[], bool] | None = None,
    ) -> WorkbookReadResult:
        source_key = str(Path(source).resolve())
        cached_key: tuple[Any, ...] | None = None
        cached_snapshot = self.snapshots.get(source_key)
        if cached_snapshot is not None:
            cached_key = self.cache.key(source, selection, read_limits, snapshot=cached_snapshot)
            cached_result = self.cache.get(cached_key)
            if cached_result is not None:
                _check_cancel(cancel_check)
                if (
                    self._authority_validated.get(cached_key) == self.authority_index.generation
                    and len(self.authority_index.measurements) >= len(cached_result.measurements)
                ):
                    self.stats.cache_hits += 1
                    return _with_read_source(cached_result, source)

                authority_keys = {
                    (item.provenance.workbook_sha256, item.id, item.provenance.range)
                    for item in cached_result.measurements
                }
                if all(key in self.authority_index.measurements for key in authority_keys):
                    self._authority_validated[cached_key] = self.authority_index.generation
                    self.stats.cache_hits += 1
                    return _with_read_source(cached_result, source)

                def restore_authority() -> WorkbookReadResult:
                    _remember_authoritative_measurements(cached_result.measurements, self.authority_index)
                    self._authority_validated[cached_key] = self.authority_index.generation
                    self.stats.cache_hits += 1
                    return _with_read_source(cached_result, source)

                return restore_authority()

        def read_uncached() -> WorkbookReadResult:
            snapshot = self._capture_in_transaction(source, cancel_check=cancel_check, max_bytes=read_limits.max_bytes)
            return spreadsheet_read(
                str(source),
                selection,
                read_limits,
                cancel_check=cancel_check,
                _cache=self.cache,
                _snapshot=snapshot,
                _authority_index=self.authority_index,
                _stats=self.stats,
            )

        result = read_uncached()
        cached_key = self.cache.key(source, selection, read_limits, snapshot=self.snapshots[source_key])
        if self.cache.get(cached_key) is not None:
            self._authority_validated[cached_key] = self.authority_index.generation
        self._prune_authority_validation()
        return result

    def analyze(
        self,
        table: TypedTable | Mapping[str, Any],
        operation: AnalysisOperation | Mapping[str, Any],
        limits: Mapping[str, Any] | None = None,
        *,
        cancel_check: Callable[[], bool] | None = None,
    ) -> AnalysisResult:
        self.ensure_open()
        _check_cancel(cancel_check)
        normalized_table = _coerce_typed_table(table)
        normalized_operation = _coerce_analysis_operation(operation)
        normalized_limits = _coerce_toolbelt_limits(limits)
        with self._lock:
            _check_cancel(cancel_check)
            cache_key = self.cache.analysis_key(normalized_table, normalized_operation, normalized_limits)
            cached_result = self.cache.get_analysis(cache_key)
            if cached_result is not None:
                self.stats.analysis_cache_hits += 1
                return cached_result
            return self._run_transaction(
                lambda: self._analyze_in_transaction(
                    normalized_table,
                    normalized_operation,
                    normalized_limits,
                    cancel_check=cancel_check,
                )
            )

    def _analyze_in_transaction(
        self,
        normalized_table: TypedTable,
        normalized_operation: AnalysisOperation,
        normalized_limits: ToolbeltLimits,
        *,
        cancel_check: Callable[[], bool] | None = None,
    ) -> AnalysisResult:
        cache_key = self.cache.analysis_key(normalized_table, normalized_operation, normalized_limits)
        cached_result = self.cache.get_analysis(cache_key)
        if cached_result is not None:
            self.stats.analysis_cache_hits += 1
            return cached_result

        def analyze_uncached() -> AnalysisResult:
            return data_analysis(
                normalized_table,
                normalized_operation,
                normalized_limits,
                cancel_check=cancel_check,
                _cache=self.cache,
                _authority_index=self.authority_index,
                _stats=self.stats,
            )

        return analyze_uncached()

    def _refresh_retained_stats(self) -> None:
        self.stats.retained_read_result_bytes = self.cache.read_bytes
        self.stats.retained_analysis_result_bytes = self.cache.analysis_bytes
        self.stats.retained_authority_record_bytes = self.authority_index.estimated_bytes

    def _prune_authority_validation(self) -> None:
        self._authority_validated = {
            key: generation
            for key, generation in self._authority_validated.items()
            if self.cache.get(key) is not None
        }

    def _enforce_budget(self) -> None:
        self._refresh_retained_stats()
        retained = (
            self.stats.snapshot_bytes
            + self.stats.retained_read_result_bytes
            + self.stats.retained_analysis_result_bytes
            + self.stats.retained_authority_record_bytes
        )
        if retained > self.max_total_bytes:
            raise ToolbeltError("workbook context retained-state budget exceeded")

    def enforce_budget(self) -> None:
        self.ensure_open()
        self._enforce_budget()

    def clear(self) -> None:
        if self._closed:
            return
        self._closed = True
        self.snapshots.clear()
        self.cache.clear()
        self.authority_index.clear()
        self._authority_validated.clear()
        self.stats.snapshot_bytes = 0
        self.stats.retained_read_result_bytes = 0
        self.stats.retained_analysis_result_bytes = 0
        self.stats.retained_authority_record_bytes = 0


@dataclass(frozen=True)
class AnalysisOperation:
    kind: str
    columns: tuple[str, ...] = ()
    key_columns: tuple[str, ...] = ()
    value_column: str | None = None
    scenarios: tuple[tuple[str, float], ...] = ()
    sensitivity: tuple[tuple[str, tuple[float, ...]], ...] = ()
    many_to_many: bool = False
    filters: tuple[tuple[str, str, Any], ...] = ()
    join_table: TypedTable | None = None
    join_keys: tuple[str, str] = ()


def _analysis_operation_identity(operation: AnalysisOperation) -> tuple[Any, ...]:
    return (
        operation.kind,
        operation.columns,
        operation.key_columns,
        operation.value_column,
        operation.scenarios,
        operation.sensitivity,
        operation.many_to_many,
        operation.filters,
        operation.join_table.material_fingerprint if operation.join_table is not None else None,
        operation.join_keys,
    )


@dataclass(frozen=True)
class AnalysisResult:
    operation: str
    rows: tuple[Mapping[str, Any], ...]
    provenance: tuple[CellProvenance, ...]
    deterministic_key: str
    warnings: tuple[str, ...] = ()
    retained_bytes: int = field(default=0, init=False, repr=False, compare=False)

    def __post_init__(self) -> None:
        object.__setattr__(
            self,
            "rows",
            tuple(
                MappingProxyType({str(key): _freeze_result_value(value) for key, value in row.items()})
                for row in self.rows
            ),
        )
        object.__setattr__(self, "warnings", tuple(self.warnings))
        object.__setattr__(self, "retained_bytes", _retained_size(self))

    def to_dict(self) -> dict[str, Any]:
        return {
            "operation": self.operation,
            "rows": [_thaw_result_value(row) for row in self.rows],
            "provenance": [_thaw_result_value(item) for item in self.provenance],
            "deterministic_key": self.deterministic_key,
            "warnings": list(self.warnings),
        }


def _freeze_result_value(value: Any) -> Any:
    if isinstance(value, Mapping):
        return MappingProxyType({str(key): _freeze_result_value(item) for key, item in value.items()})
    if isinstance(value, (list, tuple)):
        return tuple(_freeze_result_value(item) for item in value)
    if isinstance(value, set):
        return frozenset(_freeze_result_value(item) for item in value)
    return value


def _thaw_result_value(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(key): _thaw_result_value(item) for key, item in value.items()}
    if isinstance(value, tuple):
        return [_thaw_result_value(item) for item in value]
    if isinstance(value, frozenset):
        return sorted((_thaw_result_value(item) for item in value), key=repr)
    if isinstance(value, ExcelError):
        return value.value
    if hasattr(value, "__dataclass_fields__"):
        return {field_name: _thaw_result_value(getattr(value, field_name)) for field_name in value.__dataclass_fields__}
    return value


@dataclass(frozen=True)
class FormulaSpec:
    expression: str
    operation: str = "sum"


@dataclass(frozen=True)
class UnitSpec:
    units: str
    relative_tolerance: float = 1e-9


def _excel_type(value: Any) -> str:
    if isinstance(value, ExcelError):
        return "error"
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str) and value.startswith("#"):
        return "error"
    return "string"


def _coerce(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str) and value.startswith("#"):
        return ExcelError(value)
    return value


def _range_cells(ws, cell_range: str) -> list[list[Any]]:
    return [list(row) for row in ws[cell_range]]


def _range_bounds(cell_range: str) -> tuple[int, int, int, int]:
    match = re.fullmatch(r"([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?", cell_range.upper())
    if not match:
        raise ToolbeltError(f"invalid cell range: {cell_range}")
    def col(value: str) -> int:
        result = 0
        for char in value:
            result = result * 26 + ord(char) - 64
        return result
    return col(match.group(1)), int(match.group(2)), col(match.group(3) or match.group(1)), int(match.group(4) or match.group(2))


def _formula_references(formula: str | None) -> tuple[tuple[str, str], ...]:
    if not formula:
        return ()
    matches = re.findall(r"(?:'([^']+)'|([A-Za-z0-9_ ]+))!\$?([A-Z]{1,3})\$?(\d+)", formula)
    return tuple(sorted({((quoted or bare).strip(), f"{column}{row}") for quoted, bare, column, row in matches}))


def _in_selection(sheet: str, selection: SheetSelection | None) -> bool:
    if selection is None:
        return True
    if selection.sheet_names:
        return sheet in selection.sheet_names
    selected_sheets = {selected_sheet for selected_sheet, _ in (*selection.ranges, *selection.tables)}
    return not selected_sheets or sheet in selected_sheets


def _implicit_xlsx_bounds(ws) -> tuple[int, int, int, int] | None:
    used_cells = [cell for cell in getattr(ws, "_cells", {}).values() if cell.value is not None]
    if not used_cells:
        return None
    return (
        1,
        1,
        max(cell.column for cell in used_cells),
        max(cell.row for cell in used_cells),
    )


def _check_cancel(cancel_check: Callable[[], bool] | None) -> None:
    if cancel_check and cancel_check():
        raise ToolbeltError("cancelled")


def _delimited_value(value: str) -> Any:
    value = value.strip()
    if not value:
        return None
    if value.startswith("#"):
        return ExcelError(value)
    if value.lower() in {"true", "false"}:
        return value.lower() == "true"
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


def _read_delimited(
    snapshot_or_source: SourceSnapshot | Path,
    selection: SheetSelection | None,
    limits: ToolbeltLimits,
    cancel_check: Callable[[], bool] | None,
    digest: str | None = None,
) -> WorkbookReadResult:
    if isinstance(snapshot_or_source, SourceSnapshot):
        source = Path(snapshot_or_source.source)
        digest = snapshot_or_source.sha256
        content = snapshot_or_source.content
    else:
        source = Path(snapshot_or_source)
        digest = digest or _sha256_file(source)
        content = source.read_bytes()
    sheet = source.stem
    if selection and selection.sheet_names and sheet not in selection.sheet_names:
        return WorkbookReadResult(str(source), digest, (), (), ())
    delimiter = "\t" if source.suffix.lower() == ".tsv" else ","
    selected_range = next((cell_range for selected_sheet, cell_range in (selection.ranges if selection else ()) if selected_sheet == sheet), None)
    range_bounds = _range_bounds(selected_range) if selected_range else None
    row_start = range_bounds[1] if range_bounds else 1
    col_start = range_bounds[0] if range_bounds else 1
    raw_rows: list[tuple[Any, ...]] = []
    cell_count = 0
    selected_row_count = 0
    with io.StringIO(content.decode("utf-8-sig"), newline="") as handle:
        for row_number, row in enumerate(csv.reader(handle, delimiter=delimiter), 1):
            _check_cancel(cancel_check)
            if range_bounds and row_number > range_bounds[3]:
                break
            if range_bounds and row_number < range_bounds[1]:
                continue
            selected_row_count += 1
            if selected_row_count > limits.max_rows:
                raise ToolbeltError("workbook selection exceeds row limit")
            values = tuple(_delimited_value(value) for value in row)
            if range_bounds:
                values = values[range_bounds[0] - 1:range_bounds[2]]
            cell_count += len(values)
            if cell_count > limits.max_cells:
                raise ToolbeltError("workbook selection exceeds cell limit")
            raw_rows.append(values)
    if selection and selection.ranges:
        ranges = [cell_range for selected_sheet, cell_range in selection.ranges if selected_sheet == sheet]
        if ranges:
            if selected_range is None:
                raise ToolbeltError("invalid CSV selection")
    if selection and selection.tables and not any(selected_sheet == sheet and table_name == source.stem for selected_sheet, table_name in selection.tables):
        return WorkbookReadResult(str(source), digest, (), (), ())
    cell_count = sum(len(row) for row in raw_rows)
    if cell_count > limits.max_cells:
        raise ToolbeltError("workbook selection exceeds cell limit")
    hidden_rows: tuple[int, ...] = ()
    summary = SheetSummary(sheet, "visible", len(raw_rows), max((len(row) for row in raw_rows), default=0), hidden_rows, (), (source.stem,))
    measurements: list[TypedMeasurement] = []
    provenance_rows: list[tuple[CellProvenance, ...]] = []
    for row_number, row in enumerate(raw_rows, 1):
        _check_cancel(cancel_check)
        row_provenance = []
        for column_number, value in enumerate(row, 1):
            coordinate = f"{openpyxl.utils.get_column_letter(col_start + column_number - 1)}{row_start + row_number - 1}"
            locator = CellProvenance(digest, sheet, coordinate, f"{sheet}!{coordinate}")
            measurements.append(TypedMeasurement(f"{digest[:12]}:{sheet}:{coordinate}", value, _excel_type(value), None, None, FormulaStatus.VALUE, locator, trusted=True))
            row_provenance.append(locator)
        provenance_rows.append(tuple(row_provenance))
    header_rows = min(1, len(raw_rows))
    if len(raw_rows) > 1 and sum(value not in (None, "") for value in raw_rows[0]) <= 1 and sum(value not in (None, "") for value in raw_rows[1]) > 1:
        header_rows = 2
    headers = tuple(tuple("" if value is None else str(value) for value in raw_rows[index]) for index in range(header_rows))
    source_range = f"{openpyxl.utils.get_column_letter(col_start)}{row_start}:{openpyxl.utils.get_column_letter(col_start + summary.max_column - 1)}{row_start + summary.max_row - 1}" if raw_rows else ""
    table = TypedTable(f"{digest[:12]}:{sheet}:{source_range}", source.stem, sheet, headers, tuple(raw_rows[header_rows:]), tuple(provenance_rows[header_rows:]), False, source_range)
    return WorkbookReadResult(str(source), digest, (summary,), (table,), tuple(measurements))


def _preflight_xlsx_archive(snapshot: SourceSnapshot, selection: SheetSelection | None, limits: ToolbeltLimits, cancel_check: Callable[[], bool] | None) -> None:
    namespaces = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkg": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    with zipfile.ZipFile(io.BytesIO(snapshot.content)) as archive:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships.findall("pkg:Relationship", namespaces)}
        total_row_count = 0
        total_cell_count = 0
        total_merged_range_count = 0
        total_merged_cell_count = 0
        for sheet in workbook_root.findall("main:sheets/main:sheet", namespaces):
            title = sheet.attrib["name"]
            selected_sheet = not selection or not selection.sheet_names or title in selection.sheet_names
            target = targets.get(sheet.attrib.get(f"{{{namespaces['rel']}}}id"))
            if not target:
                continue
            normalized_target = target.lstrip("/")
            sheet_path = posixpath.normpath(normalized_target if normalized_target.startswith("xl/") else posixpath.join("xl", normalized_target))
            selected_ranges = [cell_range for selected_sheet, cell_range in (selection.ranges if selection else ()) if selected_sheet == title]
            bounds = [_range_bounds(cell_range) for cell_range in selected_ranges]
            if selected_sheet:
                for range_bounds in bounds:
                    if range_bounds[3] - range_bounds[1] + 1 > limits.max_rows:
                        raise ToolbeltError("workbook selection exceeds row limit")
                    if (range_bounds[3] - range_bounds[1] + 1) * (range_bounds[2] - range_bounds[0] + 1) > limits.max_cells:
                        raise ToolbeltError("workbook selection exceeds cell limit")
            with archive.open(sheet_path) as handle:
                for _, element in ET.iterparse(handle, events=("end",)):
                    _check_cancel(cancel_check)
                    if element.tag == f"{{{namespaces['main']}}}row":
                        if selected_sheet:
                            total_row_count += 1
                        if selected_sheet and total_row_count > limits.max_rows:
                            raise ToolbeltError("workbook allocation exceeds row limit")
                    elif element.tag == f"{{{namespaces['main']}}}c":
                        coordinate = element.attrib.get("r", "")
                        match = re.fullmatch(r"([A-Z]+)(\d+)", coordinate.upper())
                        if match:
                            total_cell_count += 1
                            if total_cell_count > limits.max_cells:
                                raise ToolbeltError("workbook allocation exceeds cell limit")
                    elif element.tag == f"{{{namespaces['main']}}}mergeCell":
                        total_merged_range_count += 1
                        if total_merged_range_count > limits.max_merged_ranges:
                            raise ToolbeltError("workbook merged range limit exceeded")
                        merged_ref = element.attrib.get("ref")
                        if merged_ref:
                            start_col, start_row, end_col, end_row = _range_bounds(merged_ref)
                            area = (end_col - start_col + 1) * (end_row - start_row + 1)
                            if area > limits.max_merged_cells_per_range:
                                raise ToolbeltError("workbook merged area exceeds per-range limit")
                            total_merged_cell_count += area
                            if total_merged_cell_count > limits.max_total_merged_cells:
                                raise ToolbeltError("workbook merged area exceeds total limit")
                            if total_cell_count + total_merged_cell_count > limits.max_cells:
                                raise ToolbeltError("workbook allocation exceeds cell limit")
                    element.clear()


def spreadsheet_read(
    source: str,
    selection: Mapping[str, Any] | None = None,
    limits: Mapping[str, Any] | None = None,
    *,
    cancel_check: Callable[[], bool] | None = None,
    _cache: Any = None,
    _snapshot: SourceSnapshot | None = None,
    _workbook_context: WorkbookContext | None = None,
    _authority_index: AuthorityIndex | None = None,
    _stats: WorkbookStats | None = None,
) -> WorkbookReadResult:
    """Read a bounded workbook using JSON-compatible source, selection, and limits."""
    if _workbook_context is not None:
        return _workbook_context.read(source, selection, limits, cancel_check=cancel_check)
    selection = _coerce_sheet_selection(selection)
    limits = _coerce_toolbelt_limits(limits)
    if limits.max_cells > _AUTHORITATIVE_RECORD_LIMIT:
        raise ToolbeltError("workbook cell limit exceeds authoritative retention capacity")
    source = Path(source)
    _check_cancel(cancel_check)
    if _snapshot is None:
        snapshot = SourceSnapshot.capture(source, cancel_check=cancel_check, max_bytes=limits.max_bytes)
        if _stats is not None:
            _stats.capture_count += 1
            _stats.capture_hash_count += 1
            _stats.snapshot_bytes += snapshot.byte_count
    else:
        snapshot = _snapshot
    if Path(snapshot.source).resolve() != source.resolve():
        raise ToolbeltError("source snapshot does not match requested source")
    _check_cancel(cancel_check)
    cache_key = None
    digest = None
    if _cache is not None:
        cache_key = _cache.key(source, selection, limits, snapshot=snapshot)
        cached_result = _cache.get(cache_key)
        if cached_result is not None:
            if _stats is not None:
                _stats.cache_hits += 1
            cached_result = _with_read_source(cached_result, source)
            _publish_read_result(cached_result, authority_index=_authority_index)
            return cached_result
    if snapshot.byte_count > limits.max_bytes:
        raise ToolbeltError("workbook exceeds byte limit")
    if source.suffix.lower() in {".csv", ".tsv"}:
        result = _read_delimited(snapshot, selection, limits, cancel_check)
        return _publish_read_result(
            result,
            cache=_cache,
            cache_key=cache_key,
            authority_index=_authority_index,
        )
    if source.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise ToolbeltError("unsupported spreadsheet format")
    try:
        with zipfile.ZipFile(io.BytesIO(snapshot.content)) as archive:
            expanded_bytes = sum(item.file_size for item in archive.infolist())
    except zipfile.BadZipFile as error:
        raise ToolbeltError("invalid workbook archive") from error
    if expanded_bytes > limits.max_expanded_bytes:
        raise ToolbeltError("workbook expanded content exceeds byte limit")
    _preflight_xlsx_archive(snapshot, selection, limits, cancel_check)
    if _stats is not None:
        _stats.xlsx_parse_count += 1
    digest = snapshot.sha256
    workbook = None
    cached = None
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(snapshot.content), data_only=False, read_only=False)
        if _stats is not None:
            _stats.openpyxl_load_count += 1
        cached = openpyxl.load_workbook(io.BytesIO(snapshot.content), data_only=True, read_only=True)
        if _stats is not None:
            _stats.openpyxl_load_count += 1
        sheets: list[SheetSummary] = []
        tables: list[TypedTable] = []
        measurements: list[TypedMeasurement] = []
        cell_count = 0
        for ws in workbook.worksheets:
            _check_cancel(cancel_check)
            if ws.sheet_state != "visible" and not (selection and selection.include_hidden):
                continue
            hidden_rows = tuple(row for row, dimension in ws.row_dimensions.items() if dimension.hidden)
            summary = SheetSummary(ws.title, ws.sheet_state, ws.max_row, ws.max_column, hidden_rows, tuple(str(item) for item in ws.merged_cells.ranges), tuple(ws.tables))
            sheets.append(summary)
            if not _in_selection(ws.title, selection):
                continue
            selected_tables = [table_name for sheet, table_name in (selection.tables if selection else ()) if sheet == ws.title]
            if selected_tables:
                selected_ranges = [ws.tables[name].ref for name in selected_tables if name in ws.tables]
            else:
                selected_ranges = [cell_range for sheet, cell_range in (selection.ranges if selection else ()) if sheet == ws.title]
            if not selected_ranges:
                implicit_bounds = _implicit_xlsx_bounds(ws)
                selected_ranges = [
                    f"{openpyxl.utils.get_column_letter(implicit_bounds[0])}{implicit_bounds[1]}:"
                    f"{openpyxl.utils.get_column_letter(implicit_bounds[2])}{implicit_bounds[3]}"
                ] if implicit_bounds else []
            for source_range in selected_ranges:
                min_col, min_row, max_col, max_row = _range_bounds(source_range)
                if max_row - min_row + 1 > limits.max_rows:
                    raise ToolbeltError("workbook selection exceeds row limit")
                rows: list[tuple[Any, ...]] = []
                row_provenance: list[tuple[CellProvenance, ...]] = []
                row_measurements: list[tuple[TypedMeasurement, ...]] = []
                source_rows = ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
                cached_rows = cached[ws.title].iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
                for source_row, cached_row in zip_longest(source_rows, cached_rows, fillvalue=()):
                    _check_cancel(cancel_check)
                    if not source_row:
                        raise ToolbeltError("cached workbook contains rows without source values")
                    row = source_row[0].row
                    if row in hidden_rows and not (selection and selection.include_hidden_rows):
                        continue
                    values: list[Any] = []
                    provenance: list[CellProvenance] = []
                    row_cells: list[TypedMeasurement] = []
                    for cell, cached_cell in zip_longest(source_row, cached_row, fillvalue=None):
                        _check_cancel(cancel_check)
                        formula = cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None
                        cached_value = _coerce(cached_cell.value) if formula and cached_cell is not None else None
                        raw_value = cached_value if formula and cached_value is not None else cell.value
                        value = _coerce(raw_value)
                        if formula:
                            if isinstance(cached_value, ExcelError):
                                status = FormulaStatus.ERROR
                            elif cached_value is None:
                                status = FormulaStatus.FORMULA_WITHOUT_CACHE
                            else:
                                calculation = getattr(workbook, "calculation", None)
                                stale = bool(
                                    calculation
                                    and (
                                        getattr(calculation, "calcMode", None) == "manual"
                                        or getattr(calculation, "fullCalcOnLoad", False)
                                        or getattr(calculation, "forceFullCalc", False)
                                    )
                                )
                                status = FormulaStatus.FORMULA_STALE_OR_UNKNOWN if stale else FormulaStatus.FORMULA_WITH_CACHE
                        else:
                            status = FormulaStatus.ERROR if isinstance(value, ExcelError) else FormulaStatus.VALUE
                        locator = CellProvenance(digest, ws.title, cell.coordinate, source_range)
                        measurement_id = f"{digest[:12]}:{ws.title}:{cell.coordinate}"
                        referenced_cells = _formula_references(formula)
                        references = tuple(sorted({sheet for sheet, _ in referenced_cells}))
                        measurements.append(TypedMeasurement(measurement_id, value, _excel_type(raw_value), formula, cached_value, status, locator, references, referenced_cells, True))
                        row_cells.append(measurements[-1])
                        values.append(value)
                        provenance.append(locator)
                        cell_count += 1
                        if cell_count > limits.max_cells:
                            raise ToolbeltError("workbook selection exceeds cell limit")
                    rows.append(tuple(values))
                    row_provenance.append(tuple(provenance))
                    row_measurements.append(tuple(row_cells))
                nonempty_counts = [sum(value not in (None, "") for value in row) for row in rows]
                first_dense_row = next((index for index, count in enumerate(nonempty_counts) if count > 1), None)
                if first_dense_row is not None and first_dense_row > 0 and any(count == 0 for count in nonempty_counts[:first_dense_row]):
                    rows = rows[first_dense_row:]
                    row_provenance = row_provenance[first_dense_row:]
                    row_measurements = row_measurements[first_dense_row:]
                    nonempty_counts = nonempty_counts[first_dense_row:]
                header_rows = min(1, len(rows))
                if len(rows) > 1 and nonempty_counts[0] <= 1 and nonempty_counts[1] > 1:
                    header_rows = 2
                headers = tuple(tuple("" if value is None else str(value) for value in rows[index]) for index in range(header_rows))
                table_id = f"{digest[:12]}:{ws.title}:{source_range}"
                related = tuple(
                    measurement.provenance
                    for measurement in measurements
                    if measurement.formula and measurement.provenance.sheet == ws.title and measurement.provenance.range == source_range
                )
                tables.append(TypedTable(table_id, selected_tables[0] if selected_tables else source_range, ws.title, headers, tuple(rows[header_rows:]), tuple(row_provenance[header_rows:]), ws.sheet_state != "visible", source_range, related, tuple(row_measurements[header_rows:])))
        finalized_tables = []
        for table in tables:
            related = list(table.related_provenance)
            for measurement in measurements:
                if measurement.provenance.range != table.source_range or measurement.provenance.sheet != table.sheet:
                    continue
                related.extend(CellProvenance(digest, sheet, coordinate, f"{sheet}!{coordinate}") for sheet, coordinate in measurement.cross_sheet_cells)
            finalized_tables.append(replace(table, related_provenance=tuple(dict.fromkeys(related))))
        result = WorkbookReadResult(str(source), digest, tuple(sheets), tuple(finalized_tables), tuple(measurements))
        return _publish_read_result(
            result,
            cache=_cache,
            cache_key=cache_key,
            authority_index=_authority_index,
        )
    finally:
        if cached is not None:
            cached.close()
        if workbook is not None:
            workbook.close()


def _rows(table: TypedTable) -> list[dict[str, Any]]:
    columns = table.columns
    return [dict(zip(columns, row, strict=False)) for row in table.rows]


def _ensure_homogeneous_units(rows: Sequence[dict[str, Any]], table: TypedTable, value_column: str) -> None:
    unit_column = next((column for column in table.columns if column.strip().lower() == "unit"), None)
    if unit_column is None or not rows:
        return
    contributing_rows = [row for row in rows if row.get(value_column) is not None]
    if not contributing_rows:
        return
    units = {row.get(unit_column) for row in contributing_rows}
    if None in units or "" in units:
        raise ToolbeltError("unit unavailable for numeric analysis")
    if len(units) > 1:
        raise ToolbeltError("unit mismatch across numeric analysis rows")


def data_analysis(
    table: TypedTable | Mapping[str, Any],
    operation: AnalysisOperation | Mapping[str, Any],
    limits: Mapping[str, Any] | None = None,
    *,
    cancel_check: Callable[[], bool] | None = None,
    _cache: Any = None,
    _workbook_context: WorkbookContext | None = None,
    _authority_index: AuthorityIndex | None = None,
    _stats: WorkbookStats | None = None,
) -> AnalysisResult:
    if _workbook_context is not None:
        return _workbook_context.analyze(table, operation, limits, cancel_check=cancel_check)
    table = _coerce_typed_table(table)
    operation = _coerce_analysis_operation(operation)
    limits = _coerce_toolbelt_limits(limits)
    _check_cancel(cancel_check)
    analysis_cache_key = None
    if _cache is not None and hasattr(_cache, "analysis_key"):
        analysis_cache_key = _cache.analysis_key(table, operation, limits)
        _check_cancel(cancel_check)
        cached_result = _cache.get_analysis(analysis_cache_key)
        if cached_result is not None:
            if _stats is not None:
                _stats.analysis_cache_hits += 1
            return cached_result
    rows = _rows(table)
    kind = operation.kind.lower()
    warnings: list[str] = []
    if operation.filters:
        for column, comparator, expected in operation.filters:
            if comparator == "eq":
                rows = [row for row in rows if row.get(column) == expected]
            elif comparator == "neq":
                rows = [row for row in rows if row.get(column) != expected]
            elif comparator == "gt":
                rows = [row for row in rows if row.get(column) is not None and row.get(column) > expected]
            elif comparator == "gte":
                rows = [row for row in rows if row.get(column) is not None and row.get(column) >= expected]
            elif comparator == "lt":
                rows = [row for row in rows if row.get(column) is not None and row.get(column) < expected]
            elif comparator == "lte":
                rows = [row for row in rows if row.get(column) is not None and row.get(column) <= expected]
            else:
                raise ToolbeltError(f"unsupported filter comparator: {comparator}")
    if kind == "join":
        if operation.join_table is None or len(operation.join_keys) != 2:
            raise ToolbeltError("join requires join_table and two join keys")
        right_rows = _rows(operation.join_table)
        left_key, right_key = operation.join_keys
        right_index: dict[Any, list[dict[str, Any]]] = {}
        for row in right_rows:
            right_index.setdefault(row.get(right_key), []).append(row)
        duplicate_keys = sorted((key for key, matches in right_index.items() if len(matches) > 1), key=repr)
        if duplicate_keys and not operation.many_to_many:
            raise ToolbeltError("duplicate join keys require many_to_many=True")
        result = []
        for row in rows:
            matches = right_index.get(row.get(left_key), [])
            for match in matches:
                result.append({**row, **{f"right.{name}": value for name, value in match.items()}})
                if len(result) > limits.max_result_rows:
                    raise ToolbeltError("analysis result exceeds row limit")
        if duplicate_keys:
            warnings.append("many-to-many join expanded duplicate right-side keys")
    elif kind in {"nulls", "missing"}:
        result = [{"column": column, "missing": sum(row.get(column) is None for row in rows)} for column in table.columns]
    elif kind == "duplicates":
        seen: dict[tuple[Any, ...], int] = {}
        for row in rows:
            key = tuple(row.get(column) for column in (operation.key_columns or table.columns))
            seen[key] = seen.get(key, 0) + 1
        result = [{"key": key, "count": count} for key, count in sorted(seen.items(), key=lambda item: repr(item[0])) if count > 1]
    elif kind in {"aggregate", "sum", "mean"}:
        if not operation.value_column:
            raise ToolbeltError("aggregate requires value_column")
        _ensure_homogeneous_units(rows, table, operation.value_column)
        value_index = table.columns.index(operation.value_column)
        selected_measurements = [
            row[value_index]
            for row in table.row_measurements
            if len(row) > value_index
        ]
        if any(item.formula_status in {FormulaStatus.FORMULA_WITHOUT_CACHE, FormulaStatus.FORMULA_STALE_OR_UNKNOWN, FormulaStatus.ERROR} for item in selected_measurements):
            raise ToolbeltError("BLOCKED_RECALCULATION: formula input is not trustworthy")
        values = [row.get(operation.value_column) for row in rows if row.get(operation.value_column) is not None]
        if any(isinstance(value, ExcelError) for value in values):
            raise ToolbeltError("cannot aggregate Excel errors")
        total = sum(float(value) for value in values)
        result = [{"count": len(values), "sum": total, "mean": total / len(values) if values else None}]
        if len(values) != len(rows):
            warnings.append("null values reported and excluded from aggregate")
    elif kind in {"scenario", "sensitivity"}:
        if not operation.value_column:
            raise ToolbeltError("scenario analysis requires value_column")
        _ensure_homogeneous_units(rows, table, operation.value_column)
        base = sum(float(row[operation.value_column]) for row in rows if row.get(operation.value_column) is not None)
        if kind == "scenario":
            result = [{"scenario": name, "value": base * float(multiplier)} for name, multiplier in operation.scenarios]
        else:
            result = [{"parameter": name, "value": base * float(value)} for name, values in operation.sensitivity for value in values]
    elif kind == "compare":
        if len(operation.columns) != 2:
            raise ToolbeltError("compare requires two columns")
        left_column, right_column = operation.columns
        result = []
        for row in rows:
            left = row.get(left_column)
            right = row.get(right_column)
            result.append({
                "left": left,
                "right": right,
                "delta": None if left is None or right is None else right - left,
                "status": "INCOMPLETE_INPUT" if left is None or right is None else "COMPLETE",
            })
    else:
        raise ToolbeltError(f"unsupported analysis operation: {operation.kind}")
    if len(result) > limits.max_result_rows:
        raise ToolbeltError("analysis result exceeds row limit")
    _check_cancel(cancel_check)
    provenance = tuple(dict.fromkeys((*[locator for row in table.row_provenance for locator in row], *table.related_provenance)))
    if kind == "join" and operation.join_table is not None:
        provenance = tuple(dict.fromkeys((*provenance, *[locator for row in operation.join_table.row_provenance for locator in row], *operation.join_table.related_provenance)))
    key = hashlib.sha256(repr((operation, result, provenance)).encode()).hexdigest()
    analysis_result = AnalysisResult(operation.kind, tuple(result), provenance, key, tuple(warnings))
    if _cache is not None and analysis_cache_key is not None:
        _cache.put_analysis(analysis_cache_key, analysis_result)
    return analysis_result


def _numeric(value: Any, *, allow_legacy_prose: bool = False) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if isinstance(value, ExcelError) or value is None:
        raise ToolbeltError("calculation input is unavailable")
    if not allow_legacy_prose:
        raise ToolbeltError("calculation input is not structured numeric evidence")
    match = re.search(r"=([-+]?[0-9]+(?:\.[0-9]+)?)", str(value))
    if not match:
        raise ToolbeltError("calculation input is not numeric")
    return float(match.group(1))


def _coerce_cell_provenance(value: CellProvenance | Mapping[str, Any]) -> CellProvenance:
    if isinstance(value, CellProvenance):
        return value
    if not isinstance(value, Mapping):
        raise ToolbeltError("calculation measurement provenance is unavailable")
    return CellProvenance(
        workbook_sha256=str(value.get("workbook_sha256", "")),
        sheet=str(value.get("sheet", "")),
        coordinate=str(value.get("coordinate", "")),
        range=str(value.get("range", "")),
    )


def _coerce_typed_table(value: TypedTable | Mapping[str, Any]) -> TypedTable:
    if isinstance(value, TypedTable):
        return value
    if not isinstance(value, Mapping):
        raise ToolbeltError("analysis table must be structured")
    headers = tuple(tuple(str(item) for item in row) for row in value.get("headers", ()) or ())
    rows = tuple(tuple(_coerce(item) for item in row) for row in value.get("rows", ()) or ())
    row_provenance = tuple(
        tuple(_coerce_cell_provenance(item) for item in row)
        for row in value.get("row_provenance", ()) or ()
    )
    row_measurements = tuple(
        tuple(
            item
            for item in (_coerce_calculation_input(raw) for raw in row)
            if isinstance(item, TypedMeasurement)
        )
        for row in value.get("row_measurements", ()) or ()
    )
    related_provenance = tuple(_coerce_cell_provenance(item) for item in value.get("related_provenance", ()) or ())
    return TypedTable(
        id=str(value.get("id", "")),
        name=str(value.get("name", "")),
        sheet=str(value.get("sheet", "")),
        headers=headers,
        rows=rows,
        row_provenance=row_provenance,
        hidden=bool(value.get("hidden", False)),
        source_range=str(value.get("source_range", "")),
        related_provenance=related_provenance,
        row_measurements=row_measurements,
    )


def _coerce_analysis_operation(value: AnalysisOperation | Mapping[str, Any]) -> AnalysisOperation:
    if isinstance(value, AnalysisOperation):
        return value
    if not isinstance(value, Mapping):
        raise ToolbeltError("analysis operation must be structured")
    scenarios = tuple((str(item[0]), float(item[1])) for item in value.get("scenarios", ()) or ())
    sensitivity = tuple(
        (str(item[0]), tuple(float(entry) for entry in item[1]))
        for item in value.get("sensitivity", ()) or ()
    )
    filters = tuple(tuple(item) for item in value.get("filters", ()) or ())
    return AnalysisOperation(
        kind=str(value.get("kind", "")),
        columns=tuple(str(item) for item in value.get("columns", ()) or ()),
        key_columns=tuple(str(item) for item in value.get("key_columns", ()) or ()),
        value_column=str(value["value_column"]) if value.get("value_column") is not None else None,
        scenarios=scenarios,
        sensitivity=sensitivity,
        many_to_many=bool(value.get("many_to_many", False)),
        filters=filters,
        join_table=_coerce_typed_table(value["join_table"]) if value.get("join_table") is not None else None,
        join_keys=tuple(str(item) for item in value.get("join_keys", ()) or ()),
    )


def _coerce_calculation_input(value: Evidence | CalculationRecord | TypedMeasurement | Mapping[str, Any]) -> Evidence | CalculationRecord | TypedMeasurement:
    if isinstance(value, (Evidence, CalculationRecord, TypedMeasurement)):
        return value
    if not isinstance(value, Mapping):
        raise ToolbeltError("calculation input must be structured evidence")
    if "output" in value and "key" in value:
        key_value = value["key"]
        if not isinstance(key_value, Mapping):
            raise ToolbeltError("calculation record key is unavailable")
        return CalculationRecord(
            id=str(value.get("id", "")),
            key=MeasurementKey(**{name: key_value.get(name, "") for name in ("region", "measure", "period", "units", "scenario", "version")}),
            formula=str(value.get("formula", "")),
            input_ids=[str(item) for item in value.get("input_ids", ())],
            output=float(value["output"]),
            units=str(value.get("units", "")),
            source_locators=[str(item) for item in value.get("source_locators", ())],
            verification_status=(
                value.get("verification_status", CalculationVerificationStatus.NOT_VERIFIED)
                if isinstance(value.get("verification_status", CalculationVerificationStatus.NOT_VERIFIED), CalculationVerificationStatus)
                else CalculationVerificationStatus(str(value.get("verification_status", CalculationVerificationStatus.NOT_VERIFIED)))
            ),
        )
    if "value" in value and "provenance" in value:
        status = value.get("formula_status", FormulaStatus.VALUE)
        return TypedMeasurement(
            id=str(value.get("id", "")),
            value=value.get("value"),
            data_type=str(value.get("data_type", "")),
            formula=value.get("formula"),
            cached_value=value.get("cached_value"),
            formula_status=status if isinstance(status, FormulaStatus) else FormulaStatus(str(status)),
            provenance=_coerce_cell_provenance(value["provenance"]),
            cross_sheet_references=tuple(str(item) for item in value.get("cross_sheet_references", ())),
            cross_sheet_cells=tuple(tuple(str(part) for part in item) for item in value.get("cross_sheet_cells", ())),
            trusted=False,
        )
    if "excerpt" in value and "source_locator" in value:
        locator_value = value["source_locator"]
        if not isinstance(locator_value, Mapping):
            raise ToolbeltError("evidence source locator is unavailable")
        measurement_key = value.get("measurement_key")
        return Evidence(
            id=str(value.get("id", "")),
            source_locator=SourceLocator(**{name: str(locator_value.get(name, "")) for name in ("file_version", "sheet", "range", "units", "transformation", "calculation")}),
            version=str(value.get("version", "")),
            date=str(value.get("date", "")),
            excerpt=str(value.get("excerpt", "")),
            limitations=[str(item) for item in value.get("limitations", ())],
            measurement_key=MeasurementKey(**measurement_key) if isinstance(measurement_key, Mapping) else None,
        )
    raise ToolbeltError("calculation input must be structured evidence")


def _coerce_formula_spec(value: FormulaSpec | Mapping[str, Any]) -> FormulaSpec:
    if isinstance(value, FormulaSpec):
        return value
    if not isinstance(value, Mapping):
        raise ToolbeltError("calculation formula is unavailable")
    return FormulaSpec(str(value.get("expression", "")), str(value.get("operation", "sum")))


def _coerce_unit_spec(value: UnitSpec | Mapping[str, Any]) -> UnitSpec:
    if isinstance(value, UnitSpec):
        return value
    if not isinstance(value, Mapping):
        raise ToolbeltError("calculation units are unavailable")
    return UnitSpec(str(value.get("units", "")), float(value.get("relative_tolerance", 1e-9)))


def calculation(
    inputs: Sequence[Evidence | CalculationRecord | TypedMeasurement | Mapping[str, Any]],
    formula: FormulaSpec | Mapping[str, Any],
    units: UnitSpec | Mapping[str, Any],
    approved_workbook_sha256: str | None = None,
    *,
    _authority_index: AuthorityIndex | None = None,
) -> CalculationRecord:
    inputs = tuple(_coerce_calculation_input(item) for item in inputs)
    formula = _coerce_formula_spec(formula)
    units = _coerce_unit_spec(units)
    if any(isinstance(item, TypedMeasurement) and item.formula_status in {FormulaStatus.FORMULA_WITHOUT_CACHE, FormulaStatus.FORMULA_STALE_OR_UNKNOWN, FormulaStatus.ERROR} for item in inputs):
        raise ToolbeltError("BLOCKED_RECALCULATION: formula input is not trustworthy")
    if any(isinstance(item, Evidence) for item in inputs):
        raise ToolbeltError("BLOCKED_RECALCULATION: authoritative workbook measurement required")
    measurements = [item for item in inputs if isinstance(item, TypedMeasurement)]
    if measurements:
        if not approved_workbook_sha256:
            raise ToolbeltError("BLOCKED_RECALCULATION: workbook measurement provenance is untrusted; approved workbook hash is missing or mismatched")
        resolved_measurements = []
        for item in measurements:
            if item.provenance.workbook_sha256 != approved_workbook_sha256:
                raise ToolbeltError("BLOCKED_RECALCULATION: approved workbook hash is missing or mismatched")
            measurements_index = _authority_index.measurements if _authority_index is not None else _AUTHORITATIVE_MEASUREMENTS
            authoritative = measurements_index.get((approved_workbook_sha256, item.id, item.provenance.range))
            if authoritative is None or replace(item, trusted=False) != replace(authoritative, trusted=False):
                raise ToolbeltError("BLOCKED_RECALCULATION: workbook measurement provenance is untrusted")
            resolved_measurements.append(authoritative)
        resolved_iter = iter(resolved_measurements)
        inputs = tuple(next(resolved_iter) if isinstance(item, TypedMeasurement) else item for item in inputs)
    calculations = [item for item in inputs if isinstance(item, CalculationRecord)]
    if calculations:
        if not approved_workbook_sha256:
            raise ToolbeltError("BLOCKED_RECALCULATION: approved workbook hash is missing or mismatched")
        resolved_calculations = []
        for item in calculations:
            if item.verification_status != CalculationVerificationStatus.VERIFIED or not item.input_ids or not item.source_locators:
                raise ToolbeltError("BLOCKED_RECALCULATION: authoritative calculation ancestry is unavailable")
            if any(approved_workbook_sha256 not in locator for locator in item.source_locators):
                raise ToolbeltError("BLOCKED_RECALCULATION: approved workbook hash is missing or mismatched")
            calculations_index = _authority_index.calculations if _authority_index is not None else _AUTHORITATIVE_CALCULATIONS
            authoritative = calculations_index.get((approved_workbook_sha256, item.id))
            if authoritative is None or replace(item, verification_status=CalculationVerificationStatus.NOT_VERIFIED) != replace(authoritative, verification_status=CalculationVerificationStatus.NOT_VERIFIED):
                raise ToolbeltError("BLOCKED_RECALCULATION: authoritative calculation ancestry is unavailable")
            resolved_calculations.append(authoritative)
        resolved_iter = iter(resolved_calculations)
        inputs = tuple(next(resolved_iter) if isinstance(item, CalculationRecord) else item for item in inputs)
        measurements = [item for item in inputs if isinstance(item, TypedMeasurement)]
    workbook_hashes = {item.provenance.workbook_sha256 for item in measurements}
    if len(workbook_hashes) > 1:
        raise ToolbeltError("workbook hash mismatch")
    if measurements and (not approved_workbook_sha256 or workbook_hashes != {approved_workbook_sha256}):
        raise ToolbeltError("BLOCKED_RECALCULATION: approved workbook hash is missing or mismatched")
    values = [_numeric(item.value if isinstance(item, TypedMeasurement) else item.output, allow_legacy_prose=False) for item in inputs]
    input_units = {
        item.source_locator.units for item in inputs if isinstance(item, Evidence)
    } | {
        item.units for item in inputs if isinstance(item, CalculationRecord)
    }
    if input_units and any(value != units.units for value in input_units):
        raise ToolbeltError("unit mismatch")
    operations = {"sum": sum, "subtract": lambda values: values[0] - values[1], "multiply": lambda values: operator.mul(*values), "divide": lambda values: values[0] / values[1]}
    try:
        output = float(operations[formula.operation](values))
    except KeyError as error:
        raise ToolbeltError(f"unsupported formula operation: {formula.operation}") from error
    except ZeroDivisionError as error:
        raise ToolbeltError("divide by zero") from error
    locators = list(dict.fromkeys(
        [f"{item.provenance.workbook_sha256}:{item.provenance.sheet}!{item.provenance.coordinate}" for item in inputs if isinstance(item, TypedMeasurement)]
        + [locator for item in inputs if isinstance(item, CalculationRecord) for locator in item.source_locators]
    ))
    input_ids = [item.id for item in inputs]
    identity = (formula.expression, values, units.units, input_ids, locators)
    record_id = "calculation-" + hashlib.sha256(repr(identity).encode()).hexdigest()[:16]
    record = CalculationRecord(record_id, MeasurementKey("toolbelt", formula.expression, "current", units.units), formula.expression, input_ids, output, units.units, locators, CalculationVerificationStatus.VERIFIED)
    if approved_workbook_sha256:
        _remember_authoritative_calculation(record, approved_workbook_sha256, _authority_index)
    return record


def graph_context(graph: Any, record_id: str, limit: int = 50) -> dict[str, Any]:
    from foresight_os.graph.queries import get_context
    if limit < 1:
        raise ToolbeltError("graph limit must be positive")
    return get_context(graph, record_id, limit=min(limit, 500))


def domain_validation(value: Any) -> dict[str, Any]:
    if isinstance(value, Evidence):
        return {"valid": bool(value.id and value.excerpt and value.source_locator.file_version), "kind": "evidence"}
    if isinstance(value, Mapping) and "stage_id" in value:
        raise ToolbeltError("stage result payloads are not valid domain_validation inputs")
    if isinstance(value, dict) and {"evidence", "claims"}.issubset(value):
        validate_analysis(value)
        return {"valid": True, "kind": "analysis"}
    if isinstance(value, Mapping):
        allowed_fields = {item.name for item in fields(EvidenceCandidate)}
        unsupported_fields = sorted(set(value) - allowed_fields)
        if unsupported_fields:
            raise ToolbeltError(
                "unsupported evidence candidate fields: " + ", ".join(unsupported_fields)
            )
    try:
        candidate = validate_evidence_candidate(value)
    except TypeError as error:
        raise ToolbeltError(f"invalid evidence candidate: {error}") from error
    return {"valid": True, "kind": "candidate", "candidate_id": candidate.candidate_id}


def powerbi_project_read(source: str, limits: Any = None) -> Any:
    from foresight_os.analysis.powerbi import read_powerbi_project
    from foresight_os.analysis.powerbi import PowerBIReadLimits

    if isinstance(limits, Mapping):
        limits = PowerBIReadLimits(**{name: int(limits[name]) for name in ("max_files", "max_total_bytes", "max_file_bytes") if name in limits})
    return read_powerbi_project(source, limits)


def bind_local_toolbelt() -> dict[str, Any]:
    from foresight_os.runtime.capabilities import CapabilityBinding, CapabilityBindingKind, CapabilityOwnership

    return {
        name: CapabilityBinding(function, kind=CapabilityBindingKind.TOOL, ownership=CapabilityOwnership.BORROWED, fingerprint=f"{name}-v1")
        for name, function in {
            "spreadsheet_read": spreadsheet_read,
            "data_analysis": data_analysis,
            "calculation": calculation,
            "graph_context": graph_context,
            "domain_validation": domain_validation,
            "powerbi_project_read": powerbi_project_read,
        }.items()
    }
